#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
dc_update.py - Live-Teil des MiT Datacenter Cockpits.

Was das Skript macht:
  1. News   holt RSS-/Atom-Feeds, filtert auf Rechenzentrums-Relevanz (CH/FL),
            dedupliziert und schreibt sie ins Cockpit.
  2. Geo    geocodiert Adressen und Orte ueber Nominatim (OpenStreetMap),
            mit Cache auf der Platte. Kein API-Key noetig.
  3. Excel  liest den Anrufplan neu ein, damit Termine und Ergebnisse aktuell sind.
  4. Log    schreibt im Cockpit erfasste Anrufergebnisse zurueck in den Anrufplan.
  5. Serve  startet einen lokalen Webserver, damit das Cockpit die JSON-Dateien
            ueberhaupt laden darf (file:// blockiert fetch()).
  6. Embed  backt den aktuellen Datenstand fest in die HTML-Datei, damit die
            Einzeldatei auch ohne Server und ohne Netz stimmt.

Abhaengigkeiten: Standardbibliothek. openpyxl nur fuer die Excel-Funktionen.

Aufrufe:
  python3 dc_update.py --once            einmal alles aktualisieren
  python3 dc_update.py --watch 30        alle 30 Minuten aktualisieren
  python3 dc_update.py --serve 8777      Server starten (aktualisiert im Hintergrund)
  python3 dc_update.py --test-feeds      nur Feeds pruefen, nichts schreiben
  python3 dc_update.py --merge-log       Anrufergebnisse zurueck in die Excel
  python3 dc_update.py --embed           Datenstand in die HTML backen

Harte Regel aus CLAUDE.md: keine erfundenen Werte. Feeds, die nicht antworten,
werden im Cockpit als tot markiert statt stillschweigend weggelassen. Adressen,
die Nominatim nicht findet, erscheinen nicht auf der Karte und stehen in der
Liste der ungeloesten Orte.
"""

import argparse
import datetime as dt
import gzip
import hashlib
import io
import json
import os
import re
import ssl
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "dc_data.json")
GEOCACHE = os.path.join(HERE, "geo_cache.json")
LOGCSV = os.path.join(HERE, "anruf_log.csv")
HTML = os.path.join(HERE, "MiT_DC_Cockpit.html")
CONFIG = os.path.join(HERE, "config.json")
CONFIG_EXAMPLE = os.path.join(HERE, "config.example.json")

UA = "MiT-DC-Cockpit/1.0"

# ---------------------------------------------------------------------------
# Konfiguration
# ---------------------------------------------------------------------------
STANDARD_CONFIG = {
    "_hinweis": "Kopiere diese Datei nach config.json und trage deine Werte ein. config.json gehoert NICHT ins Git.",
    "kontakt_mail": "uecoez@mobilintime.com",
    "_kontakt_mail": "Nominatim verlangt eine Kontaktadresse im User-Agent. Wird nur an nominatim.openstreetmap.org gesendet.",
    "newsapi_key": "",
    "_newsapi_key": "Optional. Konto auf newsapi.org. Leer lassen, dann laufen nur die RSS-Feeds, die brauchen keinen Key.",
    "google_maps_key": "",
    "_google_maps_key": "Nicht noetig. Die Karte laeuft auf OpenStreetMap. Feld existiert nur, falls du spaeter auf Google wechseln willst.",
    "proxy": "",
    "_proxy": "Nur setzen, wenn der MiT-Firmenproxy Pflicht ist, Format http://proxy:8080",
    "max_news": 60,
    "news_tage": 45,
    "feeds": [
        {"name": "Netzwoche",            "url": "https://www.netzwoche.ch/rss.xml",                  "aktiv": True,  "geprueft": False},
        {"name": "Netzwoche News",       "url": "https://www.netzwoche.ch/news.rss",                 "aktiv": True,  "geprueft": False},
        {"name": "inside-it.ch",         "url": "https://www.inside-it.ch/rss",                      "aktiv": True,  "geprueft": False},
        {"name": "inside-it.ch alt",     "url": "https://www.inside-it.ch/de/rss",                   "aktiv": True,  "geprueft": False},
        {"name": "IT-Markt",             "url": "https://www.it-markt.ch/rss.xml",                   "aktiv": True,  "geprueft": False},
        {"name": "Datacenter Dynamics",  "url": "https://www.datacenterdynamics.com/en/rss/",        "aktiv": True,  "geprueft": False},
        {"name": "Datacenter Dynamics Atom", "url": "https://www.datacenterdynamics.com/en/atom/",   "aktiv": True,  "geprueft": False},
        {"name": "Data Center Knowledge", "url": "https://www.datacenterknowledge.com/rss.xml",      "aktiv": True,  "geprueft": False},
        {"name": "SRF News Wirtschaft",  "url": "https://www.srf.ch/news/bnf/rss/1926",              "aktiv": True,  "geprueft": False},
        {"name": "NZZ Technologie",      "url": "https://www.nzz.ch/technologie.rss",                "aktiv": True,  "geprueft": False},
        {"name": "Aargauer Zeitung Aargau", "url": "https://www.aargauerzeitung.ch/aargau.rss",      "aktiv": True,  "geprueft": False},
        {"name": "Volksblatt Liechtenstein", "url": "https://www.volksblatt.li/rss",                 "aktiv": True,  "geprueft": False}
    ],
    "_feeds": ("Die URLs sind Kandidaten und beim Ausliefern UNGEPRUEFT. Der erste Lauf testet sie, "
               "setzt geprueft=true und aktiv=false bei toten Feeds. Im Cockpit siehst du unter "
               "Quellen, welcher Feed liefert und welcher nicht. Tote Feeds ersetzt du dort, wo die "
               "Seite ihre Feeds listet, zum Beispiel netzwoche.ch/RSS-Feeds."),
    "schlagworte": [
        "rechenzentrum", "rechenzentren", "datacenter", "data center", "data centre",
        "colocation", "hyperscale", "serverfarm", "cloud-standort",
        "notstrom", "netzersatz", "nea", "usv", "lastbank", "lastbaenke",
        "netzqualit", "oberschwingung", "en 50160",
        "batteriespeicher", "bess", "grossspeicher", "speicherkraftwerk",
        "netzanschluss", "trafostation", "umspannwerk",
        "flexbase", "vantage", "green datacenter", "digital realty", "equinix",
        "northc", "stack infrastructure", "safe host", "vaultica", "gtr",
        "implenia", "erne", "hiag", "amstein", "burkhalter", "swisscom",
        "laufenburg", "beringen", "dielsdorf", "lupfig", "volketswil",
        "glattbrugg", "winterthur", "arlesheim", "gland",
        "liechtenstein", "schaan", "vaduz", "eschen", "ruggell", "balzers"
    ],
    "ch_fl_filter": True,
    "_ch_fl_filter": "true = nur Meldungen mit Schweiz- oder FL-Bezug. false = auch internationale DC-News."
}

CH_MARKER = ["schweiz", "swiss", "switzerland", "suisse", "svizzera", "liechtenstein",
             "zuerich", "zürich", "zurich", "genf", "geneva", "basel", "bern", "aargau",
             "waadt", "wallis", "tessin", "ticino", "luzern", "st. gallen", "sankt gallen",
             "schaffhausen", "winterthur", "laufenburg", "beringen", "vaduz", "schaan"]


def lade_config():
    cfg = json.loads(json.dumps(STANDARD_CONFIG))
    if os.path.exists(CONFIG):
        try:
            with open(CONFIG, encoding="utf-8") as f:
                cfg.update(json.load(f))
        except Exception as e:
            print("config.json nicht lesbar (%s), Standardwerte aktiv." % e, file=sys.stderr)
    else:
        if not os.path.exists(CONFIG_EXAMPLE):
            with open(CONFIG_EXAMPLE, "w", encoding="utf-8") as f:
                json.dump(STANDARD_CONFIG, f, ensure_ascii=False, indent=2)
    return cfg


def speichere_config(cfg):
    ablage = {k: v for k, v in cfg.items()}
    with open(CONFIG, "w", encoding="utf-8") as f:
        json.dump(ablage, f, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# HTTP
# ---------------------------------------------------------------------------
def hole(url, cfg, timeout=25):
    """Holt eine URL und gibt (bytes, fehlertext) zurueck. Wirft nicht."""
    req = urllib.request.Request(url, headers={
        "User-Agent": "%s (%s)" % (UA, cfg.get("kontakt_mail") or "kein Kontakt hinterlegt"),
        "Accept": "application/rss+xml, application/atom+xml, application/xml, text/xml, application/json, */*",
        "Accept-Encoding": "gzip",
        "Accept-Language": "de-CH,de;q=0.9,en;q=0.6",
    })
    opener_args = []
    if cfg.get("proxy"):
        opener_args.append(urllib.request.ProxyHandler({"http": cfg["proxy"], "https": cfg["proxy"]}))
    ctx = ssl.create_default_context()
    opener_args.append(urllib.request.HTTPSHandler(context=ctx))
    opener = urllib.request.build_opener(*opener_args)
    try:
        with opener.open(req, timeout=timeout) as r:
            roh = r.read()
            if r.headers.get("Content-Encoding") == "gzip":
                try:
                    roh = gzip.decompress(roh)
                except Exception:
                    pass
            return roh, ""
    except urllib.error.HTTPError as e:
        return None, "HTTP %s" % e.code
    except urllib.error.URLError as e:
        return None, "nicht erreichbar: %s" % (e.reason,)
    except Exception as e:
        return None, "%s: %s" % (type(e).__name__, e)


# ---------------------------------------------------------------------------
# Feeds
# ---------------------------------------------------------------------------
def _text(el):
    return (el.text or "").strip() if el is not None else ""


def _strip_html(s):
    s = re.sub(r"<[^>]+>", " ", s or "")
    s = (s.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<")
          .replace("&gt;", ">").replace("&quot;", '"').replace("&#39;", "'"))
    return re.sub(r"\s+", " ", s).strip()


def _datum(s):
    if not s:
        return ""
    s = s.strip()
    for fmt in ("%a, %d %b %Y %H:%M:%S %z", "%a, %d %b %Y %H:%M:%S %Z",
                "%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%dT%H:%M:%SZ", "%Y-%m-%dT%H:%M:%S",
                "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            d = dt.datetime.strptime(s.replace("GMT", "+0000"), fmt)
            return d.date().isoformat()
        except ValueError:
            continue
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    return "%s-%s-%s" % m.groups() if m else ""


def parse_feed(roh):
    """RSS und Atom. Gibt Liste von dicts zurueck."""
    eintraege = []
    try:
        wurzel = ET.fromstring(roh)
    except ET.ParseError:
        try:
            wurzel = ET.fromstring(roh.decode("utf-8", "ignore").encode("utf-8"))
        except Exception:
            return eintraege
    ns = {"atom": "http://www.w3.org/2005/Atom", "dc": "http://purl.org/dc/elements/1.1/"}

    for item in wurzel.iter():
        tag = item.tag.split("}")[-1]
        if tag not in ("item", "entry"):
            continue
        titel = _text(item.find("title")) or _text(item.find("atom:title", ns))
        link = _text(item.find("link")) or _text(item.find("atom:link", ns))
        if not link:
            le = item.find("atom:link", ns) or item.find("link")
            if le is not None:
                link = le.get("href", "")
        datum = (_text(item.find("pubDate")) or _text(item.find("published"))
                 or _text(item.find("updated")) or _text(item.find("dc:date", ns))
                 or _text(item.find("atom:published", ns)) or _text(item.find("atom:updated", ns)))
        text = (_text(item.find("description")) or _text(item.find("summary"))
                or _text(item.find("atom:summary", ns)))
        if titel:
            eintraege.append(dict(titel=_strip_html(titel), link=link.strip(),
                                  datum=_datum(datum), text=_strip_html(text)[:400]))
    return eintraege


def relevant(e, cfg):
    blob = ("%s %s" % (e["titel"], e["text"])).lower()
    treffer = [s for s in cfg["schlagworte"] if s in blob]
    if not treffer:
        return None
    if cfg.get("ch_fl_filter"):
        dc_begriff = any(s in blob for s in ("rechenzentrum", "rechenzentren", "datacenter",
                                             "data center", "data centre", "colocation", "hyperscale"))
        ch_bezug = any(m in blob for m in CH_MARKER)
        # Firmen- und Ortsnamen aus der Pipeline zaehlen selbst als CH-Bezug
        eigene = any(s in blob for s in ("flexbase", "green datacenter", "implenia", "erne", "hiag",
                                         "swisscom", "laufenburg", "beringen", "dielsdorf", "lupfig",
                                         "volketswil", "glattbrugg", "arlesheim", "gland", "vaultica"))
        if not (ch_bezug or eigene) and not (dc_begriff and "aggreko" in blob):
            return None
    return treffer


def hole_news(cfg, daten):
    news, status = [], []
    gesehen = set()
    grenze = (dt.date.today() - dt.timedelta(days=int(cfg.get("news_tage", 45)))).isoformat()

    for feed in cfg["feeds"]:
        if not feed.get("aktiv", True):
            status.append(dict(name=feed["name"], url=feed["url"], zustand="deaktiviert",
                               detail=feed.get("detail", "beim Test tot"), anzahl=0,
                               geprueft=dt.datetime.now().isoformat(timespec="seconds")))
            continue
        roh, fehler = hole(feed["url"], cfg)
        if roh is None:
            feed["aktiv"] = False
            feed["geprueft"] = True
            feed["detail"] = fehler
            status.append(dict(name=feed["name"], url=feed["url"], zustand="tot", detail=fehler,
                               anzahl=0, geprueft=dt.datetime.now().isoformat(timespec="seconds")))
            continue
        eintraege = parse_feed(roh)
        if not eintraege:
            feed["aktiv"] = False
            feed["geprueft"] = True
            feed["detail"] = "antwortet, aber kein RSS/Atom"
            status.append(dict(name=feed["name"], url=feed["url"], zustand="kein Feed",
                               detail="antwortet, aber kein RSS/Atom", anzahl=0,
                               geprueft=dt.datetime.now().isoformat(timespec="seconds")))
            continue

        feed["geprueft"] = True
        feed["detail"] = ""
        treffer_zahl = 0
        for e in eintraege:
            schlag = relevant(e, cfg)
            if not schlag:
                continue
            if e["datum"] and e["datum"] < grenze:
                continue
            schluessel = hashlib.sha1((e["link"] or e["titel"]).encode("utf-8")).hexdigest()[:16]
            if schluessel in gesehen:
                continue
            gesehen.add(schluessel)
            treffer_zahl += 1
            news.append(dict(id=schluessel, titel=e["titel"], link=e["link"], datum=e["datum"],
                             text=e["text"], quelle=feed["name"], schlagworte=schlag[:6]))
        status.append(dict(name=feed["name"], url=feed["url"], zustand="ok",
                           detail="%d Eintraege, %d relevant" % (len(eintraege), treffer_zahl),
                           anzahl=treffer_zahl, geprueft=dt.datetime.now().isoformat(timespec="seconds")))

    news.sort(key=lambda n: (n["datum"] or "0000-00-00"), reverse=True)
    daten["news"] = news[:int(cfg.get("max_news", 60))]
    daten["quellen_status"] = status
    return len(daten["news"]), sum(1 for s in status if s["zustand"] == "ok")


# ---------------------------------------------------------------------------
# Geocoding
# ---------------------------------------------------------------------------
def lade_geocache():
    if os.path.exists(GEOCACHE):
        try:
            with open(GEOCACHE, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}


def geocode(daten, cfg):
    cache = lade_geocache()
    offen = []

    bereits = set()

    def merke(schluessel, genauigkeit):
        # Netzfehler werden nicht dauerhaft gecacht, die Adresse kommt beim
        # naechsten Lauf wieder dran. Nur echte Negativtreffer bleiben stehen.
        if not schluessel or schluessel in bereits:
            return
        eintrag = cache.get(schluessel)
        if eintrag and ("lat" in eintrag or eintrag.get("fehler") == "nicht gefunden"):
            return
        bereits.add(schluessel)
        offen.append((schluessel, genauigkeit))

    for p in daten.get("pipeline", []):
        merke(p.get("ort"), "ort")
    for b in daten.get("bestand", []):
        merke(b.get("ort"), "ort")
    for l in daten.get("liechtenstein", []):
        merke(l.get("adresse"), "adresse")

    neu = 0
    nicht_erreicht = []
    for schluessel, genauigkeit in offen:
        url = ("https://nominatim.openstreetmap.org/search?"
               + urllib.parse.urlencode({"q": schluessel, "format": "json", "limit": 1}))
        roh, fehler = hole(url, cfg, timeout=20)
        if roh is None:
            # nicht in den Cache: beim naechsten Lauf erneut versuchen
            nicht_erreicht.append(schluessel)
            print("  geo: %s -> %s" % (schluessel, fehler))
            if len(nicht_erreicht) >= 3 and len(nicht_erreicht) == len(offen[:len(nicht_erreicht)]):
                print("  geo: Nominatim antwortet nicht, Rest des Laufs uebersprungen.")
                break
        else:
            try:
                treffer = json.loads(roh.decode("utf-8"))
            except Exception:
                treffer = []
            if treffer:
                cache[schluessel] = dict(lat=float(treffer[0]["lat"]), lon=float(treffer[0]["lon"]),
                                         genauigkeit=genauigkeit,
                                         anzeige=treffer[0].get("display_name", "")[:120])
                neu += 1
            else:
                cache[schluessel] = dict(fehler="nicht gefunden")
        time.sleep(1.1)          # Nominatim-Nutzungsregel: maximal 1 Anfrage pro Sekunde

    with open(GEOCACHE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=1)
    daten["geo"] = cache
    daten["geo_offen"] = ([k for k, v in cache.items() if "lat" not in v]
                          + ["%s (nicht erreicht, naechster Lauf)" % s for s in nicht_erreicht])
    return neu, len(daten["geo_offen"])


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def merge_log():
    """Schreibt anruf_log.csv zurueck in den Tab Anrufplan.
    Laedt die Mappe OHNE data_only, sonst gehen Formeln verloren (CLAUDE.md)."""
    try:
        import openpyxl
    except ImportError:
        print("openpyxl fehlt. pip install openpyxl", file=sys.stderr)
        return 1
    import csv
    xlsx = os.path.join(HERE, "daten", "MiT_DC_Schweiz_Kontakte_Maerz2026.xlsx")
    if not os.path.exists(LOGCSV):
        print("Keine anruf_log.csv. Im Cockpit unter Log exportieren.", file=sys.stderr)
        return 1
    if not os.path.exists(xlsx):
        print("Excel nicht gefunden: %s" % xlsx, file=sys.stderr)
        return 1

    with open(LOGCSV, encoding="utf-8-sig", newline="") as f:
        zeilen = list(csv.DictReader(f, delimiter=";"))
    if not zeilen:
        print("anruf_log.csv ist leer.")
        return 0

    wb = openpyxl.load_workbook(xlsx)          # bewusst ohne data_only
    ws = wb["Anrufplan"]
    sicherung = xlsx.replace(".xlsx", "_backup_%s.xlsx" % dt.datetime.now().strftime("%Y%m%d_%H%M%S"))
    wb.save(sicherung)

    nach_name = {}
    for r in range(6, ws.max_row + 1):
        n = ws.cell(row=r, column=4).value
        if n:
            nach_name[str(n).strip().lower()] = r

    geschrieben = 0
    for z in zeilen:
        r = nach_name.get((z.get("name") or "").strip().lower())
        if not r:
            print("  nicht im Anrufplan gefunden: %s" % z.get("name"))
            continue
        ws.cell(row=r, column=11).value = z.get("erreicht", "")        # Erreicht
        ws.cell(row=r, column=12).value = z.get("ergebnis", "")        # Ergebnis
        ws.cell(row=r, column=13).value = z.get("naechster_schritt", "")
        ws.cell(row=r, column=14).value = z.get("wiedervorlage", "")
        geschrieben += 1

    wb.save(xlsx)
    print("Anrufplan aktualisiert: %d Zeilen. Sicherung: %s" % (geschrieben, os.path.basename(sicherung)))
    return 0


# ---------------------------------------------------------------------------
# HTML-Embed
# ---------------------------------------------------------------------------
START = "/*SEED_START*/"
ENDE = "/*SEED_END*/"


def embed(daten):
    if not os.path.exists(HTML):
        vorlage = os.path.join(HERE, "MiT_DC_Cockpit.template.html")
        if os.path.exists(vorlage):
            import shutil
            shutil.copyfile(vorlage, HTML)
            print("  Cockpit aus der Vorlage erzeugt.")
        else:
            print("HTML nicht gefunden: %s" % HTML, file=sys.stderr)
            return 1
    with open(HTML, encoding="utf-8") as f:
        html = f.read()
    a, b = html.find(START), html.find(ENDE)
    if a < 0 or b < 0:
        print("Marker SEED_START/SEED_END fehlen in der HTML-Datei.", file=sys.stderr)
        return 1
    neu = html[:a + len(START)] + "\nconst SEED = " + json.dumps(daten, ensure_ascii=False) + ";\n" + html[b:]
    with open(HTML, "w", encoding="utf-8") as f:
        f.write(neu)
    print("Datenstand in %s eingebacken (%.0f KB)." % (os.path.basename(HTML), len(neu) / 1024))
    return 0


# ---------------------------------------------------------------------------
# Lauf
# ---------------------------------------------------------------------------
def lauf(cfg, mit_geo=True, mit_embed=True):
    import dc_seed
    print("[%s] Aktualisierung laeuft" % dt.datetime.now().strftime("%d.%m.%Y %H:%M"))

    daten = dc_seed.baue(None)                 # Excel neu einlesen
    if os.path.exists(DATA):                   # News und Geo aus dem letzten Lauf uebernehmen
        try:
            with open(DATA, encoding="utf-8") as f:
                alt = json.load(f)
            daten["news"] = alt.get("news", [])
            daten["quellen_status"] = alt.get("quellen_status", [])
            daten["geo"] = alt.get("geo", {})
        except Exception:
            pass

    anzahl, ok = hole_news(cfg, daten)
    print("  News: %d relevant aus %d lebenden Feeds" % (anzahl, ok))

    if mit_geo:
        neu, offen = geocode(daten, cfg)
        print("  Geo: %d neu aufgeloest, %d ungeloest" % (neu, offen))

    daten["meta"]["aktualisiert"] = dt.datetime.now().isoformat(timespec="seconds")
    with open(DATA, "w", encoding="utf-8") as f:
        json.dump(daten, f, ensure_ascii=False, indent=1)
    print("  geschrieben: %s" % os.path.basename(DATA))

    speichere_config(cfg)                      # tote Feeds merken
    if mit_embed:
        embed(daten)
    return daten


def serve(port, cfg, intervall):
    import http.server
    import socketserver

    def hintergrund():
        while True:
            try:
                lauf(cfg, mit_geo=True, mit_embed=True)
            except Exception as e:
                print("  Fehler im Hintergrundlauf: %s" % e, file=sys.stderr)
            time.sleep(max(5, intervall) * 60)

    threading.Thread(target=hintergrund, daemon=True).start()

    class Handler(http.server.SimpleHTTPRequestHandler):
        def __init__(self, *a, **kw):
            super().__init__(*a, directory=HERE, **kw)

        def end_headers(self):
            self.send_header("Cache-Control", "no-store")
            super().end_headers()

        def log_message(self, *a):
            pass

    socketserver.TCPServer.allow_reuse_address = True
    with socketserver.TCPServer(("127.0.0.1", port), Handler) as srv:
        print("\nCockpit laeuft: http://127.0.0.1:%d/MiT_DC_Cockpit.html" % port)
        print("Aktualisierung alle %d Minuten. Beenden mit Strg+C.\n" % intervall)
        try:
            srv.serve_forever()
        except KeyboardInterrupt:
            print("\nbeendet")


def test_feeds(cfg):
    print("Feed-Test, es wird nichts geschrieben.\n")
    lebend = 0
    for feed in cfg["feeds"]:
        roh, fehler = hole(feed["url"], cfg)
        if roh is None:
            print("  TOT   %-28s %s  (%s)" % (feed["name"], feed["url"], fehler))
            continue
        eintraege = parse_feed(roh)
        if not eintraege:
            print("  LEER  %-28s %s  (antwortet, aber kein RSS/Atom)" % (feed["name"], feed["url"]))
            continue
        treffer = sum(1 for e in eintraege if relevant(e, cfg))
        lebend += 1
        print("  OK    %-28s %d Eintraege, %d relevant" % (feed["name"], len(eintraege), treffer))
    print("\n%d von %d Feeds liefern." % (lebend, len(cfg["feeds"])))
    if lebend == 0:
        print("Kein einziger Feed erreichbar. Pruefe Firmenproxy (config.json, Feld proxy) oder Firewall.")


def main():
    p = argparse.ArgumentParser(description="MiT Datacenter Cockpit, Updater")
    p.add_argument("--once", action="store_true", help="einmal aktualisieren")
    p.add_argument("--watch", type=int, metavar="MIN", help="alle MIN Minuten aktualisieren")
    p.add_argument("--serve", type=int, nargs="?", const=8777, metavar="PORT", help="lokalen Server starten")
    p.add_argument("--intervall", type=int, default=30, help="Minuten zwischen Laeufen bei --serve")
    p.add_argument("--test-feeds", action="store_true", help="nur Feeds pruefen")
    p.add_argument("--merge-log", action="store_true", help="anruf_log.csv in die Excel zurueckschreiben")
    p.add_argument("--embed", action="store_true", help="aktuellen Datenstand in die HTML backen")
    p.add_argument("--no-geo", action="store_true", help="Geocoding ueberspringen")
    a = p.parse_args()

    sys.path.insert(0, HERE)
    cfg = lade_config()

    if a.test_feeds:
        return test_feeds(cfg)
    if a.merge_log:
        return merge_log()
    if a.embed:
        if os.path.exists(DATA):
            with open(DATA, encoding="utf-8") as f:
                return embed(json.load(f))
        print("dc_data.json fehlt, erst --once laufen lassen.", file=sys.stderr)
        return 1
    if a.serve is not None:
        return serve(a.serve, cfg, a.intervall)
    if a.watch:
        while True:
            try:
                lauf(cfg, mit_geo=not a.no_geo)
            except Exception as e:
                print("Fehler: %s" % e, file=sys.stderr)
            time.sleep(a.watch * 60)
    lauf(cfg, mit_geo=not a.no_geo)


if __name__ == "__main__":
    sys.exit(main() or 0)
