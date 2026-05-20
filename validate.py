#!/usr/bin/env python3
"""validate.py — Pre-Import-Validator für MiT-CRM-Listen.

Nutzung:
    python validate.py kunden.xlsx
    python validate.py kunden.xlsx --report report.txt --json report.json
    python validate.py kunden.xlsx --fix kunden_fixed.xlsx

Prüft Header, Pflichtfelder, Dropdown-Werte, E-Mail-Format, Duplikate und
PLZ-Plausibilität. Exit-Code 0 = OK, 1 = Warnungen, 2 = Fehler.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path

import openpyxl

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PLZ_RE = re.compile(r"^\d{4}$")
PHONE_RE = re.compile(r"^[+0-9 ()/.\-]{6,}$")
URL_RE = re.compile(r"^https?://", re.IGNORECASE)

EXPECTED_HEADERS = [
    "Nr.", "Prio", "Segment", "Firmenname *", "Ort", "PLZ", "Kanton",
    "Ansprechpartner", "Funktion / Titel", "E-Mail", "Telefon", "Website",
    "Status", "Nächster Schritt", "Bedarf / kVA", "Hauptprodukt",
    "Follow-up Datum", "Notizen", "Internes",
]
# Optional analytics columns (introduced in Final template):
OPTIONAL_HEADERS = ["Wahrsch. %", "Wert CHF", "Letzte Aktivität"]

SEGMENTS = {
    "EVU/Netzbetreiber", "Spital/Klinik", "Pharma/Chemie", "Rechenzentrum/IT",
    "Industrie/Maschinenbau", "Bau/Infrastruktur", "Gemeinde/Öffentlich",
    "Bergbahn/Tourismus", "Events/Messen", "Elektroplaner/Engineering",
    "Forschung/Bildung", "Lebensmittel/Food",
}
PRIOS = {"A", "B", "C"}
STATUS = {"offen", "kontaktiert", "in Gespräch", "Angebot gesendet",
          "aktiv", "inaktiv"}
KANTONE = {"ZH", "BE", "LU", "AG", "SG", "GE", "BS", "BL", "SO", "TG", "VS",
           "VD", "FR", "GR", "AR", "AI", "SZ", "ZG", "TI", "NE", "UR", "OW",
           "NW", "GL", "SH", "JU"}
SEGMENT_NORMALIZE = {
    "Spital/Gesundheit": "Spital/Klinik",
    "RZ/Telekom": "Rechenzentrum/IT",
    "Elektroplaner": "Elektroplaner/Engineering",
    "Events / Messen": "Events/Messen",
}

SEVERITY_ORDER = {"error": 2, "warning": 1, "info": 0}


@dataclass
class Issue:
    severity: str  # error / warning / info
    row: int | None
    column: str | None
    code: str
    message: str

    def fmt(self) -> str:
        loc = f"row {self.row}" if self.row else "—"
        col = f"/{self.column}" if self.column else ""
        return f"[{self.severity.upper():7}] {loc}{col} {self.code}: {self.message}"


@dataclass
class Report:
    file: str
    sheet: str
    n_rows: int = 0
    issues: list[Issue] = field(default_factory=list)
    fixes: list[str] = field(default_factory=list)

    @property
    def n_errors(self) -> int:
        return sum(1 for i in self.issues if i.severity == "error")

    @property
    def n_warnings(self) -> int:
        return sum(1 for i in self.issues if i.severity == "warning")

    def add(self, *args, **kwargs) -> None:
        self.issues.append(Issue(*args, **kwargs))

    def summary(self) -> str:
        return (f"{self.file} :: {self.sheet} — "
                f"{self.n_rows} Zeilen, {self.n_errors} Fehler, "
                f"{self.n_warnings} Warnungen")


def find_sheet(wb) -> str:
    candidates = ["📋 Kunden-Datenbank", "Kunden-Datenbank", "Kunden",
                  "Customers", "CRM"]
    for c in candidates:
        if c in wb.sheetnames:
            return c
    return wb.sheetnames[0]


def detect_header_row(ws) -> int:
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, 5)]
        if "Firmenname *" in vals or "Firmenname" in vals or vals[0] == "Nr.":
            return r
    return 3


def validate(path: str, fix_path: str | None = None) -> Report:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = find_sheet(wb)
    ws = wb[sheet]
    rep = Report(file=path, sheet=sheet)

    header_row = detect_header_row(ws)
    headers = [ws.cell(row=header_row, column=c).value
               for c in range(1, ws.max_column + 1)]
    headers_clean = [h for h in headers if h]

    # Header check
    missing = [h for h in EXPECTED_HEADERS if h not in headers_clean]
    if missing:
        rep.add("error", header_row, None, "HDR001",
                f"Fehlende Pflicht-Spalten: {missing}")
    extras = [h for h in headers_clean
              if h not in EXPECTED_HEADERS and h not in OPTIONAL_HEADERS]
    if extras:
        rep.add("warning", header_row, None, "HDR002",
                f"Unbekannte Spalten: {extras}")

    col_idx = {h: i + 1 for i, h in enumerate(headers) if h}

    def col(name: str) -> int | None:
        return col_idx.get(name)

    keys_seen: dict[tuple, int] = {}
    n_data = 0
    fix_changes = []

    for r in range(header_row + 1, ws.max_row + 1):
        firma_col = col("Firmenname *")
        if firma_col is None:
            break
        firma = ws.cell(row=r, column=firma_col).value
        if not firma:
            # Skip empty trailing rows; only flag if other cells are filled
            has_other = any(ws.cell(row=r, column=c).value is not None
                            for c in range(1, ws.max_column + 1))
            if has_other:
                rep.add("error", r, "D", "REQ001",
                        "Firmenname leer aber andere Felder gesetzt")
            continue
        n_data += 1

        # Prio
        c_prio = col("Prio")
        if c_prio:
            v = ws.cell(row=r, column=c_prio).value
            if v not in PRIOS:
                rep.add("error", r, "B", "ENM001",
                        f"Prio '{v}' ungültig (erlaubt: A/B/C)")

        # Segment
        c_seg = col("Segment")
        if c_seg:
            v = ws.cell(row=r, column=c_seg).value
            if v in SEGMENT_NORMALIZE:
                new = SEGMENT_NORMALIZE[v]
                rep.add("warning", r, "C", "ENM002",
                        f"Segment '{v}' → kanonisch '{new}'")
                if fix_path:
                    ws.cell(row=r, column=c_seg).value = new
                    fix_changes.append(f"row {r}: Segment {v} → {new}")
            elif v not in SEGMENTS:
                rep.add("error", r, "C", "ENM002",
                        f"Segment '{v}' nicht in Taxonomie")

        # Kanton
        c_kt = col("Kanton")
        if c_kt:
            v = ws.cell(row=r, column=c_kt).value
            if v and v not in KANTONE:
                rep.add("error", r, "G", "ENM003",
                        f"Kanton '{v}' kein gültiges CH-Kürzel")

        # Status
        c_st = col("Status")
        if c_st:
            v = ws.cell(row=r, column=c_st).value
            if v and v not in STATUS:
                rep.add("error", r, "M", "ENM004", f"Status '{v}' ungültig")
            elif not v and fix_path:
                ws.cell(row=r, column=c_st).value = "offen"
                fix_changes.append(f"row {r}: Status → 'offen'")

        # E-Mail
        c_em = col("E-Mail")
        if c_em:
            v = ws.cell(row=r, column=c_em).value
            if v:
                s = str(v).strip()
                if not EMAIL_RE.match(s):
                    rep.add("error", r, "J", "FMT001",
                            f"E-Mail '{v}' nicht im Format user@domain.tld")
                    if fix_path:
                        ws.cell(row=r, column=c_em).value = None
                        fix_changes.append(f"row {r}: E-Mail '{v}' geleert")

        # PLZ
        c_plz = col("PLZ")
        if c_plz:
            v = ws.cell(row=r, column=c_plz).value
            if v is not None:
                s = str(v).strip()
                if not PLZ_RE.match(s):
                    rep.add("warning", r, "F", "FMT002",
                            f"PLZ '{v}' nicht 4-stellig")

        # Telefon
        c_tel = col("Telefon")
        if c_tel:
            v = ws.cell(row=r, column=c_tel).value
            if v and not PHONE_RE.match(str(v).strip()):
                rep.add("warning", r, "K", "FMT003",
                        f"Telefon '{v}' enthält unerwartete Zeichen")

        # Website
        c_url = col("Website")
        if c_url:
            v = ws.cell(row=r, column=c_url).value
            if v and not URL_RE.match(str(v).strip()):
                rep.add("info", r, "L", "FMT004",
                        f"Website '{v}' ohne http(s)://")

        # Duplikat
        ort = ws.cell(row=r, column=col("Ort") or 1).value if col("Ort") else None
        plz = ws.cell(row=r, column=col("PLZ") or 1).value if col("PLZ") else None
        key = (str(firma).strip().lower(),
               str(ort or "").strip().lower(),
               str(plz or "").strip())
        if key in keys_seen:
            rep.add("error", r, "D", "DUP001",
                    f"Duplikat von Zeile {keys_seen[key]} (Firma+Ort+PLZ)")
        else:
            keys_seen[key] = r

    rep.n_rows = n_data

    if fix_path and fix_changes:
        wb.save(fix_path)
        rep.fixes = fix_changes
    return rep


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("file", help="Pfad zur .xlsx-Datei")
    ap.add_argument("--report", help="Textreport in Datei schreiben")
    ap.add_argument("--json", dest="json_out", help="JSON-Report schreiben")
    ap.add_argument("--fix", help="Auto-Korrekturen in neue Datei speichern")
    ap.add_argument("--quiet", action="store_true",
                    help="Nur Summary ausgeben")
    args = ap.parse_args()

    if not Path(args.file).exists():
        print(f"Datei nicht gefunden: {args.file}", file=sys.stderr)
        return 2

    rep = validate(args.file, fix_path=args.fix)

    lines = [rep.summary(), ""]
    if not args.quiet:
        # Group by severity
        for sev in ("error", "warning", "info"):
            grp = [i for i in rep.issues if i.severity == sev]
            if grp:
                lines.append(f"== {sev.upper()} ({len(grp)}) ==")
                for i in grp[:50]:
                    lines.append("  " + i.fmt())
                if len(grp) > 50:
                    lines.append(f"  … und {len(grp)-50} weitere")
                lines.append("")
        if rep.fixes:
            lines.append(f"== AUTO-FIXES ({len(rep.fixes)}) ==")
            for f in rep.fixes[:20]:
                lines.append("  " + f)
            if len(rep.fixes) > 20:
                lines.append(f"  … und {len(rep.fixes)-20} weitere")

    text = "\n".join(lines)
    print(text)

    if args.report:
        Path(args.report).write_text(text, encoding="utf-8")
    if args.json_out:
        data = {
            "file": rep.file, "sheet": rep.sheet, "n_rows": rep.n_rows,
            "n_errors": rep.n_errors, "n_warnings": rep.n_warnings,
            "issues": [asdict(i) for i in rep.issues],
            "fixes": rep.fixes,
        }
        Path(args.json_out).write_text(json.dumps(data, ensure_ascii=False,
                                                  indent=2), encoding="utf-8")

    if rep.n_errors:
        return 2
    if rep.n_warnings:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
