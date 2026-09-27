"""Erzeugt Kundenliste_MiT_Schweiz.xlsx aus daten.py."""
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from daten import D, R, INST, FM, EVU, PLAN, GU, RZ, IND, SPI, INF, EVT

OUT = "Kundenliste_MiT_Schweiz.xlsx"
STAND = "27.09.2026"

NAVY, TEAL, LIGHT, GREY = "1B2430", "0E7C86", "EAF2F4", "F3F4F6"
FONT = "Arial"
thin = Side(style="thin", color="D0D5DB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

SEG = {
    INST: "Installateure (60 %)",
    FM: "FM / EVU / Planer (20 %)", EVU: "FM / EVU / Planer (20 %)", PLAN: "FM / EVU / Planer (20 %)",
    GU: "GU/TU + Industrie (20 %)", RZ: "GU/TU + Industrie (20 %)", IND: "GU/TU + Industrie (20 %)",
    SPI: "GU/TU + Industrie (20 %)", INF: "GU/TU + Industrie (20 %)", EVT: "GU/TU + Industrie (20 %)",
}
SEGMENTE = ["Installateure (60 %)", "FM / EVU / Planer (20 %)", "GU/TU + Industrie (20 %)"]
SOLL = [0.6, 0.2, 0.2]
KATEGORIEN = [INST, FM, EVU, PLAN, GU, RZ, IND, SPI, INF, EVT]
REGIONEN = [R[k] for k in "WMNZOCT"]
GEBIET = {"W": "West", "T": "Tessin"}
STATUS = ["Offen: CRM prüfen", "A-Kunde", "Bestandskunde mit Potenzial", "Verlorener Kunde", "Potenzial / Aufbau"]
TEAM = ["Burak", "Mauro", "Jörg", "Roberto"]

ROLLE = {
    INST: ("Leiter Grossprojekte / Bereichsleiter Baustrom / Einkauf", "Projektleiter"),
    FM: ("Leiter Technisches FM / Key Account Industrie & Pharma", "Technisches Facility Management"),
    EVU: ("Leiter Netzbau / Leiter Unterwerke / Asset Management", "Netzbau"),
    PLAN: ("Fachbereichsleiter Elektro / Energie", "Elektroplanung"),
    GU: ("Projektleiter Baustelleneinrichtung / Einkauf", "Baustelleneinrichtung"),
    RZ: ("Construction / Commissioning Manager, Critical Facility Ops", "Commissioning"),
    IND: ("Leiter Energie & Utilities / Technischer Einkauf", "Utilities"),
    SPI: ("Leiter Technik & Betrieb", "Leiter Technik"),
    INF: ("Leiter Energieversorgung / Gesamtprojektleiter", "Energieversorgung"),
    EVT: ("Produktionsleitung / Technische Leitung", "Technische Leitung"),
}

COLS = [  # (Header, Breite)
    ("ID", 6), ("Kernliste Top 70", 10), ("Segment", 20), ("Kategorie", 15), ("Firma", 30),
    ("Gruppe / Mutter", 22), ("Ort", 18), ("Kt.", 5), ("Region", 17), ("Gebiet", 13), ("Sprache", 9),
    ("Website", 20), ("Grösse", 24), ("Anlass / Projekt 2025-2027", 38), ("MiT-Leistungen", 26),
    ("Warum (Begründung)", 48), ("Fit", 6), ("Volumen", 8), ("Timing", 7), ("Zugang", 7),
    ("Score", 7), ("Prio Vorschlag", 9), ("Prio final", 8), ("Prio", 6), ("Kundenstatus MiT", 22),
    ("Verloren seit", 11), ("Zielrolle Ansprechperson", 34), ("Ansprechperson (Name)", 22),
    ("GL öffentlich (Quelle)", 26), ("LinkedIn Firma", 13), ("LinkedIn Personen", 14), ("Quelle", 11),
    ("Validierung / Hinweis", 36), ("Verantwortlich", 13), ("Nächster Schritt", 26),
    ("Datum nächster Schritt", 12), ("Letzter Kontakt", 11),
]
H = {name: get_column_letter(i + 1) for i, (name, _) in enumerate(COLS)}
HDR_ROW, FIRST = 4, 5


def score(r):
    return r[12] * 0.3 + r[13] * 0.3 + r[14] * 0.25 + r[15] * 0.15


seg_order = {s: i for i, s in enumerate(SEGMENTE)}
rows = sorted(D, key=lambda r: (seg_order[SEG[r[0]]], KATEGORIEN.index(r[0]), -score(r), r[1]))
LAST = FIRST + len(rows) - 1

wb = Workbook()

# ------------------------------------------------------------------ Methodik
me = wb.active
me.title = "Methodik"
kl = wb.create_sheet("Kundenliste", 0)
db = wb.create_sheet("Übersicht", 0)
li = wb.create_sheet("Listen")


def style_title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="5B6470")


style_title(me, "Methodik, Legende und Grenzen der Recherche", f"Stand {STAND}")
me.column_dimensions["A"].width = 30
me.column_dimensions["B"].width = 12
me.column_dimensions["C"].width = 95

me["A4"], me["B4"], me["C4"] = "Gewichtung Score", "Wert", "Bedeutung (Skala je Kriterium 1 = schwach bis 5 = stark)"
crit = [
    ("Fit", 0.30, "Passt der Bedarf zum MiT-Portfolio (Generator, BESS, USV, Trafo, NEA-Test, PQ, Baustrom, Event)?"),
    ("Volumen", 0.30, "Erwartbares Mietvolumen pro Jahr bzw. Hebel über Töchter/Rahmenvertrag."),
    ("Timing", 0.25, "Konkreter Anlass 2025-2027 (Bau, UW-Umbau, Commissioning, Ramp-up, Event-Saison)."),
    ("Zugang", 0.15, "Erreichbarkeit: regionale Nähe, Entscheid lokal vs. Konzern/EMEA, öffentliche Beschaffung."),
]
for i, (k, w, t) in enumerate(crit):
    r = 5 + i
    me.cell(r, 1, k); me.cell(r, 2, w).number_format = "0%"; me.cell(r, 3, t)
    me.cell(r, 2).font = Font(name=FONT, color="0000FF")
    me.cell(r, 2).fill = PatternFill("solid", fgColor="FFFF00")
me["A9"], me["B9"] = "Summe Gewichte", "=SUM(B5:B8)"
me["B9"].number_format = "0%"
me["A11"], me["B11"], me["C11"] = "Schwelle Prio A", 4.0, "Score >= Schwelle A: jetzt aktiv angehen (Termin in den nächsten 6 Wochen)."
me["A12"], me["B12"], me["C12"] = "Schwelle Prio B", 3.0, "Score >= Schwelle B: im Quartal angehen, Anlass beobachten."
me["A13"], me["C13"] = "Prio C", "Darunter: Beobachten, opportunistisch bedienen, über Partner/Töchter abdecken."
for c in ("B11", "B12"):
    me[c].font = Font(name=FONT, color="0000FF")
    me[c].fill = PatternFill("solid", fgColor="FFFF00")
    me[c].number_format = "0.0"

notes = [
    ("Bedienung", ""),
    ("Gelbe Zellen", "Eingaben. Gewichte und Schwellen hier ändern, Score und Prio in der Kundenliste rechnen neu."),
    ("Prio final", "Leer lassen = Vorschlag gilt. A/B/C eintragen = übersteuert den Vorschlag. Spalte «Prio» ist die gültige."),
    ("Kundenstatus MiT", "Alle Zeilen stehen auf «Offen: CRM prüfen». Bestands-, A- und verlorene Kunden sind öffentlich nicht"),
    ("", "recherchierbar und müssen aus dem MiT-CRM bzw. von Mauro, Jörg und Roberto eingetragen werden (Dropdown)."),
    ("Verloren seit", "Nur bei Status «Verlorener Kunde» ausfüllen (Datum). Grund in «Validierung / Hinweis»."),
    ("Kernliste Top 70", "70 Firmen im Verhältnis 60/20/20 (42/14/14) nach Vorgabe. Alle übrigen Zeilen sind Reserve bzw. Ergänzung."),
    ("", ""),
    ("Grenzen der Recherche", ""),
    ("Quellen", "Websuche vom 27.09.2026: Firmenwebsites, Handelsregister-Aggregatoren, Medienmitteilungen, Fachpresse."),
    ("", "Einige Seiten (u.a. NZZ, SRF, einzelne .ch-Domains) waren nicht direkt abrufbar: Werte teils nur aus Suchergebnissen."),
    ("Mitarbeiterzahlen", "Teils aus Vorjahren oder Sekundärquellen. «(prüfen)» bzw. «uneinheitlich» = vor externer Verwendung verifizieren."),
    ("Ansprechpersonen", "Keine Namen erfunden. Nur öffentlich genannte GL-Mitglieder mit Quelle. Die LinkedIn-Links öffnen eine"),
    ("", "Suche (Firma bzw. Firma + Zielfunktion), keine verifizierten Profile. Namen nach Prüfung in «Ansprechperson» eintragen."),
    ("Relevanz / Warum", "Einschätzung aus Sicht mobiler Energie. Belegt sind nur die Fakten in «Anlass / Projekt» und «Grösse»."),
    ("Nicht aufgenommen", "Steiner AG (Nachlassstundung), Deltalis (RZ-Betrieb eingestellt). Ohne belastbare Quelle weggelassen:"),
    ("", "Dussmann, Kaufmann, Perenzia, Nine, Swiss Fort Knox, atNorth, Colt. BKW (Konzern), Axpo Grid, Swissgrid, AEW, EKT, EWL,"),
    ("", "St.Galler Stadtwerke sind reale Zielkunden, aber noch ohne geprüfte Kennzahlen: als nächste Ergänzung vorgesehen."),
    ("Wettbewerb/Partner", "Kummler+Matter EVT vermietet eigene mobile Trafostationen: Partner und Wettbewerber zugleich."),
    ("Vertraulich", "Interne Arbeitsliste. Nicht an Kunden weitergeben."),
]
r = 15
for k, t in notes:
    me.cell(r, 1, k); me.cell(r, 3, t)
    if t == "" and k:
        me.cell(r, 1).font = Font(name=FONT, bold=True, size=11, color=TEAL)
    r += 1

for row in me.iter_rows(min_row=4, max_row=r):
    for c in row:
        if c.font.name != FONT or (c.font.color is None):
            c.font = Font(name=FONT, size=10, bold=c.font.bold, color=c.font.color)
for c in ("A4", "B4", "C4"):
    me[c].font = Font(name=FONT, bold=True, color="FFFFFF")
    me[c].fill = PatternFill("solid", fgColor=NAVY)

# ------------------------------------------------------------------ Listen
lists = {"A": ("Segment", SEGMENTE), "B": ("Kategorie", KATEGORIEN), "C": ("Region", REGIONEN),
         "D": ("Prio", ["A", "B", "C"]), "E": ("Kundenstatus", STATUS), "F": ("Team", TEAM),
         "G": ("Kernliste", ["Ja", "Nein"]), "H": ("Skala", [1, 2, 3, 4, 5]),
         "I": ("Gebiet", ["West", "Deutschschweiz", "Tessin"])}
for col, (name, vals) in lists.items():
    li[f"{col}1"] = name
    li[f"{col}1"].font = Font(name=FONT, bold=True)
    for i, v in enumerate(vals):
        li[f"{col}{i + 2}"] = v
    li.column_dimensions[col].width = 26
li.sheet_state = "hidden"

# ------------------------------------------------------------------ Kundenliste
style_title(kl, "Zielkundenliste Schweiz: Mobile Energie (Generator, BESS, USV, Trafo, NEA, PQ, Baustrom, Event)",
            f"Stand {STAND}  |  {len(rows)} Firmen  |  Kernliste 70 (60/20/20)  |  Score und Prio rechnen aus Blatt «Methodik»  |  VERTRAULICH")
kl["A3"] = "Blaue Spalten Q-T (Fit/Volumen/Timing/Zugang) sowie Prio final, Kundenstatus, Ansprechperson, Verantwortlich, nächster Schritt und Datum sind Eingabefelder."
kl["A3"].font = Font(name=FONT, size=9, color="0000FF")

for i, (name, w) in enumerate(COLS):
    c = kl.cell(HDR_ROW, i + 1, name)
    c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    kl.column_dimensions[get_column_letter(i + 1)].width = w
kl.row_dimensions[HDR_ROW].height = 32

WRAP = {"Firma", "Gruppe / Mutter", "Grösse", "Anlass / Projekt 2025-2027", "MiT-Leistungen", "Warum (Begründung)",
        "Zielrolle Ansprechperson", "GL öffentlich (Quelle)", "Validierung / Hinweis", "Ort", "Nächster Schritt", "Segment"}
CENTER = {"ID", "Kernliste Top 70", "Kt.", "Sprache", "Fit", "Volumen", "Timing", "Zugang", "Score",
          "Prio Vorschlag", "Prio final", "Prio", "Verloren seit", "Datum nächster Schritt", "Letzter Kontakt"}
INPUT = {"Fit", "Volumen", "Timing", "Zugang", "Prio final", "Kundenstatus MiT", "Verloren seit",
         "Ansprechperson (Name)", "Verantwortlich", "Nächster Schritt", "Datum nächster Schritt", "Letzter Kontakt"}

for n, rec in enumerate(rows):
    (kat, firma, mutter, ort, kt, reg, spr, web, gr, anl, prod, warum, f, v, t, z, kern, gl, quelle, val) = rec
    r = FIRST + n
    rolle, kw = ROLLE[kat]
    vals = {
        "ID": f"CH-{n + 1:03d}", "Kernliste Top 70": kern, "Segment": SEG[kat], "Kategorie": kat, "Firma": firma,
        "Gruppe / Mutter": mutter, "Ort": ort, "Kt.": kt, "Region": R[reg], "Gebiet": GEBIET.get(reg, "Deutschschweiz"),
        "Sprache": spr, "Website": web or "prüfen", "Grösse": gr, "Anlass / Projekt 2025-2027": anl or "Kein spezifischer Anlass recherchiert",
        "MiT-Leistungen": prod, "Warum (Begründung)": warum, "Fit": f, "Volumen": v, "Timing": t, "Zugang": z,
        "Score": (f"=ROUND({H['Fit']}{r}*Methodik!$B$5+{H['Volumen']}{r}*Methodik!$B$6"
                  f"+{H['Timing']}{r}*Methodik!$B$7+{H['Zugang']}{r}*Methodik!$B$8,2)"),
        "Prio Vorschlag": f'=IF({H["Score"]}{r}>=Methodik!$B$11,"A",IF({H["Score"]}{r}>=Methodik!$B$12,"B","C"))',
        "Prio final": None,
        "Prio": f'=IF({H["Prio final"]}{r}="",{H["Prio Vorschlag"]}{r},{H["Prio final"]}{r})',
        "Kundenstatus MiT": STATUS[0], "Verloren seit": None, "Zielrolle Ansprechperson": rolle,
        "Ansprechperson (Name)": None, "GL öffentlich (Quelle)": gl or None,
        "LinkedIn Firma": "Suche öffnen", "LinkedIn Personen": "Suche öffnen", "Quelle": "Link",
        "Validierung / Hinweis": val or None, "Verantwortlich": None, "Nächster Schritt": None,
        "Datum nächster Schritt": None, "Letzter Kontakt": None,
    }
    for i, (name, _) in enumerate(COLS):
        c = kl.cell(r, i + 1, vals[name])
        c.font = Font(name=FONT, size=9, color="0000FF" if name in INPUT else "000000")
        c.alignment = Alignment(wrap_text=name in WRAP, vertical="top",
                                horizontal="center" if name in CENTER else "left")
        c.border = BORDER
    kl[f"{H['Firma']}{r}"].font = Font(name=FONT, size=9, bold=True)
    kl[f"{H['Score']}{r}"].number_format = "0.00"
    for dc in ("Verloren seit", "Datum nächster Schritt", "Letzter Kontakt"):
        kl[f"{H[dc]}{r}"].number_format = "DD.MM.YYYY"
    link_font = Font(name=FONT, size=9, color="0563C1", underline="single")
    if web:
        kl[f"{H['Website']}{r}"].hyperlink = "https://" + web.split(" ")[0]
        kl[f"{H['Website']}{r}"].font = link_font
    kl[f"{H['LinkedIn Firma']}{r}"].hyperlink = (
        "https://www.linkedin.com/search/results/companies/?keywords=" + quote_plus(firma.split(" (")[0]))
    kl[f"{H['LinkedIn Personen']}{r}"].hyperlink = (
        "https://www.linkedin.com/search/results/people/?keywords=" + quote_plus(f'{firma.split(" (")[0]} {kw}'))
    kl[f"{H['Quelle']}{r}"].hyperlink = quelle
    kl[f"{H['Quelle']}{r}"].comment = Comment(quelle, "Recherche")
    for lc in ("LinkedIn Firma", "LinkedIn Personen", "Quelle"):
        kl[f"{H[lc]}{r}"].font = link_font
    kl.row_dimensions[r].height = 48

ref = f"A{HDR_ROW}:{get_column_letter(len(COLS))}{LAST}"
tab = Table(displayName="Kunden", ref=ref)
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
kl.add_table(tab)
kl.freeze_panes = f"{H['Gruppe / Mutter']}{FIRST}"


def dv(formula, col, prompt=None, allow_blank=True):
    d = DataValidation(type="list", formula1=formula, allow_blank=allow_blank, showErrorMessage=True)
    d.error = "Bitte Wert aus der Liste wählen."
    if prompt:
        d.prompt, d.showInputMessage = prompt, True
    kl.add_data_validation(d)
    d.add(f"{H[col]}{FIRST}:{H[col]}{LAST + 200}")


dv("=Listen!$A$2:$A$4", "Segment")
dv("=Listen!$B$2:$B$11", "Kategorie")
dv("=Listen!$C$2:$C$8", "Region")
dv("=Listen!$I$2:$I$4", "Gebiet")
dv("=Listen!$G$2:$G$3", "Kernliste Top 70")
dv("=Listen!$D$2:$D$4", "Prio final", "Leer = Vorschlag gilt. A/B/C übersteuert.")
dv("=Listen!$E$2:$E$6", "Kundenstatus MiT", "Aus CRM übernehmen.")
dv("=Listen!$F$2:$F$5", "Verantwortlich")
for col in ("Fit", "Volumen", "Timing", "Zugang"):
    dv("=Listen!$H$2:$H$6", col, "Skala 1 bis 5, siehe Blatt Methodik.")
dd = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
dd.error = "Bitte Datum TT.MM.JJJJ eingeben."
kl.add_data_validation(dd)
for col in ("Verloren seit", "Datum nächster Schritt", "Letzter Kontakt"):
    dd.add(f"{H[col]}{FIRST}:{H[col]}{LAST + 200}")

PRIO_FILL = {"A": ("C6EFCE", "006100"), "B": ("FFEB9C", "9C5700"), "C": ("E7E6E6", "595959")}
for col in ("Prio Vorschlag", "Prio final", "Prio"):
    rng = f"{H[col]}{FIRST}:{H[col]}{LAST}"
    for p, (bg, fg) in PRIO_FILL.items():
        kl.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{p}"'],
                                                      fill=PatternFill("solid", fgColor=bg),
                                                      font=Font(name=FONT, bold=True, color=fg)))
kl.conditional_formatting.add(
    f"{H['Kernliste Top 70']}{FIRST}:{H['Kernliste Top 70']}{LAST}",
    CellIsRule(operator="equal", formula=['"Ja"'], fill=PatternFill("solid", fgColor=LIGHT), font=Font(name=FONT, bold=True, color=TEAL)))
kl.conditional_formatting.add(
    f"{H['Kundenstatus MiT']}{FIRST}:{H['Kundenstatus MiT']}{LAST}",
    CellIsRule(operator="equal", formula=['"Verlorener Kunde"'], fill=PatternFill("solid", fgColor="FFC7CE")))
kl.conditional_formatting.add(
    f"{H['Kundenstatus MiT']}{FIRST}:{H['Kundenstatus MiT']}{LAST}",
    CellIsRule(operator="equal", formula=['"A-Kunde"'], fill=PatternFill("solid", fgColor="C6EFCE")))
# überfälliger nächster Schritt
dcol = H["Datum nächster Schritt"]
kl.conditional_formatting.add(
    f"{dcol}{FIRST}:{dcol}{LAST}",
    FormulaRule(formula=[f'AND({dcol}{FIRST}<>"",{dcol}{FIRST}<TODAY())'], fill=PatternFill("solid", fgColor="FFC7CE")))
kl.sheet_view.zoomScale = 90

# ------------------------------------------------------------------ Übersicht
style_title(db, "Übersicht Zielkunden Schweiz", f"Stand {STAND}  |  alle Zahlen per Formel aus «Kundenliste»  |  VERTRAULICH")
rng = lambda col: f"Kundenliste!${H[col]}${FIRST}:${H[col]}${LAST}"
SEGR, KATR, REGR, PRR, KERNR, STR, GEBR = (rng("Segment"), rng("Kategorie"), rng("Region"), rng("Prio"),
                                          rng("Kernliste Top 70"), rng("Kundenstatus MiT"), rng("Gebiet"))


def header(ws, row, labels, col0=1):
    for i, l in enumerate(labels):
        c = ws.cell(row, col0 + i, l)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 30


def section(ws, row, text):
    ws.cell(row, 1, text).font = Font(name=FONT, bold=True, size=12, color=TEAL)


def body(ws, r1, r2, c1, c2):
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
        for c in row:
            c.font = Font(name=FONT, size=10, bold=c.font.bold)
            c.border = BORDER
            if c.column > c1:
                c.alignment = Alignment(horizontal="center")


# KPI-Kacheln
kpis = [("Firmen gesamt", f"=COUNTA({rng('Firma')})"),
        ("Kernliste", f'=COUNTIF({KERNR},"Ja")'),
        ("Prio A", f'=COUNTIF({PRR},"A")'),
        ("Prio B", f'=COUNTIF({PRR},"B")'),
        ("Prio C", f'=COUNTIF({PRR},"C")'),
        ("Status offen (CRM)", f'=COUNTIF({STR},"Offen: CRM prüfen")')]
for i, (lab, f) in enumerate(kpis):
    col = 1 + i
    lc, vc = db.cell(4, col, lab), db.cell(5, col, f)
    lc.font = Font(name=FONT, size=9, color="5B6470")
    vc.font = Font(name=FONT, size=20, bold=True, color=NAVY)
    for c in (lc, vc):
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.alignment = Alignment(horizontal="center", vertical="center")
db.row_dimensions[5].height = 34

# Segmentmix
section(db, 7, "Segmentmix gemäss Vorgabe 60/20/20")
header(db, 8, ["Segment", "Firmen gesamt", "Kernliste", "Soll-Anteil", "Ist-Anteil Kernliste", "Abweichung", "Prio A", "Prio B", "Prio C"])
for i, s in enumerate(SEGMENTE):
    r = 9 + i
    db.cell(r, 1, s)
    db.cell(r, 2, f'=COUNTIF({SEGR},$A{r})')
    db.cell(r, 3, f'=COUNTIFS({SEGR},$A{r},{KERNR},"Ja")')
    db.cell(r, 4, SOLL[i]).number_format = "0%"
    db.cell(r, 5, f"=IF($C$12=0,0,C{r}/$C$12)").number_format = "0%"
    db.cell(r, 6, f"=E{r}-D{r}").number_format = "+0%;-0%;0%"
    for j, p in enumerate("ABC"):
        db.cell(r, 7 + j, f'=COUNTIFS({SEGR},$A{r},{PRR},"{p}")')
db.cell(12, 1, "Total").font = Font(name=FONT, bold=True)
for c in range(2, 10):
    L = get_column_letter(c)
    if c in (4, 5):
        db.cell(12, c, f"=SUM({L}9:{L}11)").number_format = "0%"
    elif c == 6:
        continue
    else:
        db.cell(12, c, f"=SUM({L}9:{L}11)")
body(db, 9, 12, 1, 9)

# Region x Prio
section(db, 14, "Regionen x Prio")
header(db, 15, ["Region", "Prio A", "Prio B", "Prio C", "Total", "davon Kernliste", "Installateure", "FM/EVU/Planer", "GU/TU+Industrie"])
for i, reg in enumerate(REGIONEN):
    r = 16 + i
    db.cell(r, 1, reg)
    for j, p in enumerate("ABC"):
        db.cell(r, 2 + j, f'=COUNTIFS({REGR},$A{r},{PRR},"{p}")')
    db.cell(r, 5, f"=SUM(B{r}:D{r})")
    db.cell(r, 6, f'=COUNTIFS({REGR},$A{r},{KERNR},"Ja")')
    for j, s in enumerate(SEGMENTE):
        db.cell(r, 7 + j, f'=COUNTIFS({REGR},$A{r},{SEGR},"{s}")')
rt = 16 + len(REGIONEN)
db.cell(rt, 1, "Total").font = Font(name=FONT, bold=True)
for c in range(2, 10):
    L = get_column_letter(c)
    db.cell(rt, c, f"=SUM({L}16:{L}{rt - 1})")
body(db, 16, rt, 1, 9)

# Gebiet
g0 = rt + 2
section(db, g0, "Verkaufsgebiet")
header(db, g0 + 1, ["Gebiet", "Prio A", "Prio B", "Prio C", "Total"])
for i, g in enumerate(["West", "Deutschschweiz", "Tessin"]):
    r = g0 + 2 + i
    db.cell(r, 1, g)
    for j, p in enumerate("ABC"):
        db.cell(r, 2 + j, f'=COUNTIFS({GEBR},$A{r},{PRR},"{p}")')
    db.cell(r, 5, f"=SUM(B{r}:D{r})")
body(db, g0 + 2, g0 + 4, 1, 5)

# Kategorie
k0 = g0 + 6
section(db, k0, "Kategorien")
header(db, k0 + 1, ["Kategorie", "Prio A", "Prio B", "Prio C", "Total", "Ø Score"])
for i, k in enumerate(KATEGORIEN):
    r = k0 + 2 + i
    db.cell(r, 1, k)
    for j, p in enumerate("ABC"):
        db.cell(r, 2 + j, f'=COUNTIFS({KATR},$A{r},{PRR},"{p}")')
    db.cell(r, 5, f"=SUM(B{r}:D{r})")
    db.cell(r, 6, f'=IFERROR(AVERAGEIF({KATR},$A{r},{rng("Score")}),0)').number_format = "0.00"
body(db, k0 + 2, k0 + 1 + len(KATEGORIEN), 1, 6)

# Kundenstatus
s0 = k0 + 3 + len(KATEGORIEN)
section(db, s0, "Kundenstatus MiT (aus CRM zu befüllen)")
header(db, s0 + 1, ["Status", "Anzahl", "davon Prio A"])
for i, s in enumerate(STATUS):
    r = s0 + 2 + i
    db.cell(r, 1, s)
    db.cell(r, 2, f'=COUNTIF({STR},$A{r})')
    db.cell(r, 3, f'=COUNTIFS({STR},$A{r},{PRR},"A")')
body(db, s0 + 2, s0 + 1 + len(STATUS), 1, 3)

db.column_dimensions["A"].width = 30
for c in "BCDEFGHI":
    db.column_dimensions[c].width = 15

ch = BarChart()
ch.type, ch.grouping, ch.overlap = "bar", "stacked", 100
ch.title = "Firmen pro Region nach Prio"
ch.add_data(Reference(db, min_col=2, max_col=4, min_row=15, max_row=15 + len(REGIONEN)), titles_from_data=True)
ch.set_categories(Reference(db, min_col=1, min_row=16, max_row=15 + len(REGIONEN)))
for s, col in zip(ch.series, ("2E8B57", "E0A526", "A6A6A6")):
    s.graphicalProperties.solidFill = col
    s.graphicalProperties.line.solidFill = col
ch.height, ch.width = 8, 16
ch.y_axis.majorGridlines = None
db.add_chart(ch, "K7")

ch2 = BarChart()
ch2.type = "col"
ch2.title = "Kernliste: Soll vs. Ist"
ch2.add_data(Reference(db, min_col=4, max_col=5, min_row=8, max_row=11), titles_from_data=True)
ch2.set_categories(Reference(db, min_col=1, min_row=9, max_row=11))
for s, col in zip(ch2.series, ("A6A6A6", TEAL)):
    s.graphicalProperties.solidFill = col
ch2.y_axis.number_format = "0%"
ch2.y_axis.majorGridlines = None
ch2.height, ch2.width = 8, 16
db.add_chart(ch2, "K25")
db.sheet_view.showGridLines = False

for ws in (kl, me):
    ws.sheet_view.showGridLines = False if ws is me else True
kl.sheet_properties.tabColor = TEAL
db.sheet_properties.tabColor = NAVY
wb.active = 0
wb.save(OUT)
print(OUT, len(rows), "Zeilen")
