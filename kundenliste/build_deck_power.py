# -*- coding: utf-8 -*-
"""Top-Player Power Schweiz 2027 im MiT-Master (Skill mit-praesentation).
Zahlen live aus daten_power.py, identisch zur Excel-Logik (Rundung auf 100 CHF)."""
import glob
import math
import os
import sys

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ASSETS = os.environ.get("MIT_ASSETS") or (glob.glob(
    "/root/.claude/skills/synced/*/mit-praesentation/assets") or [""])[0]
sys.path.insert(0, ASSETS)
from mit_deck import Deck, RED, DARK, DARKER, LIGHT, WHITE, BLACK, SOFT, F_BODY, F_TITLE, L, R, cols  # noqa: E402

from daten_power import P, COMP, RATES, rate_gen, rate_lb, rate_bess, RZ, INF, NETZ, SPI, IND, EVT, KAN  # noqa: E402

MID = RGBColor(0xBD, 0xBD, 0xBD)
PALE = RGBColor(0xE6, 0xE6, 0xE6)
HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- Zahlen
G, LB, B = rate_gen(), rate_lb(), rate_bess()
rows = []
for r in P:
    pot = round(r[9] * r[10] * G + r[11] * r[12] * LB + r[13] * r[14] * B, -2)
    exp = round(pot * r[15] * r[16], -2)
    rows.append(dict(seg=r[0], name=r[1], ein=r[2], pot=pot, exp=exp, a27=r[15], p=r[16], raw=r))
rows.sort(key=lambda x: -x["exp"])
TOT_POT = sum(x["pot"] for x in rows)
TOT_EXP = sum(x["exp"] for x in rows)
N_A = sum(1 for x in rows if x["exp"] >= 25000)
SEGS = [RZ, INF, NETZ, SPI, IND, EVT, KAN]
seg_pot = {s: sum(x["pot"] for x in rows if x["seg"] == s) for s in SEGS}
seg_exp = {s: sum(x["exp"] for x in rows if x["seg"] == s) for s in SEGS}
seg_n = {s: sum(1 for x in rows if x["seg"] == s) for s in SEGS}


def chf(v):
    return "CHF " + f"{v:,.0f}".replace(",", "'")


def mio(v, d=1):
    return f"{v / 1e6:.{d}f}".replace(".", ",") + " Mio."


def k(v):
    return f"{v / 1000:.0f}k"


d = Deck(os.path.join(HERE, "Top_Player_Power_Praesentation.pptx"),
         footer="MiT Strom CH · Top-Player Power 2027 · intern · 27.09.2026")


def style_chart(ch, fs=11, legend=True):
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(fs)
        ch.legend.font.name = F_BODY
        ch.legend.font.color.rgb = DARK
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(fs)
    ca.tick_labels.font.name = F_BODY
    ca.tick_labels.font.color.rgb = DARK
    ca.format.line.color.rgb = RGBColor(0xD6, 0xD6, 0xD6)


def labels(ser, color, fs=10, pos=XL_LABEL_POSITION.OUTSIDE_END, fmt='0,"k"'):
    dl = ser.data_labels
    dl.show_value = True
    dl.position = pos
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.font.size = Pt(fs)
    dl.font.bold = True
    dl.font.name = F_BODY
    dl.font.color.rgb = color


def text(s, x, y, w, h, t, size=12, color=DARK, bold=False, font=F_BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = d.tb(s, x, y, w, h, align=align, anchor=anchor)
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = d.para(tf, first=(i == 0), space_before=0 if i == 0 else 3, line=1.05, align=align)
        d.run(p, ln, size, color, font, bold=bold)
    return tf


def oval(s, x, y, dia, fill, label=None, fs=12, fg=WHITE):
    sh = d.rect(s, x, y, dia, dia, fill, MSO_SHAPE.OVAL)
    if label:
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        d.run(p, label, fs, fg, F_TITLE)
    return sh


def line(s, x1, y1, x2, y2, color=MID, w=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    return c


IMG = os.path.join(HERE, "bilder")


def pic(s, name, x, y, w, h, fx=0.5, fy=0.5):
    """Bild einfügen und auf Rahmen zuschneiden (fx/fy = Fokus 0..1)."""
    from PIL import Image
    path = os.path.join(IMG, name)
    iw, ih = Image.open(path).size
    target = w / h
    src = iw / ih
    shp = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if src > target:      # zu breit: links/rechts kappen
        cut = 1 - target / src
        shp.crop_left, shp.crop_right = cut * fx, cut * (1 - fx)
    else:                 # zu hoch: oben/unten kappen
        cut = 1 - src / target
        shp.crop_top, shp.crop_bottom = cut * fy, cut * (1 - fy)
    shp.line.fill.background()
    return shp


def caption(s, x, y, w, t):
    text(s, x, y, w, 0.25, t, 8.5, DARK)


# ================================================================ 1 Cover
s = d.cover(["Top-Player Power", "Schweiz 2027"], "Wo 2027 die Power-Aufträge über CHF 50'000 liegen",
            eyebrow="Mobil in Time AG · An Aggreko Company · Strom Schweiz",
            author="Burak Ücöz · Area Sales Engineer · 27.09.2026",
            stats=[(mio(TOT_POT), "CHF Potenzial brutto"), (mio(TOT_EXP, 2), "CHF Erwartung 2027")])
pic(s, "generator.png", 9.35, 4.95, 2.95, 1.25, fx=0.75, fy=0.62)
d.notes(s, "Neu aufgebaut: nur Projekte mit mehr als CHF 50'000 Power-Potenzial. Jede Zahl ist in der Excel "
           "nachvollziehbar und rechnet mit unseren eigenen Sätzen neu.")

# ================================================================ 2 Kurzfassung
s = d.content("01 · Kurzfassung", f"{len(rows)} Projekte, {mio(TOT_EXP, 2)} CHF realistisch")
d.kpis(s, [(str(len(rows)), "Top-Player über CHF 50'000"), (mio(TOT_POT).replace(" Mio.", ""), "Mio. CHF Potenzial brutto"),
           (mio(TOT_EXP, 2).replace(" Mio.", ""), "Mio. CHF Erwartung 2027 (gewichtet)"), (str(seg_n[RZ]), "Rechenzentren im Fokus"),
           (str(N_A), "Prio A: Erwartung ab CHF 25'000")])
d.band(s, 4.55, "Die grössten Einzeltickets liegen in Rechenzentren. Das sicherste Volumen liefern Tunnel, "
                "Netz und Spitäler mit langen Laufzeiten.", bold_prefix="Kernaussage:", h=1.4)
d.notes(s, "Erwartung 2027 = Potenzial × Anteil 2027 × Gewinnwahrscheinlichkeit. Das ist die Zahl, mit der wir planen. "
           "Das Brutto-Potenzial zeigt, was möglich ist, wenn wir alles gewinnen.")

# ================================================================ 3 Treiber
s = d.content("02 · Warum jetzt", "Drei Treiber machen 2027 zum Power-Jahr")
xs, cw = cols(3)
drv = [
    ("rechenzentrum.png", 0.5, "Rechenzentrums-Boom", ["EKZ meldet über 100 Anschlussanfragen von Rechenzentren (RZ).",
      "6 von 9 neuen EKZ-Unterwerken entstehen für RZ.", "STACK 36 MW, Digital Realty 15 MW, Vantage 100+ MW."], RED),
    ("foto_einspeisung.jpg", 0.2, "Netzengpass", ["Engpass im Vornetz von Axpo und Swissgrid.",
      "Neue Unterwerke kommen später als die Gebäude.", "Dazwischen: Generator- oder BESS-Überbrückung."], DARK),
    ("foto_logger.jpg", 0.45, "Commissioning", ["IST (Integrated Systems Test): Last ≈ IT-Last.",
      "Stufen 25 / 50 / 75 / 100 %, dann Volllast über Stunden.", "4 bis 6 Wochen, Mängel erzwingen Re-Tests."], DARK),
]
for (img, fy, h, t, f), x in zip(drv, xs):
    pic(s, img, x, 2.05, cw, 1.6, fy=fy)
    d.rect(s, x, 3.65, cw, 2.55, f)
    text(s, x + 0.28, 3.82, cw - 0.5, 0.4, h, 15, WHITE, True)
    text(s, x + 0.28, 4.3, cw - 0.5, 1.9, t, 12, WHITE if f == RED else SOFT)
text(s, L, 6.35, 11.9, 0.3, "Quellen: ekz.ch (2025), stackinfra.com, investor.digitalrealty.com (27.08.2026), anvilfield.com, sunbeltsolomon.com. "
     "Fotos: eigene Aufnahmen, Logos unkenntlich.", 8.5, DARK)

# ================================================================ 3b Portfolio
s = d.content("03 · Portfolio Power", "Was wir vermieten, auf einen Blick")
xs, cw = cols(5)
tiles = [
    ("generator.png", 0.72, 0.6, "Generator", "Diesel oder HVO, 30 bis 2'100 kVA. Baustrom, Netzersatz, Events. Leistung in kVA/MVA (Scheinleistung)."),
    ("lastbank_detail.png", 0.5, 0.6, "Lastbank", "Künstliche Last bis 6,25 MW für NEA-Tests und RZ-Commissioning (IST)."),
    ("bess_detail.png", 0.55, 0.5, "BESS", "Battery Energy Storage System: leise, emissionsfrei, kappt Lastspitzen, Hybrid mit Generator."),
    ("foto_trafo.jpg", 0.5, 0.5, "Mobiler Trafo", "Provisorische Einspeisung MS/NS bei Unterwerk- und Stationsumbau."),
    ("foto_messung.jpg", 0.5, 0.5, "NEA-Test & Messung", "Netzersatzanlage unter Last prüfen, Netzqualität messen vor und nach der Umschaltung."),
]
for (img, fx, fy, h, t), x in zip(tiles, xs):
    pic(s, img, x, 2.05, cw, 1.75, fx=fx, fy=fy)
    d.rect(s, x, 3.8, cw, 2.6, LIGHT)
    d.rect(s, x, 3.8, 0.6, 0.045, RED)
    text(s, x + 0.18, 3.98, cw - 0.3, 0.4, h, 14, BLACK, True)
    text(s, x + 0.18, 4.45, cw - 0.32, 2.1, t, 11, DARK)
caption(s, L, 6.5, 11.9, "Illustrationen: eigene Grafik im MiT-Design. Fotos Trafo und Messung: eigene Aufnahmen, Typenschild unkenntlich. Portfolio laut mobilintime.com und aggreko.com.")

# ================================================================ 4 Pipeline
s = d.content("04 · Pipeline", "Wo das Geld liegt: Potenzial und Erwartung")
cd = CategoryChartData()
cats = ["Rechenzentren", "Infrastruktur", "Netz / UW", "Spitäler", "Industrie", "Events", "Kanal-Partner"]
cd.categories = cats
cd.add_series("Potenzial brutto", [seg_pot[x] for x in SEGS])
cd.add_series("Erwartung 2027", [seg_exp[x] for x in SEGS])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(L), Inches(2.0), Inches(8.3), Inches(4.7), cd)
ch = gf.chart
style_chart(ch, 11)
ch.plots[0].gap_width = 60
ch.plots[0].overlap = -10
for ser, col in zip(ch.series, (DARK, RED)):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
    labels(ser, col, 10)
top_seg = max(SEGS, key=lambda x: seg_exp[x])
box_x = 9.05
d.rect(s, box_x, 2.1, 3.35, 2.15, RED)
text(s, box_x + 0.25, 2.3, 2.9, 0.4, "Rechenzentren", 13, WHITE, True)
text(s, box_x + 0.25, 2.72, 2.9, 1.5, [f"{mio(seg_pot[RZ])} CHF brutto", "grösste Einzeltickets,",
                                       "aber Beschaffung oft EMEA"], 12, WHITE)
d.rect(s, box_x, 4.45, 3.35, 2.15, DARK)
text(s, box_x + 0.25, 4.65, 2.9, 0.4, "Infrastruktur", 13, WHITE, True)
text(s, box_x + 0.25, 5.07, 2.9, 1.5, [f"{k(seg_exp[INF])} CHF Erwartung 2027", "Tunnel-Backup über 52 Wochen:",
                                       "planbarstes Volumen"], 12, SOFT)
d.notes(s, "Grau ist das Brutto-Potenzial, rot die gewichtete Erwartung 2027. Rechenzentren haben das grösste Potenzial, "
           "Infrastruktur die höchste Erwartung, weil die Laufzeiten lang und die Termine fix sind.")

# ================================================================ 5 RZ-Blasen-Zeitstrahl
s = d.content("05 · Rechenzentren", f"Die Commissioning-Welle: {mio(seg_pot[RZ])} CHF")
x0, x1 = 1.3, 11.2
m0, m1 = 2026 * 12 + 7, 2028 * 12 + 12   # Jul 2026 bis Dez 2028
X = lambda y, m: x0 + (y * 12 + m - m0) / (m1 - m0) * (x1 - x0)
axis_y = 6.35
line(s, x0, axis_y, 12.2, axis_y, DARK, 1.75)
for yr in (2027, 2028):
    xx = X(yr, 1)
    line(s, xx, axis_y - 0.08, xx, axis_y + 0.08, DARK, 1.5)
for yr, mm in ((2026, 10), (2027, 7), (2028, 7)):
    text(s, X(yr, mm) - 0.6, axis_y + 0.08, 1.2, 0.3, str(yr), 12, DARK, True, F_TITLE, PP_ALIGN.CENTER)
UP, LO = 3.0, 4.75
dcs = [  # name, jahr, monat, MW, wann, lane-y, label-seite
    ("STACK Beringen", 2026, 10, 36, "Herbst 2026", LO, "r"),
    ("Green Lupfig", 2026, 12, 12, "2026/27", UP, "r"),
    ("NorthC Arlesheim", 2027, 6, 4.5, "Mitte 2027", LO, "r"),
    ("NorthC Genf", 2028, 4, 4.5, "Q2 2028", LO, "l"),
    ("Digital Realty ZUR4", 2028, 6, 15, "2028", UP, "l"),
    ("Vantage ZRH3", 2028, 10, 100, "Phase 1 2028", LO, "t"),
]
for nm, yy, mm, mw, when, cy, side in dcs:
    dia = 0.38 + 0.215 * math.sqrt(mw)
    cx = X(yy, mm)
    line(s, cx, cy + dia / 2, cx, axis_y, MID, 1)
    oval(s, cx - dia / 2, cy - dia / 2, dia, RED if mw >= 30 else DARK, f"{mw:g} MW", 11 if dia < 1.0 else (13 if dia < 1.6 else 18))
    if side == "r":
        text(s, cx + dia / 2 + 0.1, cy - 0.25, 1.35, 0.55, [nm, when], 10.5, BLACK)
    elif side == "l":
        text(s, cx - dia / 2 - 1.5, cy - 0.25, 1.4, 0.55, [nm, when], 10.5, BLACK, align=PP_ALIGN.RIGHT)
    else:
        text(s, cx - 1.0, cy - dia / 2 - 0.55, 2.0, 0.5, [nm, when], 10.5, BLACK, align=PP_ALIGN.CENTER)
text(s, x0, 1.95, 7.0, 0.3, "Kreisfläche ≈ IT-Leistung. Zusätzlich: Green Dielsdorf (Bau 2026-28), FlexBase Laufenburg (2028, unsicher).",
     9.5, DARK)
d.notes(s, "Hot Leads sind STACK Beringen und Green Lupfig: Commissioning läuft jetzt bzw. 2027 in Phasen. "
           "Digital Realty und Vantage kommen 2028, der Zugang muss aber 2027 aufgebaut werden, am besten schon mit Baustrom.")

# ================================================================ 6 Top 10
s = d.content("06 · Top 10", "Die zehn Projekte mit der höchsten Erwartung")
top = rows[:10]
cd = CategoryChartData()
cd.categories = [x["name"].replace(" (Neue Axenstrasse)", "").replace(" (nächstes RZ)", "") for x in top][::-1]
cd.add_series("Erwartung 2027", [x["exp"] for x in top][::-1])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(L), Inches(1.95), Inches(8.2), Inches(4.8), cd)
ch = gf.chart
style_chart(ch, 10.5, legend=False)
ch.has_title = False
ch.plots[0].gap_width = 45
ser = ch.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = RED
for i, x in enumerate(top[::-1]):
    pt = ser.points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = RED if x["seg"] == RZ else DARK
labels(ser, DARK, 10)
d.rect(s, 9.0, 2.05, 3.4, 4.6, LIGHT)
text(s, 9.25, 2.25, 3.0, 0.4, "Lesehilfe", 13, BLACK, True)
text(s, 9.25, 2.75, 2.95, 3.8, ["Rot = Rechenzentrum, grau = übrige.", "",
                               "Tunnel führen, weil Backup-Generatoren über ein ganzes Jahr laufen.", "",
                               "Vantage hat das grösste Potenzial (" + k([x for x in rows if x['name'].startswith('Vantage')][0]['pot'])
                               + " CHF), aber nur 10 % Wahrscheinlichkeit."], 11.5, DARK)
d.notes(s, "Die Reihenfolge ergibt sich aus der Erwartung 2027. Ändern sich Wahrscheinlichkeiten, ändert sich das Ranking automatisch in der Excel.")

# ================================================================ 7 Produkt-Matrix
s = d.content("07 · Produktbedarf", "Was die Top-Player an Power brauchen")
prods = ["Generator", "Lastbank", "BESS", "Mobiler Trafo", "NEA-Test / USV", "Tank / HVO"]
mat = {  # 3 = Kernbedarf, 2 = häufig, 1 = gelegentlich, 0 = kaum
    "Rechenzentren": [3, 3, 2, 1, 3, 2],
    "Infrastruktur": [3, 0, 2, 1, 0, 3],
    "Netz / Unterwerk": [3, 0, 1, 3, 1, 1],
    "Spitäler": [2, 2, 2, 1, 3, 1],
    "Industrie": [3, 1, 1, 1, 2, 1],
    "Events": [3, 0, 3, 0, 0, 2],
    "Kanal-Partner": [3, 1, 1, 1, 2, 1],
}
cx0, cy0, cw, chh, lw = L + 2.35, 2.35, 1.55, 0.52, 2.3
for j, pnm in enumerate(prods):
    text(s, cx0 + j * (cw + 0.05), cy0 - 0.4, cw, 0.35, pnm, 11, BLACK, True, F_BODY, PP_ALIGN.CENTER)
fills = {3: RED, 2: DARK, 1: MID, 0: LIGHT}
for i, (sg, vals) in enumerate(mat.items()):
    y = cy0 + i * (chh + 0.05)
    text(s, L, y + 0.12, lw, 0.35, sg, 12, BLACK, True)
    for j, v in enumerate(vals):
        d.rect(s, cx0 + j * (cw + 0.05), y, cw, chh, fills[v])
ly = cy0 + 7 * (chh + 0.05) + 0.15
for i, (v, t) in enumerate(((3, "Kernbedarf"), (2, "häufig"), (1, "gelegentlich"), (0, "kaum"))):
    d.rect(s, cx0 + i * 2.0, ly, 0.3, 0.22, fills[v])
    text(s, cx0 + i * 2.0 + 0.4, ly - 0.02, 1.5, 0.3, t, 10.5, DARK)
d.notes(s, "Generator ist überall Kernprodukt. Lastbank ist das RZ-Produkt, der mobile Trafo das Netzprodukt, "
           "BESS gewinnt bei Events und Nachtbaustellen, weil es leise ist. Einschätzung aus den Szenarien der Excel.")

# ================================================================ 8 Rechenweg
s = d.content("08 · Rechenweg", "So entsteht jede Zahl in der Liste")
bx = [("Generator", f"MVA × Wochen × {chf(G)}"), ("Lastbank", f"MW × Wochen × {chf(LB)}"),
      ("BESS", f"MW × Monate × {chf(B)}")]
for i, (h, t) in enumerate(bx):
    x = L + i * 2.95
    d.rect(s, x, 2.15, 2.6, 1.3, DARK)
    text(s, x + 0.2, 2.3, 2.2, 0.35, h, 14, WHITE, True)
    text(s, x + 0.2, 2.72, 2.3, 0.6, t, 11.5, SOFT)
    if i < 2:
        text(s, x + 2.6, 2.45, 0.35, 0.6, "+", 26, DARK, False, F_TITLE, PP_ALIGN.CENTER)
text(s, L + 8.8, 2.45, 0.4, 0.6, "=", 26, DARK, False, F_TITLE, PP_ALIGN.CENTER)
d.rect(s, L + 9.3, 2.15, 2.63, 1.3, RED)
text(s, L + 9.5, 2.3, 2.3, 0.35, "Potenzial", 14, WHITE, True)
text(s, L + 9.5, 2.72, 2.3, 0.6, "× Anteil 2027 × Wahrscheinlichkeit = Erwartung", 11, WHITE)
st_ = [x for x in rows if x["name"].startswith("STACK")][0]
r_ = st_["raw"]
d.rect(s, L, 3.75, R - L, 1.35, LIGHT)
text(s, L + 0.3, 3.9, 11.3, 0.35, "Beispiel STACK Beringen", 13, BLACK, True)
text(s, L + 0.3, 4.3, 11.3, 0.8,
     f"{r_[9]:g} MVA × {r_[10]} Wochen Generator + {r_[11]:g} MW × {r_[12]} Wochen Lastbank = {chf(st_['pot'])} Potenzial.  "
     f"× {r_[15]:.0%} in 2027 × {r_[16]:.0%} Gewinnchance = {chf(st_['exp'])} Erwartung 2027.".replace(".0%", "%"),
     12.5, DARK)
d.band(s, 5.35, "Sätze sind US-Markt-Richtwerte, umgerechnet mit 0,828 CHF/USD (24.09.2026). Öffentliche CH-Preise ab 100 kVA "
                "gibt es nicht. Mit MiT-Sätzen in der Excel rechnet alles neu. Trafo, Transport, Service und Treibstoff sind nicht enthalten.",
       bold_prefix="Wichtig:", fill=DARK, h=1.3, fs=12.5)
d.notes(s, "Das Modell ist bewusst konservativ. Sobald wir unsere eigenen Sätze eintragen, wird die Zahl belastbar.")

# ================================================================ 9 Cross-Sell
s = d.content("09 · Cross-Sell", "Power öffnet die Tür, Wärme und Kälte folgen")
cx, cy = 6.43, 4.35
fuel = RATES["diesel_lh"] * RATES["diesel_chf"]
cool = RATES["cool_usd_month"] / RATES["cool_mw"] * RATES["fx"]
sat = [
    (1.2, 2.05, "Bauheizung und Trocknung", ["Tunnel, GU-Baustellen, RZ-Rohbau im Winter", "Gotthard, Sisikon, Bachem, Dielsdorf"]),
    (8.45, 2.05, "Temporäre Kühlung", ["RZ-IST, Spital-Umzug, Pharma-Validierung",
                                       "Richtwert ca. " + chf(round(cool, -2)) + " pro MW und Monat"]),
    (1.2, 5.1, "Zelt- und Eventheizung", ["WEF Davos, Ski-WM Crans-Montana", "Heizzentralen sind MiT-Kerngeschäft"]),
    (8.45, 5.1, "BESS-Hybrid und HVO", [f"1 MVA bei 75 % Last: ca. {chf(round(fuel))} Diesel pro Stunde",
                                        "Hybrid senkt Laufzeit, HVO passt zu NorthC"]),
]
for x, y, h, t in sat:
    line(s, cx, cy, x + 1.85, y + 0.65, MID, 1.5)
for x, y, h, t in sat:
    d.rect(s, x, y, 3.7, 1.4, DARK)
    text(s, x + 0.22, y + 0.15, 3.3, 0.35, h, 13.5, WHITE, True)
    text(s, x + 0.22, y + 0.55, 3.3, 0.85, t, 11, SOFT)
oval(s, cx - 1.0, cy - 1.0, 2.0, RED, "Power", 24)
d.notes(s, "Wir gehen mit Power rein und nehmen Wärme und Kälte mit. Beim RZ-IST ist Kühlung oft nicht fertig, "
           "auf Winterbaustellen braucht es Bauheizung, und WEF und Ski-WM brauchen Zeltheizung. Dieselpreis TCS 18.09.2026: 2,41 CHF/l.")

# ================================================================ 10 Infra / Netz / Spital / Industrie
s = d.content("10 · Infrastruktur, Netz, Spital, Industrie", "Das sichere Volumen: lange Laufzeiten")
non = [x for x in rows if x["seg"] in (INF, NETZ, SPI, IND)][:10]
tbl = [["Projekt", "Einstieg über", "Potenzial CHF", "Erwartung 2027"]]
for x in non:
    tbl.append([x["name"], x["ein"].split(";")[-1].strip(), f"{x['pot']:,.0f}".replace(",", "'"), f"{x['exp']:,.0f}".replace(",", "'")])
d.table(s, tbl, colw=[3.3, 2.5, 1.35, 1.35], w=8.3, y=2.05, h=0.42 + 10 * 0.4, fs=10.5, head_fs=10.5)
pic(s, "tunnel.png", 9.0, 2.05, 3.4, 2.1, fy=0.6)
pic(s, "spital.png", 9.0, 4.35, 3.4, 2.1, fx=0.6, fy=0.55)
d.notes(s, "Tunnel, Grimsel und Bachem bringen ganzjährige Mieten. Bei KSA, Axpo Niederurnen und CKW sind die Termine 2027 fix. "
           "Öffentliche Beschaffung bei ASTRA, SBB, ewz, USZ beachten.")

# ================================================================ 11 Events
s = d.content("11 · Events", "Winter 2027 zuerst, ESAF 2028 jetzt offerieren")
evs = [x for x in rows if x["seg"] == EVT]
ev_txt = {
    "WEF Annual Meeting 2027": ["18. bis 22.01.2027, Davos", "Pavillons und Sicherheit bei knappem Ortsnetz", "Aufbau ab Dezember"],
    "FIS Ski-WM Crans-Montana 2027": ["01. bis 14.02.2027", "Broadcast, Fan-Zonen, Hospitality", "Offerte auf Französisch"],
    "ESAF 2028 Thun": ["25. bis 27.08.2028, Thuner Allmend", "Arena, Festgelände, Camping", "Konzept 2027, Umsatz 2028"],
}
pic(s, "event.png", L, 2.05, 3.9, 3.45, fx=0.4)
ev_order = sorted(evs, key=lambda x: ("Ski-WM" not in x["name"], "WEF" not in x["name"]))
for i, x in enumerate(ev_order):
    y = 2.05 + i * 1.18
    f = RED if "Ski-WM" in x["name"] else DARK
    d.rect(s, 4.6, y, 7.8, 1.08, f)
    text(s, 4.85, y + 0.14, 4.3, 0.4, x["name"].replace(" Annual Meeting", ""), 14, WHITE, True)
    text(s, 4.85, y + 0.52, 4.6, 0.5, " · ".join(ev_txt.get(x["name"], [])[:2]), 11, WHITE if f == RED else SOFT)
    text(s, 9.6, y + 0.14, 2.6, 0.4, chf(x["pot"]), 16, WHITE, True, F_TITLE, PP_ALIGN.RIGHT)
    text(s, 9.6, y + 0.58, 2.6, 0.4, ev_txt.get(x["name"], ["", "", ""])[2], 10.5, WHITE if f == RED else SOFT, align=PP_ALIGN.RIGHT)
d.band(s, 5.75, "Sommer-Openairs (St.Gallen, Frauenfeld, Gurten, Paléo, Montreux) liegen einzeln unter CHF 50'000: als Saisonpaket anbieten.",
       fill=LIGHT, h=0.85, fs=12.5)
d.notes(s, "Szenario-Grössen bei Events sind Annahmen. Die Ski-WM ist rot markiert, weil sie 2027 stattfindet und die Beschaffung jetzt läuft.")

# ================================================================ 12 Wettbewerb
s = d.content("12 · Wettbewerb", "Wer sonst mobile Power anbietet")
tbl = [["Anbieter", "Angebot laut eigener Website"]] + [[a, b] for a, b, _ in COMP]
d.table(s, tbl, colw=[3.2, 8.7], y=2.05, h=0.42 + len(COMP) * 0.38, fs=11, head_fs=11)
d.band(s, 5.35, "Lastbänke bis 6,25 MW mit Commissioning Level 1 bis 5, BESS, Kühlung und Heizung aus einer Hand, "
                "mit der Aggreko-Flotte im Rücken.", bold_prefix="Unser Unterschied:", h=1.3, fs=13)
d.notes(s, "Quelle Aggreko-Portfolio: aggreko.com (Data Centre Commissioning, Load Banks). Im Markt sind vor allem "
           "Generator-Vermieter bis 2 MVA; kaum jemand bietet das komplette Commissioning-Paket.")

# ================================================================ 13 90 Tage
s = d.content("13 · Die nächsten 90 Tage", "Vom Modell zum Auftrag bis Januar 2027")
d.steps(s, [
    ("01", "Oktober", ["Hot Leads: STACK, Green Lupfig, KSA, Ski-WM.", "Commissioning- und Umzugstermine klären."]),
    ("02", "November", ["Rahmengespräche Burkhalter, VINCI, Equans.", "WEF-Offerte abgeben."]),
    ("03", "Dezember", ["MiT-Sätze ins Modell, CRM-Status eintragen.", "Pipeline neu rechnen."]),
    ("04", "Januar", ["Review mit Mauro, Jörg, Roberto.", "RZ 2028 und Tunnel-Rahmen aufgleisen."]),
], y=2.3)
d.band(s, 5.55, "Termine sind ein Vorschlag, bitte im Team bestätigen.", fill=LIGHT, h=0.85, fs=12.5)
d.notes(s, "Reihenfolge nach Dringlichkeit: Commissioning und Events laufen, bevor die grossen RZ 2028 kommen.")

# ================================================================ Glossar
s = d.content("Anhang · Glossar", "Abkürzungen in diesem Deck")
g1 = [["Kürzel", "Bedeutung"],
      ["BESS", "Battery Energy Storage System, Batteriespeicher"],
      ["NEA", "Netzersatzanlage (Notstromaggregat)"],
      ["USV", "Unterbrechungsfreie Stromversorgung"],
      ["IST / L1-L5", "Integrated Systems Test, RZ-Gesamttest = Level 5"],
      ["kVA / MVA", "Scheinleistung (Nennleistung Generator)"],
      ["kW / MW", "Wirkleistung; bei cos phi 0,8: 1 MVA ≈ 0,8 MW"],
      ["IT-MW", "Elektrische Leistung der Server im RZ"],
      ["UW / GIS", "Unterwerk / gasisolierte Schaltanlage"],
      ["HS / MS / NS", "Hoch-, Mittel-, Niederspannung"],
      ["HVO", "Hydriertes Pflanzenöl, erneuerbarer Diesel"],
      ["TBM / BSA", "Tunnelbohrmaschine / Tunnel-Sicherheitstechnik"]]
g2 = [["Kürzel", "Bedeutung"],
      ["RZ", "Rechenzentrum"],
      ["GU / TU / ARGE", "General-, Totalunternehmer / Arbeitsgemeinschaft"],
      ["IBN / PM", "Inbetriebnahme / Projektmanagement"],
      ["IVöB / simap", "Beschaffungsrecht / Ausschreibungsportal"],
      ["EMEA", "Europa, Nahost, Afrika: zentraler Einkauf"],
      ["GMP", "Good Manufacturing Practice (Pharma-Regeln)"],
      ["ASTRA / SBB", "Bundesamt für Strassen / Schweizerische Bundesbahnen"],
      ["EKZ / ewz / CKW / IWB", "Stromversorger Kt. Zürich, Stadt Zürich, Zentralschweiz, Basel"],
      ["KWO / KSA / USZ / KSSG", "Kraftwerke Oberhasli; Spitäler Aarau, Zürich, St.Gallen"],
      ["WEF / FIS / ESAF", "World Economic Forum / Ski-Weltverband / Schwingfest"],
      ["CRM / FX", "Kundendatenbank / Wechselkurs"]]
d.table(s, g1, colw=[1.6, 4.2], x=L, w=5.85, y=2.05, h=0.42 + 11 * 0.38, fs=10, head_fs=10.5)
d.table(s, g2, colw=[1.9, 3.95], x=L + 6.08, w=5.85, y=2.05, h=0.42 + 11 * 0.38, fs=10, head_fs=10.5)

# ================================================================ 14 Schluss
d.closing("Der eine Hebel", ["Rechenzentren früh", "im Bau besetzen,", "nicht erst beim IST."],
          ["Burak Ücöz · Area Sales Engineer", "Mobil in Time AG · An Aggreko Company"])

from notizen import build_notes  # noqa: E402
NOTES = build_notes(dict(n=len(rows), pot=mio(TOT_POT), exp=chf(TOT_EXP), na=N_A, rz_pot=mio(seg_pot[RZ]),
                         inf_exp=chf(seg_exp[INF]), fuel=f"{RATES['diesel_lh'] * RATES['diesel_chf']:.0f}",
                         top10=f"{sum(x['exp'] for x in rows[:10]) / TOT_EXP:.0%}".replace("%", " %")))
for sl, n in zip(d.prs.slides, NOTES):
    d.notes(sl, n)
assert len(NOTES) == len(d.prs.slides), (len(NOTES), len(d.prs.slides))
print(d.save(title="Top-Player Power Schweiz 2027", author="Burak Ücöz"))
print(len(rows), TOT_POT, TOT_EXP, N_A)
