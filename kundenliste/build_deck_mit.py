# -*- coding: utf-8 -*-
"""Zielkunden Schweiz 2026/27 im MiT-Master (Skill mit-praesentation).
Zahlen aus Kundenliste_MiT_Schweiz.xlsx, Stand 27.09.2026."""
import glob
import os
import sys

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ASSETS = os.environ.get("MIT_ASSETS") or (glob.glob(
    "/root/.claude/skills/synced/*/mit-praesentation/assets") or [""])[0]
sys.path.insert(0, ASSETS)
from mit_deck import Deck, RED, DARK, DARKER, LIGHT, WHITE, BLACK, F_BODY, F_TITLE, L, R, cols  # noqa: E402

GREY_C = RGBColor(0xBD, 0xBD, 0xBD)
HERE = os.path.dirname(os.path.abspath(__file__))
d = Deck(os.path.join(HERE, "Zielkunden_Schweiz_Praesentation.pptx"),
         footer="MiT Strom CH · Zielkunden Schweiz 2026/27 · intern · 27.09.2026")


def stacked_bar(s, x, y, w, h, cats, series, fs=11):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(fs)
    ch.legend.font.name = F_BODY
    ch.legend.font.color.rgb = DARK
    ch.plots[0].gap_width = 45
    ch.plots[0].overlap = 100
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(fs)
    ca.tick_labels.font.name = F_BODY
    ca.tick_labels.font.color.rgb = DARK
    ca.format.line.color.rgb = RGBColor(0xD6, 0xD6, 0xD6)
    for ser, col, lab in zip(ch.series, (RED, DARK, GREY_C), (WHITE, WHITE, DARK)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
        ser.data_labels.show_value = True
        ser.data_labels.position = XL_LABEL_POSITION.CENTER
        ser.data_labels.font.size = Pt(fs - 1)
        ser.data_labels.font.bold = True
        ser.data_labels.font.name = F_BODY
        ser.data_labels.font.color.rgb = lab
        ser.data_labels.number_format = '0;;;'
        ser.data_labels.number_format_is_linked = False
    return ch


def box(s, x, y, w, h, head, lines, fill=DARK, hs=14, bs=12):
    d.rect(s, x, y, w, h, fill)
    tf = d.tb(s, x + 0.28, y + 0.22, w - 0.56, 0.4)
    d.run(tf.paragraphs[0], head, hs, WHITE, F_BODY, bold=True)
    tf = d.tb(s, x + 0.28, y + 0.62, w - 0.56, h - 0.8)
    for j, t in enumerate(lines):
        p = d.para(tf, first=(j == 0), space_before=0 if j == 0 else 3, line=1.05)
        d.run(p, t, bs, WHITE if fill == RED else RGBColor(0xE6, 0xE6, 0xE6))


# ------------------------------------------------------------------ 1 Cover
s = d.cover(["Zielkunden Schweiz", "2026/27"], "Mobile Energie: Kundenliste nach Vorgabe 60/20/20",
            eyebrow="Mobil in Time AG · An Aggreko Company · Strom Schweiz",
            author="Burak Ücöz · Area Sales Engineer · 27.09.2026",
            stats=[("125", "Firmen recherchiert"), ("70", "in der Kernliste")])
d.notes(s, "Roberto und ich haben die Vorgabe gemacht, hier ist die Umsetzung für die ganze Schweiz. "
           "125 Firmen recherchiert, 70 als Kernliste im Verhältnis 60/20/20. Die Details stehen in der Excel.")

# ------------------------------------------------------------------ 2 Auftrag
s = d.content("01 · Auftrag", "Vorgabe umgesetzt, und etwas mehr")
d.cards(s, [
    ("Vorgabe", ["50 bis 70 Kunden mit Namen.",
                 "Mix: Installateure 60 %, FM / Stadtwerke / Planer 20 %, GU/TU + Industrie 20 %.",
                 "Status: Bestandskunde, verlorener Kunde, Potenzial / Aufbau.",
                 "Lieber mehr als weniger, damit wir priorisieren können."]),
    ("Umgesetzt", ["125 Firmen recherchiert, jede mit Quelle.",
                   "Kernliste 70: exakt 42 / 14 / 14.",
                   "Score aus vier Kriterien, Prio A/B/C mit Begründung.",
                   "Status-Spalte vorbereitet, der CRM-Abgleich ist offen."], RED),
], cols_n=2, h=3.3, bs=14)
d.band(s, 5.7, "Der MiT-Kundenstatus (A-Kunde, Bestand, verloren, Aufbau) steht nicht im Netz. Er kommt aus dem CRM.",
       bold_prefix="Offen:", fill=LIGHT, h=0.95, fs=13.5)
d.notes(s, "Den MiT-Kundenstatus liefert keine Webrecherche. Alle Firmen stehen auf 'Offen: CRM prüfen'. "
           "Das ist der Teil, den wir gemeinsam machen müssen.")

# ------------------------------------------------------------------ 3 Kennzahlen
s = d.content("02 · Auf einen Blick", "36 Firmen sind reif für den Erstkontakt")
d.kpis(s, [("125", "Firmen gesamt"), ("70", "Kernliste 60/20/20"), ("36", "Prio A: Score ab 4,0"),
           ("69", "Prio B: Score ab 3,0"), ("20", "Prio C: beobachten")])
d.band(s, 4.55, "A-Firmen haben einen belegten Anlass 2026/27 oder sind Gruppen mit Rahmenvertrags-Hebel. "
                "Sie gehören in den nächsten sechs Wochen an den Tisch.", bold_prefix="Kernaussage:", h=1.4)
d.notes(s, "A: jetzt angehen, Termin in sechs Wochen. B: im Quartal angehen, Anlass beobachten. "
           "C: beobachten, über Partner oder Töchter abdecken.")

# ------------------------------------------------------------------ 4 Segmentmix
s = d.content("03 · Segmentmix", "Mix 60/20/20 erfüllt, A-Dichte bei GU/TU")
stacked_bar(s, L, 2.0, 6.6, 4.1, ["Installateure", "FM / EVU / Planer", "GU/TU + Industrie"],
            [("Prio A", (11, 7, 18)), ("Prio B", (20, 25, 24)), ("Prio C", (11, 3, 6))])
d.table(s, [["Segment", "Gesamt", "Kern", "Soll"],
            ["Installateure", "42", "42", "60 %"],
            ["FM / EVU / Planer", "35", "14", "20 %"],
            ["GU/TU + Industrie", "48", "14", "20 %"],
            ["Total", "125", "70", "100 %"]],
        colw=[2.3, 1.0, 0.9, 0.9], x=7.35, w=5.05, y=2.1, h=2.2, fs=11)
d.rect(s, 7.35, 4.6, 5.05, 1.95, LIGHT)
tf = d.tb(s, 7.6, 4.8, 4.55, 1.6)
d.run(tf.paragraphs[0], "Die Hälfte aller A-Firmen sitzt im kleinsten Kernsegment. Dort liegen die Anlässe: "
                        "Rechenzentren, Pharma-Neubauten, Tunnel, Flughäfen.", 13, DARK)
d.notes(s, "Kernliste exakt 42/14/14. Die Reserve von 55 Firmen liegt vor allem bei EVU/Planern und GU/Industrie.")

# ------------------------------------------------------------------ 5 Methodik
s = d.content("04 · Methodik", "Vier Kriterien entscheiden die Prio")
d.cards(s, [("Fit · 30 %", ["Passt der Bedarf zum MiT-Portfolio: Generator, BESS, USV, Trafo, NEA-Test, PQ, Baustrom, Event?"]),
            ("Volumen · 30 %", ["Mietvolumen pro Jahr und Hebel über Töchter und Rahmenverträge."]),
            ("Timing · 25 %", ["Konkreter Anlass 2025 bis 2027: Bau, UW-Umbau, Commissioning, Event-Saison."]),
            ("Zugang · 15 %", ["Regional erreichbar? Entscheid lokal oder im Konzern? Öffentliche Beschaffung?"])],
        cols_n=4, h=2.85, hs=15, bs=12)
d.band(s, 5.25, "A ab Score 4,0 · B ab 3,0 · C darunter. Skala je Kriterium 1 bis 5. Gewichte und Schwellen "
                "sind in der Excel änderbar, «Prio final» übersteuert den Vorschlag.",
       bold_prefix="Schwellen:", fill=LIGHT, h=1.3, fs=13)
d.notes(s, "Das Scoring ist transparent und diskutierbar. Wer eine Firma anders sieht, setzt Prio final.")

# ------------------------------------------------------------------ 6 Regionen
s = d.content("05 · Regionen", "Zürich und Nordwestschweiz: 22 von 36 A")
regs = ["Tessin", "Zentralschweiz", "Ostschweiz", "Mittelland/Bern", "West (Romandie)", "Nordwestschweiz", "Zürich"]
stacked_bar(s, L, 1.95, 7.0, 4.75, regs, [("Prio A", (1, 3, 2, 4, 4, 11, 11)),
                                          ("Prio B", (6, 4, 8, 9, 17, 10, 15)),
                                          ("Prio C", (6, 3, 1, 3, 6, 1, 0))], fs=10.5)
for i, (h, t, f) in enumerate([
        ("Zürich 26 · NWCH 22", ["RZ-Cluster, Pharma Basel, Konzernzentralen: hier entscheidet sich das Jahr."], RED),
        ("West 27, davon 4 A", ["Aufbaugebiet mit vielen Namen und wenig Anlass-Treffern. Französisch ist Pflicht."], DARK),
        ("Tessin 13, davon 1 A", ["Kleine Installateure, Gotthard, Events. Über Partner und Marti bedienen."], DARK)]):
    box(s, 7.8, 2.05 + i * 1.55, 4.6, 1.4, h, t, f, hs=13.5, bs=11.5)
d.notes(s, "Region nach Hauptsitz bzw. relevanter Baustelle. Lonza Visp zählt zu West, ist aber Oberwallis und deutschsprachig.")

# ------------------------------------------------------------------ 7 Installateure
s = d.content("06 · Installateure · 60 %", "Rahmenverträge schlagen Einzelbaustellen")
rows = [["Prio A", "Region", "Score", "Hebel"],
        ["Burkhalter Gruppe", "Zürich", "4,45", "83 Gesellschaften, ein Rahmen für Dutzende Töchter"],
        ["ETAVIS AG (VINCI)", "Zürich", "4,45", "Infrastruktur und RZ, Tür zu Actemium und Gfeller"],
        ["Equans Switzerland", "Zürich", "4,45", "Energie, Verkehr, Industrie, FM unter einem Dach"],
        ["BKW Building Solutions", "Bern", "4,20", "Dach von AEK, ISP, swisspro, Arnold, inelectro"],
        ["Actemium Schweiz", "Bern", "4,15", "Shutdowns, Inbetriebnahmen, Lastbanktests"],
        ["Arnold AG (BKW)", "Bern", "4,15", "HS- und Kabelnetzbau: Trafo-Miete, Netzersatz"],
        ["CKW Gebäudetechnik", "Luzern", "4,15", "Grösster Installateur der Zentralschweiz"],
        ["EKZ Eltop", "Zürich", "4,15", "Ganzer Kanton Zürich, PV und Speicher"],
        ["ETAVIS Kriegel+Schaffner", "Basel", "4,15", "Pharma-Grossprojekte Basel"],
        ["Selmoni Gruppe", "Basel", "4,15", "Life Sciences: USV und NEA bei Shutdowns"],
        ["Egg-Telsa", "Genf", "4,15", "Grösster Unabhängiger in Genf, Cap2030 GVA"]]
d.table(s, rows, colw=[2.5, 1.0, 0.75, 4.35], w=8.6, y=2.05, h=0.42 + 11 * 0.375, fs=10.5, head_fs=10.5)
box(s, 9.3, 2.05, 3.1, 4.55, "42 Installateure",
    ["11 A · 20 B · 11 C", "", "Burkhalter, VINCI und BKW decken zusammen 19 Firmen der Liste ab.",
     "", "Einkauf zentral klären, dann Töchter regional bearbeiten."], RED, hs=15, bs=12)
d.notes(s, "Burkhalter-Töchter in der Liste: Sedelec, Mérinat, Grichting & Valterio, EAGB, Hufschmid, Celio. "
           "VINCI: ETAVIS mit Töchtern, Gfeller, Actemium. BKW: AEK, ISP, swisspro, Arnold, inelectro.")

# ------------------------------------------------------------------ 8 FM / EVU / Planer
s = d.content("07 · FM, Stadtwerke, Planer · 20 %", "Netzumbau braucht Trafos, Planer schreiben aus")
d.cards(s, [
    ("Stadtwerke / EVU", ["CKW: ca. 2'500 Trafostationen", "EKZ: 17'187 km Netz in Erneuerung",
                          "ewz: Unterwerke, z.B. Oerlikon", "IWB: Netzknoten Basel verstärkt"], RED),
    ("Facility Manager", ["CBRE GWS: Pharmamandate Basel, NEA-Lasttests",
                          "Equans FM: Fusion 02/2026, Lieferanten werden neu gewählt"]),
    ("Elektroplaner", ["Amstein + Walthert: >1'100 MA, legt NEA und USV für Spitäler und RZ fest"]),
], cols_n=3, h=3.0, bs=12)
d.band(s, 5.4, "Planer kaufen nicht selbst. Ziel ist, dass Mietlösungen in NEA-, USV- und Provisoriums-"
               "Ausschreibungen genannt werden.", bold_prefix="7 A-Firmen im Segment.", fill=LIGHT, h=1.15, fs=13)
d.notes(s, "EVU-Beschaffung ist teils öffentlich (IVöB), das braucht Vorlauf. Bei Planern zählt Präsenz in der Projektierung.")

# ------------------------------------------------------------------ 9 RZ-Welle
s = d.content("08 · Rechenzentren und Grossprojekte", "Die Commissioning-Welle 2026 bis 2028")
d.steps(s, [
    ("26", "2026", ["STACK Beringen, 36 MW", "Green Lupfig, 12 MW", "NorthC Genf, Baustart", "Cap2030 GVA, Etappe 1"]),
    ("27", "2027", ["NorthC Arlesheim, Mitte 2027", "Flughafen ZH, Tower ab 2027", "USZ Campus MITTE1|2"]),
    ("28", "2028", ["Digital Realty ZUR4, 15 MW", "FlexBase Laufenburg (Zeitplan umstritten)", "Vantage ZRH2, 24 MW"]),
], y=2.2)
d.band(s, 5.5, "Wer im Bau den Baustrom liefert, sitzt beim Commissioning mit Lastbank und NEA-Test am Tisch.",
       bold_prefix="Kernaussage:", h=1.1, fs=14)
d.notes(s, "Einstieg über den GU (Implenia bei Green, ERNE bei FlexBase) oder direkt beim Betreiber. "
           "Hyperscaler sprechen oft Englisch, Beschaffung teils über EMEA.")

# ------------------------------------------------------------------ 10 GU/TU + Industrie
s = d.content("09 · GU/TU und Industrie · 20 %", "18 A-Firmen mit belegtem Anlass")
rows = [["Firma", "Typ", "Score", "Anlass (Quelle in der Excel)"],
        ["Implenia", "GU", "4,70", "GU Green-RZ Lupfig/Dielsdorf, Sisikoner Tunnel bis 2034"],
        ["Digital Realty", "RZ", "4,70", "Spatenstich ZUR4 am 26.08.2026"],
        ["Green Datacenter", "RZ", "4,70", "Lupfig Start 2026, Dielsdorf RZ 3"],
        ["STACK Infrastructure", "RZ", "4,55", "Beringen SH, 36 MW"],
        ["Losinger Marazzi", "GU", "4,45", "Terminal GVA Cap2030, ca. 600 Mio. CHF"],
        ["Marti-Gruppe", "GU", "4,45", "2. Gotthardröhre, Hauptlos Süd"],
        ["ERNE / Frutiger", "GU", "4,40", "FlexBase Laufenburg / Sisikoner Tunnel"],
        ["NorthC", "RZ", "4,40", "Arlesheim und Genf parallel im Bau"],
        ["Bachem", "Industrie", "4,40", "Bau K Ramp-up, Werk Sisslerfeld 750 Mio. CHF"],
        ["Roche / Lonza", "Industrie", "4,30 / 4,00", "Bau 12 Basel / Ibex Visp Ramp-up 2026"],
        ["FlexBase / Vantage", "RZ", "4,30", "BESS 1,6 GWh plus KI-RZ / ZRH2 24 MW"],
        ["USZ, GVA, SBB Energie, OASG", "div.", "4,00 bis 4,15", "Spital-Umbau, Flughafen, Umformer, Openair"]]
d.table(s, rows, colw=[3.0, 1.2, 1.4, 6.3], y=2.05, h=0.42 + 12 * 0.35, fs=10.5, head_fs=10.5)
d.notes(s, "Fakten mit Quelle belegt, siehe Spalten Anlass und Quelle in der Excel. Industrie-Einkauf läuft oft zentral: "
           "Zugang über Installateur oder FM vor Ort (ETAVIS K+S, Selmoni, CBRE).")

# ------------------------------------------------------------------ 11 Korrekturen
s = d.content("10 · Was die Recherche korrigiert", "Sechs Punkte, bevor wir anrufen")
d.cards(s, [
    ("Steiner AG raus", ["Nachlassstundung, kein TU-Neugeschäft in der Deutschschweiz seit 2023."]),
    ("Compagnoni unabhängig", ["Elektro Compagnoni ist keine Burkhalter-Tochter, sondern Familienfirma."]),
    ("Kummler+Matter", ["Vermietet eigene mobile Trafostationen: Partner und Wettbewerber zugleich."], RED),
    ("Stahl bewusst C", ["Gerlafingen und Swiss Steel laufen mit Staatshilfe: Bonität prüfen."]),
    ("Öffentliche Beschaffung", ["ewz, USZ, SBB, CSCS: IVöB-Verfahren und Fristen einplanen."]),
    ("BKW am Lauberhorn", ["BKW ist Sponsor: Konkurrenzlage vor der Ansprache klären."]),
], cols_n=3, rows=2, h=2.1, gap_y=0.22, hs=15, bs=12)
d.notes(s, "Zusätzlich: Mitarbeiterzahlen teils aus Sekundärquellen, in der Excel mit 'prüfen' markiert. "
           "Ansprechpersonen bewusst nicht erfunden, die LinkedIn-Spalten öffnen nur die Suche.")

# ------------------------------------------------------------------ 12 Nächste Schritte
s = d.content("11 · Nächste Schritte (Vorschlag)", "Vom Ranking zur Telefonliste in sechs Wochen")
d.steps(s, [
    ("01", "CRM-Abgleich", ["KW 40-41 · bis 09.10.2026", "Status je Firma: A-Kunde, Bestand, verloren mit Datum, Aufbau."]),
    ("02", "Zuteilung", ["KW 41 · bis 09.10.2026", "Verantwortliche und Prio final je Firma setzen."]),
    ("03", "Erstkontakt A", ["KW 42-45 · 12.10.-06.11.", "Alle 36 A-Firmen: Kontakt verifizieren, Termin anfragen."]),
    ("04", "Review", ["KW 46 · ab 09.11.2026", "Pipeline prüfen, B-Firmen mit neuem Anlass hochstufen."]),
], y=2.3)
d.band(s, 5.6, "Termine sind ein Vorschlag, bitte mit Mauro und Jörg bestätigen.", fill=LIGHT, h=0.9, fs=13)
d.notes(s, "Konkret brauche ich von Mauro und Jörg bis 09.10.2026 den Status je Firma aus dem CRM, "
           "bei verlorenen Kunden das Datum und den Grund, falls bekannt.")

# ------------------------------------------------------------------ 13 Schluss
d.closing("Der eine Hebel", ["Der CRM-Abgleich", "macht aus A/B/C", "eine Telefonliste."],
          ["Burak Ücöz · Area Sales Engineer", "Mobil in Time AG · An Aggreko Company"])

print(d.save(title="Zielkunden Schweiz 2026/27", author="Burak Ücöz"))
