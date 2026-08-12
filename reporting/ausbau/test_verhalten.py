#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Verhaltenstest auf einer Wegwerfkopie — beweist, dass die neuen Teile
LIVE rechnen (nicht nur im heutigen Datenstand stimmen).

Szenario A: Bei einem aktiven Deal wird ein Projektzeitraum nachgetragen
            (01.08.2026 – 10.08.2026).
  A1  Die Erfassungsliste im Definitionsblatt schrumpft von 46 auf 45.
  A2  Der Hinweis im CEO Report zeigt «45 von 46 … offen».
  A3  Die Startmonat-Tabelle zaehlt den Deal im August (Anzahl 1, Volumen).
  A4  Die Dauer (Spalte G) betraegt 10 Kalendertage.
  A5  Das Startmonat-Diagramm (_data!AA9 = August) zeigt das Volumen.

Szenario B: Projektende VOR Projektstart.
  B1  Kurzstatus BB meldet «⚠ Ende vor Start».
  B2  Der Deal zaehlt NICHT in die Monatstabelle (AY=0), Liste bleibt 46.

Szenario C: Eine WON-Zeile erhaelt PO-Nr., Belegdatum und die Vertragsart
            «Einzelauftrag».
  C1  «davon belegt» steigt (sofern Einstand vollstaendig), «ohne Angabe»
      in der Vertragsarten-Tabelle sinkt um 1.
  C2  Executive PDF Top-Offerten bleiben unberuehrt.
"""
import datetime
import shutil
import subprocess
import sys
import openpyxl

QUELLE = 'CH_MiT_Strom_Customer_CEO_CFO_MASTER__2_.xlsx'
KOPIE = 'test_szenario.xlsx'
RECALC = '/root/.claude/skills/synced/xlsx/scripts/recalc.py'
fehler = []


def pruef(name, bedingung, detail=''):
    print(('  ✔ ' if bedingung else '  ✘ ') + name + (f'  [{detail}]' if detail else ''))
    if not bedingung:
        fehler.append(name)


def recalc():
    subprocess.run(['pkill', '-9', '-f', 'soffice.bin'], capture_output=True)
    r = subprocess.run([sys.executable, RECALC, KOPIE, '420'],
                       capture_output=True, text=True)
    assert '"status": "success"' in r.stdout, r.stdout + r.stderr


# ---------------------------------------------------------------- Szenario A
shutil.copy(QUELLE, KOPIE)
wb = openpyxl.load_workbook(KOPIE)
p = wb['MiT Strom Pipeline']
ziel = None
for r in range(6, 72):
    if p[f'B{r}'].value and p[f'R{r}'].value == 'offered':
        ziel = r
        break
vol_ziel = p[f'I{ziel}'].value
p[f'O{ziel}'] = datetime.date(2026, 8, 1)
p[f'P{ziel}'] = datetime.date(2026, 8, 10)
wb.save(KOPIE)
recalc()

V = openpyxl.load_workbook(KOPIE, data_only=True)
vp, vd, vc = V['MiT Strom Pipeline'], V['📋 Definitionen & Klärung'], V['CEO Report']
print(f'Szenario A (Zeile {ziel}, Volumen {vol_ziel}):')
belegte = sum(1 for r in range(82, 132) if vd[f'B{r}'].value not in (None, ''))
pruef('A1 Erfassungsliste schrumpft auf 45', belegte == 45, f'{belegte}')
pruef('A2 CEO-Hinweis zeigt 45 von 46',
      '45 von 46' in str(vc['A250'].value), str(vc['A250'].value)[:60])
pruef('A3 Monatstabelle August: 1 Deal + Volumen',
      vc['E242'].value == 1 and abs((vc['F242'].value or 0) - vol_ziel) < 0.01,
      f"{vc['E242'].value} / {vc['F242'].value}")
pruef('A4 Dauer = 10 Kalendertage', vp[f'G{ziel}'].value == 10,
      str(vp[f'G{ziel}'].value))
pruef('A5 Diagrammdaten August', abs((V['_data']['AA9'].value or 0) - vol_ziel) < 0.01,
      str(V['_data']['AA9'].value))

# ---------------------------------------------------------------- Szenario B
shutil.copy(QUELLE, KOPIE)
wb = openpyxl.load_workbook(KOPIE)
p = wb['MiT Strom Pipeline']
p[f'O{ziel}'] = datetime.date(2026, 8, 10)
p[f'P{ziel}'] = datetime.date(2026, 8, 1)
wb.save(KOPIE)
recalc()
V = openpyxl.load_workbook(KOPIE, data_only=True)
vp, vd, vc = V['MiT Strom Pipeline'], V['📋 Definitionen & Klärung'], V['CEO Report']
print('Szenario B:')
pruef('B1 Kurzstatus meldet Ende vor Start',
      'Ende vor Start' in str(vp[f'BB{ziel}'].value), str(vp[f'BB{ziel}'].value))
belegte = sum(1 for r in range(82, 132) if vd[f'B{r}'].value not in (None, ''))
pruef('B2 Liste bleibt 46, August bleibt leer',
      belegte == 46 and (vc['E242'].value or 0) == 0, f'{belegte} / {vc["E242"].value}')

# ---------------------------------------------------------------- Szenario C
shutil.copy(QUELLE, KOPIE)
wb = openpyxl.load_workbook(KOPIE)
p = wb['MiT Strom Pipeline']
won = next(r for r in range(6, 72) if p[f'R{r}'].value == 'WON')
p[f'AF{won}'] = 'PO-2026-0815'                      # Auftrags-/PO-Nr. (AF)
p[f'AG{won}'] = datetime.date(2026, 8, 1)           # Belegdatum (AG)
p[f'AB{won}'] = 'Einzelauftrag'                     # Vertragsart
wb.save(KOPIE)
recalc()
V = openpyxl.load_workbook(KOPIE, data_only=True)
vc = V['CEO Report']
vpdf = V['📑 Executive PDF']
print(f'Szenario C (WON-Zeile {won}):')
pruef('C1 Vertragsart «ohne Angabe» sinkt auf 11, «Einzelauftrag» = 1',
      vc['D259'].value == 11 and vc['D255'].value == 1,
      f"{vc['D259'].value} / {vc['D255'].value}")
pruef('C1b Vertragsarten-TOTAL weiterhin = WON gemeldet',
      '✔' in str(vc['F260'].value), str(vc['F260'].value)[:50])
pruef('C2 Top-Offerte 1 unveraendert', vpdf['C29'].value == 'Rimus & Sstrada',
      str(vpdf['C29'].value))

print()
print(f'Verhaltenstest: {len(fehler)} Fehler')
sys.exit(1 if fehler else 0)
