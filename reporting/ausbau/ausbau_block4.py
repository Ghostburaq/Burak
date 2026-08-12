#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BLOCK 4 — Diagramme im einheitlichen Design.

Design: Dunkelblau 15243A als Hauptfarbe, Orange F47B20 als Akzent,
keine 3D-Effekte, keine Legende wo eine Farbe reicht, Datenbeschriftung
direkt am Balken, Tausendertrennzeichen.

4.1  Balken horizontal: aktive Pipeline nach Segment, absteigend.
4.2  Balken horizontal: aktive Pipeline nach Kanton, Top 10.
4.3  Saeulen: Volumen nach Status — WON in Orange.
4.4  Saeulen: gewichtete Pipeline zerlegt in WON und offene Offerten.
4.5  Saeulen (vorbereitet): Volumen nach Startmonat — fuellt sich,
     sobald Projektzeitraeume erfasst sind.

Alle Diagrammquellen sind Formelbereiche im versteckten Blatt _data
(Formeln auf die Pipeline, Zeilen 6–860) — keine statischen Wertelisten,
keine Ganzspaltenbezuege. Die beiden Dashboard-Diagramme bekommen
dasselbe Design.
"""
import copy
import openpyxl
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Font, Alignment
from ausbau_util import DATEI, sp

BLAU = '15243A'
ORANGE = 'F47B20'
PALETTE = [ORANGE, BLAU, '3E5C76', '7C93AD', 'B45309', '5A748C',
           '2C425C', 'D98C4A', '8A9BB0', 'C0CCD9']

wb = openpyxl.load_workbook(DATEI)
data = wb['_data']
diag = wb['📊 Diagramme']
dash = wb['Dashboard']
P = "'MiT Strom Pipeline'!"

# ============================================== Datenbereiche in _data
# Kopfzeilen zur Selbstdokumentation des versteckten Blatts.
KOEPFE = {'J1': 'Volumen aktiv', 'Q1': '_RankSeg', 'R1': 'Segment sortiert',
          'S1': 'Volumen sortiert', 'T13': 'Volumen aktiv', 'U13': '_RankKt',
          'V1': 'Kanton Top10', 'W1': 'Volumen Top10',
          'X1': 'Zerlegung', 'Y1': 'Gewichtet CHF',
          'Z1': 'Monat', 'AA1': 'Volumen Startmonat'}
for koord, text in KOEPFE.items():
    assert data[koord].value is None, f'_data!{koord} belegt: {data[koord].value!r}'
    data[koord] = text
    data[koord].font = Font(name='Calibri', size=10, bold=True)

# Die Segmentliste G2:G27 war unvollstaendig: «Datacenter» und
# «Elektroplaner» kommen in der Pipeline vor, fehlten aber in der Liste —
# das alte Segment-Diagramm zeigte deshalb zu wenig. Beide werden ergaenzt
# (mit denselben Formeln wie die uebrigen Zeilen).
assert data['G28'].value is None and data['G29'].value is None
for r, seg in ((28, 'Datacenter'), (29, 'Elektroplaner')):
    data[f'G{r}'] = seg
    data[f'H{r}'] = f'=SUMIF({sp("D")},G{r},{sp("I")})'
    data[f'I{r}'] = f'=COUNTIF({sp("D")},G{r})'

# 4.1 Segment: aktives Volumen (J), Rangschluessel (Q), absteigend sortiert (R/S)
for r in range(2, 30):
    data[f'J{r}'] = (f'=SUMPRODUCT(--({sp("D")}=G{r}),{sp("AP")},{sp("AL")})')
    data[f'Q{r}'] = f'=J{r}+(50-ROW())*0.0001'
    k = 'ROW()-1'
    data[f'R{r}'] = (f'=IFERROR(IF(LARGE($Q$2:$Q$29,{k})<0.5,"",'
                     f'INDEX($G$2:$G$29,MATCH(LARGE($Q$2:$Q$29,{k}),$Q$2:$Q$29,0))),"")')
    data[f'S{r}'] = (f'=IF(R{r}="","",'
                     f'INDEX($J$2:$J$29,MATCH(LARGE($Q$2:$Q$29,{k}),$Q$2:$Q$29,0)))')
    data[f'S{r}'].number_format = '#,##0'

# Wachhund: taucht kuenftig ein Segment auf, das in der Liste fehlt,
# zeigt das Diagramm einen eigenen Restbalken statt still zu wenig.
data['R30'] = (f'=IF(ROUND(SUMPRODUCT({sp("AP")},{sp("AL")})-SUM($J$2:$J$29),2)=0,'
               f'"","Übrige (Segment fehlt in der Liste)")')
data['S30'] = (f'=IF(R30="","",SUMPRODUCT({sp("AP")},{sp("AL")})-SUM($J$2:$J$29))')
data['S30'].number_format = '#,##0'

# 4.2 Kanton: aktives Volumen (T), Rangschluessel (U), Top 10 (V/W)
for r in range(14, 40):
    data[f'T{r}'] = (f'=SUMPRODUCT(--({sp("C")}=A{r}),{sp("AP")},{sp("AL")})')
    data[f'U{r}'] = f'=T{r}+(45-ROW())*0.0001'
for r in range(2, 12):
    k = 'ROW()-1'
    data[f'V{r}'] = (f'=IFERROR(IF(LARGE($U$14:$U$39,{k})<0.5,"",'
                     f'INDEX($A$14:$A$39,MATCH(LARGE($U$14:$U$39,{k}),$U$14:$U$39,0))),"")')
    data[f'W{r}'] = (f'=IF(V{r}="","",'
                     f'INDEX($T$14:$T$39,MATCH(LARGE($U$14:$U$39,{k}),$U$14:$U$39,0)))')
    data[f'W{r}'].number_format = '#,##0'

# 4.4 Zerlegung der gewichteten Pipeline
data['X2'] = 'Gewonnene Aufträge (Faktor 100 %)'
data['Y2'] = f'=SUMPRODUCT({sp("AP")},--({sp("R")}="WON"),{sp("AJ")},{sp("AL")})'
data['X3'] = 'Offene Offerten (Faktoren 0–90 %)'
data['Y3'] = f'=SUMPRODUCT({sp("AP")},--({sp("R")}<>"WON"),{sp("AJ")},{sp("AL")})'
data['Y2'].number_format = '#,##0'
data['Y3'].number_format = '#,##0'

# 4.5 Startmonat (Bezugsjahr in AB1, wie im CEO Report: aktuelles Jahr)
data['AB1'] = '=YEAR(TODAY())'
MONATE = ['Jan', 'Feb', 'Mär', 'Apr', 'Mai', 'Jun',
          'Jul', 'Aug', 'Sep', 'Okt', 'Nov', 'Dez']
for i, m in enumerate(MONATE):
    r = 2 + i
    data[f'Z{r}'] = m
    data[f'AA{r}'] = (f'=SUMPRODUCT({sp("AP")},--({sp("AX")}=DATE($AB$1,{i + 1},1)),'
                      f'{sp("AL")})')
    data[f'AA{r}'].number_format = '#,##0'

# ============================================== Diagramm-Werkzeuge
def ref(col, r1, r2):
    return Reference(data, min_col=openpyxl.utils.column_index_from_string(col),
                     min_row=r1, max_row=r2)

def balken(cat_col, cat_r, val_col, val_r, richtung, titel=None,
           punktfarben=None, num='#,##0'):
    """Ein Balken-/Saeulendiagramm im einheitlichen Design."""
    ch = BarChart()
    ch.type = richtung                    # 'bar' = horizontal, 'col' = Saeulen
    ch.grouping = 'clustered'
    ch.gapWidth = 60
    ch.legend = None
    if titel:
        ch.title = titel
    ch.add_data(ref(val_col, val_r[0], val_r[1]), titles_from_data=False)
    ch.set_categories(ref(cat_col, cat_r[0], cat_r[1]))
    s = ch.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill=BLAU)
    if punktfarben:
        s.data_points = [DataPoint(idx=i, spPr=GraphicalProperties(solidFill=f)) for i, f in punktfarben]
    ch.dataLabels = DataLabelList(showVal=True, showLegendKey=False,
                                  showCatName=False, showSerName=False)
    ch.dataLabels.numFmt = num
    ch.dataLabels.dLblPos = 'outEnd'
    ch.y_axis.numFmt = num
    if richtung == 'bar':
        # Groesster Balken oben: Kategorienachse umkehren,
        # Wertachse bleibt unten (crosses am Maximum).
        ch.x_axis.scaling.orientation = 'maxMin'
        ch.y_axis.crosses = 'max'
    return ch

# ============================================== 📊 Diagramme neu bestuecken
alte = list(diag._charts)
assert len(alte) == 5, f'Erwartet 5 Diagramme, gefunden {len(alte)}'
anker = [ch.anchor for ch in alte]
diag._charts.clear()

TITEL_ZELLEN = {
    'A4': '🏭  Aktive Pipeline nach Segment — Volumen CHF, absteigend',
    'J4': '🗺  Aktive Pipeline nach Kanton — Top 10, Volumen CHF',
    'A50': '📊  Volumen nach Status — gewonnene Aufträge (WON) in Orange',
    'J50': '⚖️  Gewichtete Pipeline — zerlegt in WON und offene Offerten',
    'A95': '📅  Volumen nach Startmonat — füllt sich, sobald Projektzeiträume erfasst sind',
}
for koord, text in TITEL_ZELLEN.items():
    z = diag[koord]
    z.value = text
    z.font = Font(name='Calibri', size=13, bold=True, color='FF15243A')
    z.alignment = Alignment(horizontal='left', vertical='center')

c41 = balken('R', (2, 30), 'S', (2, 30), 'bar')
c42 = balken('V', (2, 11), 'W', (2, 11), 'bar')
c43 = balken('B', (2, 11), 'D', (2, 11), 'col',
             punktfarben=[(0, ORANGE)])
c44 = balken('X', (2, 3), 'Y', (2, 3), 'col',
             punktfarben=[(0, ORANGE)])
c45 = balken('Z', (2, 13), 'AA', (2, 13), 'col')

for ch, a in zip((c41, c42, c43, c44, c45), anker):
    ch.anchor = a
    diag.add_chart(ch)

# ============================================== Dashboard: gleiches Design
alte_dash = list(dash._charts)
assert len(alte_dash) == 2
dash._charts.clear()

# Ring: Status-Verteilung (Anzahl) — mehrere Farben, darum mit Legende.
ring = DoughnutChart()
ring.title = 'Pipeline — Status (Anzahl)'
ring.add_data(Reference(data, min_col=3, min_row=2, max_row=11),
              titles_from_data=False)
ring.set_categories(Reference(data, min_col=2, min_row=2, max_row=11))
ring.dataLabels = DataLabelList(showVal=True, showLegendKey=False,
                                showCatName=False, showSerName=False)
ring.dataLabels.numFmt = '0'
ring.series[0].data_points = [
    DataPoint(idx=i, spPr=GraphicalProperties(solidFill=PALETTE[i % len(PALETTE)])) for i in range(10)]
ring.legend.position = 'r'

# Balken: Status nach Volumen — WON in Orange (erste Kategorie).
vol = balken('B', (2, 11), 'D', (2, 11), 'bar',
             titel='Pipeline — Status nach Volumen CHF',
             punktfarben=[(0, ORANGE)])

for ch, alt in zip((ring, vol), alte_dash):
    ch.anchor = alt.anchor
    dash.add_chart(ch)

wb.save(DATEI)
print('Block 4 geschrieben: _data erweitert (Segment/Kanton/Zerlegung/Monate), '
      '5 Diagramme auf «📊 Diagramme» neu, 2 Dashboard-Diagramme im selben Design.')
