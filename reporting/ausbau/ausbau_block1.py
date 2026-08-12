#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BLOCK 1 — Nachweis und Klarheit.

Der CEO Report traegt in den Zeilen 201–251 bereits die Abschnitte
4 (gewichtete Bandtabelle), 5 (Qualitaet & Nachweis) und 6 (Startmonat).
Dieser Block ergaenzt chirurgisch, was der Auftrag zusaetzlich verlangt:

1.1  Im Abschnitt «5. QUALITÄT & NACHWEIS» fehlt zu den CHF-Werten die
     Anzahl. Neu: Spalte D «Anzahl» fuer WON gemeldet / davon belegt /
     noch zu belegen / brutto / bereinigt. Werte und Layout der uebrigen
     Zeilen bleiben unveraendert.
1.2  Neue Tabelle «7. VERTRAGSARTEN» (WON-Volumen nach Vertragsart);
     «Abrufbereitschaft» und «ohne Angabe» orange hervorgehoben.
     Die Fusszeile wandert dafuer von Zeile 251 ans neue Blattende.
1.3  Pipeline brutto und bereinigt stehen im Abschnitt 5 bereits
     untereinander (Zeilen 224/225). Zusaetzlich zeigen die Infozeilen
     A8 in CEO Report und Dashboard beide Werte nebeneinander.
1.4  Offene Pflichtangaben im Definitionsblatt: existiert live
     (Zeilen 15–23) — wird nur verifiziert, nicht angefasst.

Alles sind Formeln auf die Pipeline (Zeilen 6–860) — keine hart
getippten Summen. Kostenspalten J–N werden nirgends ausgewertet.
"""
import copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from ausbau_util import DATEI, merge, sp, F_ORANGE, CHF2

GRUEN = PatternFill('solid', fgColor='FF166534')   # Abschnittstitel im CEO Report
KOPF = PatternFill('solid', fgColor='FF1E273A')    # Tabellenkoepfe im CEO Report
ZEILE = PatternFill('solid', fgColor='FFF8FAFC')
TOTAL = PatternFill('solid', fgColor='FFDDE5EF')

wb = openpyxl.load_workbook(DATEI)
ceo = wb['CEO Report']
dash = wb['Dashboard']
d = wb['📋 Definitionen & Klärung']

# ---------------------------------------------------------------- 1.4 pruefen
ANKER_14 = {
    'B17': 'WON ohne Auftrags-/PO-Nummer und Datum',
    'B18': 'WON ohne vollständigen Einstand',
    'B19': 'Aktive Deals ohne erfassten Projektzeitraum',
    'B21': 'WON ohne Vertragsart (Einzelauftrag / Rahmenabruf / Abrufbereitschaft)',
    'B22': 'Kunden mehrfach in der Pipeline, Variante noch nicht geklärt',
}
for koord, text in ANKER_14.items():
    assert d[koord].value == text, f'1.4-Anker fehlt: {koord} = {d[koord].value!r}'
    f = d['C' + koord[1:]].value
    assert isinstance(f, str) and f.startswith('='), f'1.4: C{koord[1:]} ist keine Formel'
assert str(d['C23'].value).startswith('=SUM'), '1.4: Gesamturteil C23 ist keine Formel'

# ------------------------------------------- 1.1 Spalte D «Anzahl» ergaenzen
assert ceo['A218'].value.startswith('📋'), 'Abschnitt 5 nicht an Zeile 218'
assert ceo['A219'].value == 'Kennzahl' and ceo['E219'].value == 'Wert'

ANZAHL = {
    220: f'=COUNTIF({sp("R")},"WON")',                                  # WON gemeldet
    222: f'=SUMPRODUCT({sp("AT")})',                                    # davon belegt
    223: f'=COUNTIF({sp("R")},"WON")-SUMPRODUCT({sp("AT")})',           # noch zu belegen
    224: f'=SUMPRODUCT({sp("AP")})',                                    # Pipeline brutto
    225: f'=SUMPRODUCT({sp("AP")},{sp("AU")})',                         # bereinigt
}
for r in range(219, 230):
    bereich = f'A{r}:D{r}'
    if bereich in {str(m) for m in ceo.merged_cells.ranges}:
        ws_font = copy.copy(ceo[f'A{r}'].font)
        ceo.unmerge_cells(bereich)
        merge(ceo, f'A{r}:C{r}')
        ceo[f'A{r}'].font = ws_font

kz = ceo['E219']
ceo['D219'] = 'Anzahl'
ceo['D219'].font = copy.copy(kz.font)
ceo['D219'].fill = copy.copy(kz.fill)
ceo['D219'].alignment = Alignment(horizontal='center', vertical='center')
ceo['D219'].border = copy.copy(kz.border)

muster = ceo['E220']
for r in range(220, 230):
    z = ceo[f'D{r}']
    z.value = ANZAHL.get(r)
    z.font = Font(name='Arial', size=10, bold=(r in ANZAHL), color='FF1F3864')
    z.fill = copy.copy(muster.fill)
    z.border = copy.copy(muster.border)
    z.alignment = Alignment(horizontal='center', vertical='center')
    z.number_format = '0' if r in ANZAHL else 'General'

# Klarstellung im Label der Nachweiszeile (Definition von «belegt»).
ceo['F222'] = ('Beleg = Auftrags-/PO-Nummer UND Belegdatum UND vollständiger '
               'Einstand (Hilfsspalte AT). Nur dieser Teil ist berichtsfähig.')

# --------------------------------- Fusszeile ans Blattende verschieben
fuss = ceo['A251'].value
assert isinstance(fuss, str) and fuss.startswith('="Stand: ')
fuss_font = copy.copy(ceo['A251'].font)
fuss_align = copy.copy(ceo['A251'].alignment)
ceo.unmerge_cells('A251:L251')
ceo['A251'] = None
ceo.row_dimensions[251].height = 15

# ------------------------------------------------- 1.2 Vertragsarten-Tabelle
def titel(koord, text):
    z = ceo[koord]
    z.value = text
    z.font = Font(name='Arial', size=11, bold=True, color='FFFFFFFF')
    z.fill = GRUEN
    z.alignment = Alignment(horizontal='left', vertical='center')

def tkopf(koord, text, horizontal='center'):
    z = ceo[koord]
    z.value = text
    z.font = Font(name='Arial', size=9, bold=True, color='FFFFFFFF')
    z.fill = KOPF
    z.alignment = Alignment(horizontal=horizontal, vertical='center', wrap_text=True)

def wert(koord, inhalt, fmt='General', fett=False, farbe='FF111111',
         fuellung=ZEILE, horizontal='left', groesse=10, wrap=False):
    z = ceo[koord]
    z.value = inhalt
    z.font = Font(name='Arial', size=groesse, bold=fett, color=farbe)
    z.fill = fuellung
    z.number_format = fmt
    z.alignment = Alignment(horizontal=horizontal, vertical='center', wrap_text=wrap)

titel('A253', '📑  7. VERTRAGSARTEN — WON-Volumen nach Vertragsart')
merge(ceo, 'A253:N253')
ceo.row_dimensions[253].height = 24

tkopf('A254', 'Vertragsart', 'left'); merge(ceo, 'A254:C254')
tkopf('D254', 'Anzahl')
tkopf('E254', 'WON-Volumen CHF')
tkopf('F254', 'Hinweis', 'left'); merge(ceo, 'F254:L254')
ceo.row_dimensions[254].height = 19.5

VERTRAG = [
    ('Einzelauftrag', '="Einzelauftrag"', False,
     'Fester Auftrag, Bestellung oder PO liegt vor — zählt voll.'),
    ('Rahmenvertrag – Abruf bestätigt', '="Rahmenvertrag – Abruf bestätigt"', False,
     'Bestätigter, terminierter Abruf — zählt mit dem Betrag des Abrufs.'),
    ('Abrufbereitschaft', '="Abrufbereitschaft"', True,
     'Kein bestätigter Abruf — zählt nicht als gesicherter Umsatz, gehört in die Offert-Pipeline.'),
    ('Option / Reservation', '="Option / Reservation"', False,
     'Unverbindliche Reservation ohne Vertrag — zählt nicht.'),
    ('ohne Angabe', '=""', True,
     'Vertragsart in Spalte AB erfassen — bis dahin ist das Volumen nicht belastbar.'),
]
for i, (label, bed, orange, hinweis) in enumerate(VERTRAG):
    r = 255 + i
    fu = F_ORANGE if orange else ZEILE
    tf = 'FFFFFFFF' if orange else 'FF111111'
    hf = 'FFFFFFFF' if orange else 'FF555555'
    wert(f'A{r}', label, fett=orange, farbe=tf, fuellung=fu); merge(ceo, f'A{r}:C{r}')
    wert(f'D{r}', f'=SUMPRODUCT(--({sp("R")}="WON"),--({sp("AB")}{bed}))',
         fmt='0', fett=orange, farbe=tf, fuellung=fu, horizontal='center')
    wert(f'E{r}', f'=SUMPRODUCT(--({sp("R")}="WON"),--({sp("AB")}{bed}),{sp("AL")})',
         fmt=CHF2, fett=orange, farbe=tf, fuellung=fu, horizontal='right')
    wert(f'F{r}', hinweis, groesse=9, farbe=hf, fuellung=fu, wrap=True)
    merge(ceo, f'F{r}:L{r}')
    ceo.row_dimensions[r].height = 19.5

wert('A260', 'TOTAL WON', fett=True, fuellung=TOTAL); merge(ceo, 'A260:C260')
wert('D260', '=SUM(D255:D259)', fmt='0', fett=True, fuellung=TOTAL, horizontal='center')
wert('E260', '=SUM(E255:E259)', fmt=CHF2, fett=True, fuellung=TOTAL, horizontal='right')
wert('F260',
     '=IF(ROUND(E260-E220,2)=0,"✔  stimmt mit «WON gemeldet» in Abschnitt 5 überein",'
     '"⛔  Abweichung zu «WON gemeldet»: "&TEXT(E260-E220,"#,##0.00")&" CHF — Ursache klären")',
     groesse=9, fett=True, fuellung=TOTAL, wrap=True)
merge(ceo, 'F260:L260')
ceo.row_dimensions[260].height = 19.5
ceo.conditional_formatting.add(
    'F260', FormulaRule(formula=['ABS($E$260-$E$220)>0.005'], stopIfTrue=True,
                        font=Font(name='Arial', size=9, bold=True, color='FFFFFFFF'),
                        fill=PatternFill('solid', bgColor='FFC00000')))

# Fusszeile neu setzen (kommt nach dem letzten Abschnitt; Block 2 haengt
# darunter die Kontrollrechnung an und schiebt sie erneut).
ceo['A262'] = fuss
ceo['A262'].font = fuss_font
ceo['A262'].alignment = fuss_align
merge(ceo, 'A262:L262')
ceo.row_dimensions[262].height = 13.5

# ------------------------------------------------ 1.3 Infozeilen ergaenzen
ZUSATZ = ('&"   ·   Pipeline brutto: "&TEXT(SUMPRODUCT(' + sp('AP') + ',' + sp('AL') +
          '),"#,##0")&" CHF  ·  bereinigt: "&TEXT(SUMPRODUCT(' + sp('AP') + ',' +
          sp('AU') + ',' + sp('AL') + '),"#,##0")&" CHF"')
for blatt in (ceo, dash):
    a8 = blatt['A8'].value
    assert isinstance(a8, str) and a8.startswith('=') and 'brutto' not in a8
    blatt['A8'] = a8 + ZUSATZ

ceo.print_area = "'CEO Report'!$A$1:$N$262"

wb.save(DATEI)
print('Block 1 geschrieben: Anzahl-Spalte, Vertragsarten (Zeilen 253–260), '
      'Fusszeile verschoben, Infozeilen ergänzt, 1.4 verifiziert.')
