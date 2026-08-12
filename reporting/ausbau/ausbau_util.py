#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gemeinsame Bausteine fuer die Ausbau-Skripte (Block 1 bis 4).

Regel 1: Gewonnene Auftraege bleiben auf 100 Prozent — die Bandtabelle auf
         «⚖️ Wahrscheinlichkeit» (B17:D22) wird von keinem Skript beruehrt.
Regel 2: Keine Marge im Reporting — kein Skript liest oder schreibt die
         Kostenspalten J–N in einen Bericht.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DATEI = 'CH_MiT_Strom_Customer_CEO_CFO_MASTER__2_.xlsx'
PIPE = 'MiT Strom Pipeline'

# Einheitliches Design (Vorgabe Block 4, gilt auch fuer Hervorhebungen)
BLAU = 'FF15243A'      # Hauptfarbe dunkelblau
ORANGE = 'FFF47B20'    # Akzent
BLAU_HELL = 'FFEDF1F7'
GRAU_TEXT = 'FF555555'

F_BLAU = PatternFill('solid', fgColor=BLAU)
F_ORANGE = PatternFill('solid', fgColor=ORANGE)
F_KOPF = PatternFill('solid', fgColor='FF1E273A')
F_ZEILE = PatternFill('solid', fgColor='FFF8FAFC')
F_ZEILE2 = PatternFill('solid', fgColor=BLAU_HELL)
F_TOTAL = PatternFill('solid', fgColor='FFDDE5EF')

DUENN = Side(style='thin', color='FFB8C2CC')
RAHMEN = Border(left=DUENN, right=DUENN, top=DUENN, bottom=DUENN)

CHF = '#,##0'
CHF2 = '#,##0.00'
PROZ = '0.0%'


def kopf(ws, zelle, text, font='Arial', groesse=11):
    """Abschnittstitel: dunkelblau, weiss, fett."""
    z = ws[zelle]
    z.value = text
    z.font = Font(name=font, size=groesse, bold=True, color='FFFFFFFF')
    z.fill = F_BLAU
    z.alignment = Alignment(horizontal='left', vertical='center')
    return z


def tab_kopf(ws, zelle, text, font='Arial', horizontal='center'):
    """Tabellenkopf: dunkler Balken, weiss, fett, klein."""
    z = ws[zelle]
    z.value = text
    z.font = Font(name=font, size=9, bold=True, color='FFFFFFFF')
    z.fill = F_KOPF
    z.alignment = Alignment(horizontal=horizontal, vertical='center', wrap_text=True)
    z.border = RAHMEN
    return z


def zelle(ws, koord, wert, font='Arial', groesse=10, fett=False, farbe='FF111111',
          fuellung=None, fmt=None, horizontal='left', wrap=False):
    z = ws[koord]
    z.value = wert
    z.font = Font(name=font, size=groesse, bold=fett, color=farbe)
    if fuellung is not None:
        z.fill = fuellung
    if fmt:
        z.number_format = fmt
    z.alignment = Alignment(horizontal=horizontal, vertical='center', wrap_text=wrap)
    z.border = RAHMEN
    return z


def merge(ws, bereich):
    """Verbund setzen — nur wenn er nicht schon existiert."""
    if bereich not in {str(m) for m in ws.merged_cells.ranges}:
        ws.merge_cells(bereich)


def unmerge_bereich(ws, r_von, r_bis):
    """Alle Verbuende aufloesen, die ganz in den Zeilen r_von..r_bis liegen."""
    weg = [str(m) for m in ws.merged_cells.ranges
           if m.min_row >= r_von and m.max_row <= r_bis]
    for m in weg:
        ws.unmerge_cells(m)
    return weg


def raum_leeren(ws, r_von, r_bis, c_von=1, c_bis=14):
    """Inhalt und Stil eines Zeilenbands zuruecksetzen (nach unmerge)."""
    leer_font = Font(name='Arial', size=10)
    for r in range(r_von, r_bis + 1):
        for c in range(c_von, c_bis + 1):
            z = ws.cell(row=r, column=c)
            z.value = None
            z.font = leer_font
            z.fill = PatternFill()
            z.border = Border()
            z.number_format = 'General'
            z.alignment = Alignment()


P = "'MiT Strom Pipeline'!"


def sp(col, fest=True):
    """Pipelinebereich einer Spalte, Zeile 6 bis 860, absolut."""
    d = '$' if fest else ''
    return f"{P}{d}{col}{d}6:{d}{col}{d}860"
