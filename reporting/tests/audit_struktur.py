#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""E) Strukturaudit: Verknuepfungen, Namen, Dropdowns, Formate, Diagramme,
Druckbereiche, verbundene Zellen, Schriften."""
import openpyxl, sys, re
from openpyxl.utils import range_boundaries, get_column_letter

F = 'CH_MiT_Strom_Customer_CEO_CFO_MASTER.xlsx'
wb = openpyxl.load_workbook(F)
V = openpyxl.load_workbook(F, data_only=True)
bad = []


def note(kind, msg):
    bad.append(f'[{kind}] {msg}')


# 1) Benannter Bereich zeigt auf die Skala
dn = wb.defined_names.get('Wahrscheinlichkeit_Skala')
if dn is None:
    note('NAME', 'Wahrscheinlichkeit_Skala fehlt')
else:
    sheet, ref = dn.attr_text.split('!')
    sheet = sheet.strip("'")
    if sheet not in wb.sheetnames:
        note('NAME', f'zeigt auf unbekanntes Blatt {sheet}')
    else:
        vals = [V[sheet][f'B{r}'].value for r in range(6, 11)]
        if vals != [0, 0.1, 0.3, 0.6, 0.9]:
            note('NAME', f'Skalenwerte unerwartet: {vals}')

# 2) Datenpruefung
dvs = {str(d.sqref): d for d in wb['MiT Strom Pipeline'].data_validations.dataValidation}
if 'S6:S860' not in dvs:
    note('DROPDOWN', 'Dropdown auf S6:S860 fehlt')
else:
    d = dvs['S6:S860']
    if d.formula1 != 'Wahrscheinlichkeit_Skala':
        note('DROPDOWN', f'Quelle ist {d.formula1}')
    if d.errorStyle != 'warning':
        note('DROPDOWN', f'errorStyle={d.errorStyle} (Altwerte wuerden blockiert)')
for need in ('R6:R860', 'C6:C860', 'T6:T860'):
    if need not in dvs:
        note('DROPDOWN', f'bestehendes Dropdown {need} verloren')

# 3) Bedingte Formatierung
cf = [str(k.sqref) for k in wb['MiT Strom Pipeline'].conditional_formatting._cf_rules]
for need in ('I6:I860', 'P6:P860', 'R6:R860', 'S6:S860'):
    if need not in cf:
        note('FORMAT', f'bedingte Formatierung {need} fehlt')

# 4) Diagramme + Quellbezuege
dia = wb['📊 Diagramme']
if len(dia._charts) != 5:
    note('DIAGRAMM', f'{len(dia._charts)} Diagramme statt 5')
for ch in dia._charts:
    for s in ch.series:
        f = s.val.numRef.f if (s.val and s.val.numRef) else None
        if not f:
            note('DIAGRAMM', 'Datenreihe ohne Quellbezug')
            continue
        sh = f.split('!')[0].strip("'")
        if sh not in wb.sheetnames:
            note('DIAGRAMM', f'Quelle zeigt auf unbekanntes Blatt: {f}')
        rng = f.split('!')[1]
        if any(V[sh][c[0].coordinate].value is None
               for c in V[sh][rng.replace('$', '')]):
            pass  # leere Zellen sind erlaubt

# 5) Druckbereiche muessen alle sichtbaren Inhalte abdecken
for ws in wb.worksheets:
    if ws.title == '_data' or not ws.print_area:
        continue
    pa = ws.print_area[0] if isinstance(ws.print_area, list) else ws.print_area
    mc, mr, xc, xr = range_boundaries(pa.split('!')[1])
    # letzte Zeile/Spalte mit sichtbarem Inhalt (Hilfsspalten ausgenommen)
    lastr = lastc = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value in (None, ''):
                continue
            col = get_column_letter(c.column)
            if ws.title == 'MiT Strom Pipeline' and c.column > 35:   # ab AJ = Hilfsspalten
                continue
            if ws.title in ('Dashboard', 'CEO Report') and col == 'U':
                continue
            if ws.title == '📄 Report' and col in ('J', 'K', 'L'):
                continue
            if ws.column_dimensions[col].hidden:      # ausgeblendet -> druckt nie
                continue
            # Unterhalb des Druckbereichs liegt nur vorformatierter Formelvorrat
            # (liefert ""). Deals dort loesen die Kontrolle auf dem Blatt
            # "Wahrscheinlichkeit" aus - siehe Verhaltenstest unter_druckbereich.
            if ws.title == 'MiT Strom Pipeline' and c.row > 90:
                continue
            lastr = max(lastr, c.row); lastc = max(lastc, c.column)
    if lastr > xr:
        note('DRUCK', f'{ws.title}: Inhalt bis Zeile {lastr}, Druckbereich nur bis {xr}')
    if lastc > xc:
        note('DRUCK', f'{ws.title}: Inhalt bis Spalte {get_column_letter(lastc)}, '
                      f'Druckbereich nur bis {get_column_letter(xc)}')

# 6) Verbundene Zellen duerfen keinen Inhalt verstecken
for ws in wb.worksheets:
    for m in ws.merged_cells.ranges:
        mc, mr, xc, xr = m.bounds
        for r in range(mr, xr + 1):
            for c in range(mc, xc + 1):
                if (r, c) == (mr, mc):
                    continue
                cell = V[ws.title].cell(row=r, column=c)
                if cell.value not in (None, ''):
                    note('VERBUND', f'{ws.title}!{cell.coordinate} verdeckt durch {m}: {cell.value!r}')

# 7) Schriftarten
fonts = set()
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font and c.font.name:
                fonts.add(c.font.name)
if not fonts <= {'Arial', 'Calibri', 'Cambria', 'Times New Roman', 'Consolas'}:
    note('SCHRIFT', f'unerwartete Schriftarten: {fonts}')

# 8) Zahlenformate der neuen Bereiche
# Zeilen der Skala und der Bandtabelle werden gesucht, damit die Pruefung
# nicht bricht, sobald eine Stufe oder ein Band dazukommt.
_WV = V['⚖️ Wahrscheinlichkeit']
_SK_VON = 6
_SK_BIS = next(r for r in range(_SK_VON, 30)
               if not isinstance(_WV[f'B{r}'].value, (int, float))) - 1
_BD_VON = next(r for r in range(_SK_BIS + 1, 40) if _WV[f'A{r}'].value == '0 – 44 %')
_BD_TOT = next(r for r in range(_BD_VON, 40) if _WV[f'A{r}'].value == 'TOTAL aktiv')
FMT = [('⚖️ Wahrscheinlichkeit', [f'B{_SK_VON}', f'B{_SK_BIS}', f'D{_BD_VON}', f'D{_BD_TOT-1}'], '0%'),
       ('⚖️ Wahrscheinlichkeit', [f'F{_BD_VON}', f'G{_BD_VON}', f'F{_BD_TOT}', f'G{_BD_TOT}'], '#,##0'),
       ('⚖️ Wahrscheinlichkeit', [f'H{_BD_VON}', f'H{_BD_TOT}'], '0.0%'),
       ('Dashboard', ['F208', 'G208', 'F212', 'G212'], '#,##0'),
       ('Dashboard', ['D208', 'I208'], None),
       ('CEO Report', ['F208', 'G208'], '#,##0')]
for sheet, cells, want in FMT:
    for co in cells:
        got = wb[sheet][co].number_format
        if want and got != want:
            note('FORMAT', f'{sheet}!{co}: Zahlenformat {got!r} statt {want!r}')

# 9) Prozentwerte muessen als Bruch gespeichert sein.
#    Skala und Bandtabelle werden gesucht statt fest verdrahtet - die Zahl der
#    Stufen und Baender kann sich aendern.
_W = V['⚖️ Wahrscheinlichkeit']
_skala_von = 6
_skala_bis = next(r for r in range(_skala_von, 30)
                  if not isinstance(_W[f'B{r}'].value, (int, float))) - 1
_band_von = next(r for r in range(_skala_bis + 1, 40) if _W[f'A{r}'].value == '0 – 44 %')
_band_bis = next(r for r in range(_band_von, 40) if _W[f'A{r}'].value == 'TOTAL aktiv') - 1
for r in range(_skala_von, _skala_bis + 1):
    v = _W[f'B{r}'].value
    if not (0 <= v <= 1):
        note('WERT', f'Skalenwert B{r}={v} ist kein Bruch')
for r in range(_band_von, _band_bis + 1):
    for col in ('B', 'C', 'D'):
        v = _W[f'{col}{r}'].value
        if not isinstance(v, (int, float)) or not (0 <= v <= 1):
            note('WERT', f'{col}{r}={v} ist kein Bruch')

# 10) Blattreihenfolge / Register
if wb.sheetnames != ['📑 Executive PDF', '📄 Report', 'CEO Report', 'Dashboard',
                     '📊 Diagramme', 'MiT Strom Pipeline', '📋 Definitionen & Klärung',
                     '🔍 Herleitung & Formeln', '⚖️ Wahrscheinlichkeit', '_data']:
    note('BLATT', f'Reihenfolge: {wb.sheetnames}')
if wb['_data'].sheet_state != 'hidden':
    note('BLATT', '_data ist nicht mehr ausgeblendet')

# 11) Keine Formel darf auf ein leeres Blatt / falsche Spalte zeigen
#     Stichprobe: alle Verweise auf die Pipeline muessen in A..BC liegen
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not (isinstance(v, str) and v.startswith('=')):
                continue
            for col in re.findall(r"'MiT Strom Pipeline'!\$?([A-Z]{1,2})\$?\d", v):
                idx = openpyxl.utils.column_index_from_string(col)
                if idx > 55:                       # bis BC reichen die Hilfsspalten
                    note('BEZUG', f'{ws.title}!{c.coordinate} zeigt auf Spalte {col} (> BC)')

# 12) Neue Fachspalten, Dropdowns und benannte Bereiche
P_ = wb['MiT Strom Pipeline']
for co, want in (('Y5', 'Nettoumsatz'), ('Z5', 'Einstand'), ('AA5', 'Abwicklung'),
                 ('AB5', 'Vertragsart'), ('AC5', 'Deal-Gruppe'), ('AD5', 'Variante'),
                 ('AE5', 'Offert-Nr.'), ('AF5', 'Auftrag / PO-Nr.'),
                 ('AG5', 'Beleg-Datum'), ('AH5', 'Prüfstatus'), ('AI5', 'Wahr. %'),
                 ('N5', 'Übrige Kosten')):
    if want not in str(P_[co].value):
        note('SPALTE', f'{co}: erwartet «{want}», gefunden {P_[co].value!r}')
dvs2 = {str(x.sqref): x for x in P_.data_validations.dataValidation}
for need in ('AA6:AA860', 'AB6:AB860', 'AD6:AD860', 'AG6:AG860'):
    if need not in dvs2:
        note('DROPDOWN', f'Auswahlliste {need} fehlt')
for nm in ('MwSt_Satz', 'Netto_Faktor', 'Schwelle_Aufteilung', 'Wahrscheinlichkeit_Skala'):
    if wb.defined_names.get(nm) is None:
        note('NAME', f'benannter Bereich {nm} fehlt')
# MwSt darf in Spalte N nicht mehr als Kosten stehen
for r in range(6, 121):
    v = P_[f'N{r}'].value
    if isinstance(v, str) and '1.081' in v:
        note('MWST', f'N{r} enthält weiterhin eine MwSt-Formel')
cf2 = [str(k.sqref) for k in P_.conditional_formatting._cf_rules]
for need in ('AH6:AH860', 'O6:P860', 'AI6:AI860'):
    if need not in cf2:
        note('FORMAT', f'bedingte Formatierung {need} fehlt')

# 13) Erklaerungsblatt: kein Beschriftungstext darf als Formel gelesen werden
E_ = wb['🔍 Herleitung & Formeln']
for r_ in E_.iter_rows():
    for c_ in r_:
        v_ = c_.value
        if isinstance(v_, str) and v_.startswith('='):
            if re.search(r'[()!$]', v_) or v_[1:].strip() in wb.defined_names:
                continue
            note('TEXT-ALS-FORMEL', f'Herleitung!{c_.coordinate}: {v_[:40]}')

# 14) Auslieferungsstand: jede erfasste Wahrscheinlichkeit liegt auf der Skala
# 100 % gehoert zur Skala, ist aber gewonnenen Auftraegen vorbehalten.
SKALA_SET = {0, 0.1, 0.3, 0.6, 0.9, 1.0}
Vp = V['MiT Strom Pipeline']
ab = 0
for r in range(6, 861):
    if Vp[f'B{r}'].value in (None, ''):
        continue
    sv = Vp[f'S{r}'].value
    if isinstance(sv, (int, float)) and round(sv, 6) not in SKALA_SET:
        note('SKALA', f'S{r} = {sv} liegt nicht auf 0/10/30/60/90/100 %')
        ab += 1
    # Der bisherige Wert muss erhalten geblieben sein, wo umgeschluesselt wurde
    alt = Vp[f'AI{r}'].value
    if isinstance(alt, (int, float)) and not isinstance(sv, (int, float)):
        note('UMSCHLUESSELUNG', f'AI{r} hat einen Altwert, S{r} ist aber leer')
_AUSSER = next(r for r in range(_BD_TOT, 60)
               if isinstance(_WV[f'A{r}'].value, str)
               and _WV[f'A{r}'].value.startswith('Aktive Deals mit Wahrscheinlichkeit'))
if ab == 0 and _WV[f'E{_AUSSER}'].value not in (0, None):
    note('SKALA', 'Zaehler der Deals ausserhalb der Skala steht nicht auf 0')

print('=' * 70)
if bad:
    print(f'{len(bad)} Befunde:')
    for b in bad:
        print('  ', b)
else:
    print('Struktur OK — Namen, Dropdowns, Formate, Diagramme, Druckbereiche,')
    print('verbundene Zellen, Schriften und Bezüge sind alle in Ordnung.')
sys.exit(1 if bad else 0)
