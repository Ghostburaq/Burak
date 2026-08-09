#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
F) Abdeckungspruefung: ist jede Rueckfrage von Maria und Oliver im File
   beantwortet - mit Antworttext UND einer Live-Zahl, die wirklich rechnet?
"""
import openpyxl, sys

F = sys.argv[1] if len(sys.argv) > 1 else 'CH_MiT_Strom_Customer_CEO_CFO_MASTER.xlsx'
wbf = openpyxl.load_workbook(F)
wbv = openpyxl.load_workbook(F, data_only=True)
DEF = '📋 Definitionen & Klärung'
bad = []

# Jede Frage: (Wer, Stichwort in der Frage, Begriffe die in der Antwort stehen muessen)
FRAGEN = [
    ('Maria', 'Margenplanung', ['MwSt', 'Nettoumsatz', 'Einstand']),
    ('Maria', 'Verantwortung', ['Innendienst', 'Einstand']),
    ('Maria', 'MiT CH', ['Abwicklung', 'Standard']),
    ('Maria', 'WON-Volumen', ['netto', 'belegt']),
    ('Maria', 'Offerten erstellt', ['Offert-Nr', 'Beleg-Datum', 'Prüfstatus']),
    ('Maria', 'Datacenter', ['Deal-Gruppe', 'Alternative']),
    ('Maria', 'Abrufbereitschaft', ['Abruf', 'Offert-Pipeline']),
    ('Oliver', 'ausserhalb der Aggreko-Skala', ['0/10/30/60/90', 'umgeschlüsselt', '45–59']),
]

df, dv = wbf[DEF], wbv[DEF]
# Frageblock einlesen: Spalte B = Frage, C = Antwort, G = Live-Zahl
block = []
for r in range(1, df.max_row + 1):
    frage = df[f'B{r}'].value
    if not isinstance(frage, str) or not frage.strip():
        continue
    antwort = df[f'C{r}'].value
    live_f = df[f'G{r}'].value
    live_v = dv[f'G{r}'].value
    if isinstance(antwort, str) and isinstance(live_f, str) and live_f.startswith('='):
        block.append((r, frage, antwort, live_v))

print(f'Gefundene Frage-Antwort-Zeilen mit Live-Zahl: {len(block)}')
for wer, stichwort, begriffe in FRAGEN:
    # Nur der Fragetext zaehlt - sonst trifft ein Stichwort zufaellig in einer
    # anderen Antwort und die Luecke bliebe unentdeckt.
    treffer = [b for b in block if stichwort.lower() in b[1].lower()]
    if not treffer:
        bad.append(f'[{wer}] Frage «{stichwort}» hat keine Antwortzeile mit Live-Zahl')
        continue
    r, frage, antwort, live = treffer[0]
    fehlend = [t for t in begriffe if t.lower() not in antwort.lower()]
    if fehlend:
        bad.append(f'[{wer}] Antwort zu «{stichwort}» (Zeile {r}) nennt nicht: {fehlend}')
    if live in (None, '', 0):
        bad.append(f'[{wer}] Live-Zahl zu «{stichwort}» (Zeile {r}) ist leer: {live!r}')
    else:
        print(f'  ✔ [{wer:<6}] {stichwort:<18} Zeile {r:>3}  →  {str(live)[:88]}')

# Die offenen Punkte muessen eine Gesamtaussage tragen
gesamt = [r for r in range(1, df.max_row + 1) if df[f'B{r}'].value == 'Gesamturteil']
if not gesamt:
    bad.append('Gesamturteil zu den offenen Punkten fehlt')
else:
    r = gesamt[0]
    if dv[f'D{r}'].value in (None, ''):
        bad.append(f'Gesamturteil in Zeile {r} ist leer')
    else:
        print(f'  ✔ [beide ] Gesamturteil       Zeile {r:>3}  →  {str(dv[f"D{r}"].value)[:88]}')

print()
if bad:
    print(f'{len(bad)} Lücken:')
    for b in bad:
        print('  ', b)
else:
    print('Alle Rückfragen von Maria und Oliver sind im File beantwortet — je mit')
    print('Antworttext und einer Live-Zahl, die aus der Pipeline rechnet.')
sys.exit(1 if bad else 0)
