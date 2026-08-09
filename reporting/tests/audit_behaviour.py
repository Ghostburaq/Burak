#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
D) Verhaltenstests: Rohdaten veraendern, neu berechnen und pruefen, ob die
ganze Mappe korrekt mitzieht. Jedes Szenario wird zusaetzlich mit dem
unabhaengigen Modell (audit_values.py) gegengerechnet.
"""
import openpyxl, subprocess, shutil, sys, json, os

MASTER = 'CH_MiT_Strom_Customer_CEO_CFO_MASTER.xlsx'
RECALC = '/root/.claude/skills/xlsx/scripts/recalc.py'
results = []


def build(name, mutate):
    f = f'test_{name}.xlsx'
    shutil.copy(MASTER, f)
    wb = openpyxl.load_workbook(f)          # Formeln behalten
    mutate(wb)
    wb.save(f)
    out = subprocess.run([sys.executable, RECALC, f, '300'],
                         capture_output=True, text=True)
    try:
        j = json.loads(out.stdout)
    except Exception:
        return f, {'status': 'PARSE-FEHLER', 'raw': out.stdout[-300:]}
    return f, j


def verify(f, name, extra=None):
    """Fehlerscan + unabhaengige Nachrechnung + Zusatzerwartungen."""
    msgs = []
    V = openpyxl.load_workbook(f, data_only=True)
    ERRS = ('#REF!', '#NAME?', '#DIV/0!', '#VALUE!', '#N/A', '#NULL!', '#NUM!', 'Err:')
    bad = [f'{ws.title}!{c.coordinate}={c.value}'
           for ws in V.worksheets for row in ws.iter_rows() for c in row
           if isinstance(c.value, str) and any(e in c.value for e in ERRS)]
    if bad:
        msgs.append(f'FEHLERWERTE ({len(bad)}): ' + ', '.join(bad[:5]))

    r = subprocess.run([sys.executable, 'audit_values.py', f],
                       capture_output=True, text=True)
    tail = [l for l in r.stdout.splitlines() if l.strip()]
    model = tail[0] if tail else '?'
    devi = tail[1] if len(tail) > 1 else '?'
    if r.returncode != 0:
        msgs.append('MODELL-ABWEICHUNG:\n      ' + '\n      '.join(tail[1:12]))

    if extra:
        msgs += extra(V)
    results.append((name, model, devi, msgs))
    return msgs


P = 'MiT Strom Pipeline'


# ---------------------------------------------------------------- Szenarien
def s1(wb):
    """Neuer Deal in der ersten freien Zeile (72)."""
    ws = wb[P]
    vals = dict(A=999, B='Testkunde Neu AG', C='ZH', D='Industrie', E='2x 500 kVA',
                F=1000, G=90, H='Sep', I=250000, J=100000, K=5000, L=2000,
                M=3000, N=0, R='offered', S=0.6, T='Warm', U='Test',
                V='Nächster Schritt Test', W='Sep', X='Testzeile')
    for k, v in vals.items():
        ws[f'{k}72'] = v


def s1x(V):
    m = []
    p = V[P]
    if p['AJ72'].value != 0.5:
        m.append(f'Faktor 60% -> erwartet 0.5, ist {p["AJ72"].value}')
    if abs((p['Q72'].value or 0) - 125000) > 0.01:
        m.append(f'Gew.Wert erwartet 125000, ist {p["Q72"].value}')
    # Deal muss im Offerte-Block des Dashboards auftauchen
    names = [V['Dashboard'][f'B{r}'].value for r in range(75, 137)]
    if 'Testkunde Neu AG' not in names:
        m.append('neuer Deal fehlt im Offerte-Block des Dashboards')
    if V['⚖️ Wahrscheinlichkeit']['E18'].value != 3:
        m.append(f'Band 60-89% erwartet 3 Deals, ist {V["⚖️ Wahrscheinlichkeit"]["E18"].value}')
    return m


def s2(wb):
    """Grosster WON-Deal wird zu LOST."""
    wb[P]['R9'] = 'LOST'


def s2x(V):
    m = []
    if V['Dashboard']['A7'].value != 11:
        m.append(f'WON-Anzahl erwartet 11, ist {V["Dashboard"]["A7"].value}')
    if V[P]['AK9'].value not in (None, ''):
        m.append('AK9 muss leer sein, wenn nicht mehr WON')
    top = V['📑 Executive PDF']['C14'].value
    if top == 'Ice Hockey Championship':
        m.append('Top-1 WON haette wechseln muessen')
    return m


def s3(wb):
    """Wahrscheinlichkeit leeren (aktiver Deal)."""
    wb[P]['S6'] = None


def s3x(V):
    m = []
    if V[P]['Q6'].value != 'tbd':
        m.append(f'Q6 erwartet "tbd", ist {V[P]["Q6"].value!r}')
    if V[P]['AJ6'].value != 0:
        m.append(f'AJ6 erwartet 0, ist {V[P]["AJ6"].value}')
    if V[P]['AQ6'].value != 0:
        m.append('AQ6 muss 0 sein (nicht auf Skala)')
    return m


def s4(wb):
    """Wahrscheinlichkeit als 60 statt 0.6 erfasst (Salesforce-Import)."""
    wb[P]['S6'] = 60


def s4x(V):
    m = []
    if V[P]['AJ6'].value != 0.5:
        m.append(f'60 (statt 0.6) -> Faktor erwartet 0.5, ist {V[P]["AJ6"].value}')
    if abs((V[P]['Q6'].value or 0) - 195000) > 0.01:
        m.append(f'Q6 erwartet 195000, ist {V[P]["Q6"].value}')
    return m


BOUND = [(0.0, 0.0), (0.10, 0.0), (0.30, 0.0), (0.44, 0.0), (0.4499, 0.0),
         (0.45, 0.3), (0.50, 0.3), (0.59, 0.3), (0.5999, 0.3),
         (0.60, 0.5), (0.75, 0.5), (0.89, 0.5), (0.8999, 0.5),
         (0.90, 0.9), (0.95, 0.9), (1.00, 0.9)]


def s5(wb):
    """Alle Bandgrenzen in freien Zeilen durchtesten."""
    ws = wb[P]
    for i, (prob, _) in enumerate(BOUND):
        r = 100 + i
        ws[f'B{r}'] = f'Grenzwert {prob}'
        ws[f'C{r}'] = 'ZH'
        ws[f'I{r}'] = 100000
        ws[f'R{r}'] = 'offered'
        ws[f'S{r}'] = prob


def s5x(V):
    # Hinweis: dieses Szenario setzt bewusst Werte ausserhalb der Skala, um die
    # Bandgrenzen zu pruefen. Die Skalenkontrolle der Mappe muss das melden.
    m = []
    for i, (prob, want) in enumerate(BOUND):
        r = 100 + i
        got = V[P][f'AJ{r}'].value
        if abs((got or 0) - want) > 1e-9:
            m.append(f'S={prob} -> Faktor erwartet {want}, ist {got}')
        q = V[P][f'Q{r}'].value
        if abs((q or 0) - 100000 * want) > 0.01:
            m.append(f'S={prob} -> Gew.Wert erwartet {100000*want}, ist {q}')
    return m


def s6(wb):
    """Alle Deals loeschen - darf keine Division durch 0 o.ae. ausloesen."""
    ws = wb[P]
    for r in range(6, 72):
        for col in ('A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K',
                    'L', 'M', 'R', 'S', 'T', 'U', 'V', 'W', 'X'):
            ws[f'{col}{r}'] = None
        ws[f'N{r}'] = None


def s6x(V):
    m = []
    if V['Dashboard']['I6'].value not in (0, None):
        m.append(f'Gesamtpipeline erwartet 0, ist {V["Dashboard"]["I6"].value}')
    if V['⚖️ Wahrscheinlichkeit']['G20'].value not in (0, None):
        m.append('gewichtetes Total muss 0 sein')
    if V['⚖️ Wahrscheinlichkeit']['E28'].value not in (0, None):
        m.append('Skala-Zaehler muss 0 sein')
    if '✔' not in str(V['⚖️ Wahrscheinlichkeit']['E27'].value):
        m.append('Selbstkontrolle muss auch bei leerer Pipeline gruen sein')
    if V['Dashboard']['B11'].value not in (None, ''):
        m.append('WON-Block muss leer sein')
    return m


def s7(wb):
    """Volumen als Text ("tbd") bei aktivem Deal."""
    wb[P]['I6'] = 'tbd'


def s7x(V):
    m = []
    if V[P]['Q6'].value != 'tbd':
        m.append(f'Q6 erwartet "tbd", ist {V[P]["Q6"].value!r}')
    if V[P]['O6'].value not in (None, ''):
        m.append(f'Marge O6 muss leer sein, ist {V[P]["O6"].value!r}')
    if V[P]['AL6'].value != 0:
        m.append('AL6 muss 0 sein bei Text-Volumen')
    return m




def s8(wb):
    """Deal mit Namen aber ohne Status."""
    ws = wb[P]
    ws['B72'] = 'Ohne Status AG'
    ws['I72'] = 50000
    ws['S72'] = 0.6


def s8x(V):
    m = []
    p = V[P]
    if p['AP72'].value != 0:
        m.append('Deal ohne Status darf nicht als aktiv zaehlen')
    if p['Q72'].value != 25000:          # 50'000 x Faktor 50 % (S=60 %)
        m.append(f'Q72 erwartet 25000, ist {p["Q72"].value}')
    return m


def s9(wb):
    """Zwei WON-Deals mit exakt gleichem Volumen (Doppel-Risiko Top-Liste)."""
    ws = wb[P]
    ws['I13'] = 380000          # gleich wie Ice Hockey (Zeile 9)


def s9x(V):
    m = []
    namen = [V['📄 Report'][f'B{r}'].value for r in range(29, 39)]
    doppelt = [n for n in set(namen) if n and namen.count(n) > 1]
    if doppelt:
        m.append(f'Deal doppelt in der Top-WON-Liste: {doppelt}')
    ex = [V['📑 Executive PDF'][f'C{r}'].value for r in range(14, 19)]
    dop2 = [n for n in set(ex) if n and ex.count(n) > 1]
    if dop2:
        m.append(f'Deal doppelt in Top-5 Executive PDF: {dop2}')
    return m


def s10(wb):
    """Deal in der allerletzten Zeile des Formelbereichs (860)."""
    ws = wb[P]
    ws['B860'] = 'Letzte Zeile AG'
    ws['C860'] = 'BE'
    ws['D860'] = 'Industrie'
    ws['I860'] = 90000
    ws['R860'] = 'WON'
    ws['S860'] = 0.9


def s10x(V):
    m = []
    if V['Dashboard']['A7'].value != 13:
        m.append(f'WON-Anzahl erwartet 13, ist {V["Dashboard"]["A7"].value}')
    if abs((V[P]['Q860'].value or 0) - 81000) > 0.01:
        m.append(f'Q860 erwartet 81000, ist {V[P]["Q860"].value}')
    namen = [V['Dashboard'][f'B{r}'].value for r in range(11, 73)]
    if 'Letzte Zeile AG' not in namen:
        m.append('Deal in Zeile 860 erscheint nicht im WON-Block')
    return m



def s11(wb):
    """Deal unterhalb des Druckbereichs (Zeile 91) - muss Alarm ausloesen."""
    ws = wb[P]
    ws['B91'] = 'Ausserhalb Druckbereich AG'
    ws['C91'] = 'ZH'
    ws['I91'] = 40000
    ws['R91'] = 'offered'
    ws['S91'] = 0.3


def s11x(V):
    m = []
    W = V['⚖️ Wahrscheinlichkeit']
    if W['E30'].value != 1:
        m.append(f'Zaehler erwartet 1, ist {W["E30"].value}')
    if '⚠' not in str(W['E31'].value):
        m.append(f'Warnung fehlt: {W["E31"].value!r}')
    if V['Dashboard']['I7'].value != 47:
        m.append('Deal zaehlt trotzdem korrekt zur Pipeline (I7 erwartet 47)')
    return m


def s12(wb):
    """Verlorener Deal behaelt 90 % - U4 darf ihn nicht mitzaehlen."""
    ws = wb[P]
    ws['R6'] = 'LOST'
    ws['S6'] = 0.9


def s12x(V):
    m = []
    u4 = V[P]['U4'].value
    g20 = V['⚖️ Wahrscheinlichkeit']['G20'].value
    if abs((u4 or 0) - (g20 or 0)) > 0.01:
        m.append(f'U4={u4} weicht von der gewichteten Pipeline {g20} ab')
    if abs((V[P]['Q6'].value or 0) - 351000) > 0.01:
        m.append(f'Q6 des LOST-Deals erwartet 351000, ist {V[P]["Q6"].value}')
    return m


def s13(wb):
    """WON-Deal vollstaendig belegen: PO-Nr., Datum, Einstand, Vertragsart."""
    ws = wb[P]
    ws['AF9'] = 'PO-2026-0815'
    ws['AG9'] = __import__('datetime').date(2026, 5, 12)
    ws['AB9'] = 'Einzelauftrag'


def s13x(V):
    m = []
    p = V[P]
    if p['AV9'].value != 1:
        m.append('Beleg-Kennzeichen AV9 muesste 1 sein')
    if p['AW9'].value != 1:
        m.append('WON-belegt AW9 muesste 1 sein')
    if p['AH9'].value != '✅ belegt & kalkuliert':
        m.append(f'Pruefstatus erwartet belegt, ist {p["AH9"].value!r}')
    if abs((V['Dashboard']['E221'].value or 0) - 380000) > 1:
        m.append(f'belegtes WON erwartet 380000, ist {V["Dashboard"]["E221"].value}')
    return m


def s14(wb):
    """Variante ausschliessen: bereinigtes Volumen muss sinken."""
    ws = wb[P]
    ws['AC62'] = 'DPR-Loadbank'
    ws['AC63'] = 'DPR-Loadbank'
    ws['AD62'] = 'Führend'
    ws['AD63'] = 'Alternative – zählt nicht'


def s14x(V):
    m = []
    brutto = V['Dashboard']['E223'].value or 0
    berein = V['Dashboard']['E224'].value or 0
    if abs((brutto - berein) - 331889.15) > 1:
        m.append(f'Bereinigung erwartet -331889, ist {brutto - berein}')
    if V[P]['AX63'].value != 0:
        m.append('AX63 muesste 0 sein (Alternative)')
    return m


def s15(wb):
    """MwSt-Schalter auf netto: Nettoumsatz muss dem Volumen entsprechen."""
    d = wb['📋 Definitionen & Klärung']
    # Schalterzelle relativ zum benannten Bereich Netto_Faktor bestimmen
    fak = wb.defined_names['Netto_Faktor'].attr_text.split('!')[1].replace('$', '')
    schalter = f"C{int(fak[1:]) - 3}"
    d[schalter] = 'netto exkl. MwSt'


def s15x(V):
    m = []
    p = V[P]
    if abs((p['Y6'].value or 0) - 390000) > 0.01:
        m.append(f'Netto bei Schalter «netto» erwartet 390000, ist {p["Y6"].value}')
    if abs((p['O6'].value or 0) - 118000) > 0.01:
        m.append(f'Marge erwartet 118000, ist {p["O6"].value}')
    return m


def s16(wb):
    """Einstand unvollstaendig: keine Marge, aber Hinweis."""
    ws = wb[P]
    ws['K6'] = None            # Transport loeschen
    ws['AF6'] = 'PO-X'
    ws['AG6'] = __import__('datetime').date(2026, 1, 1)


def s16x(V):
    m = []
    p = V[P]
    if p['AR6'].value != 0:
        m.append('Einstand duerfte nicht als vollstaendig gelten')
    if p['O6'].value not in (None, ''):
        m.append(f'Marge muesste leer sein, ist {p["O6"].value!r}')
    if 'Einstand' not in str(p['AH6'].value):
        m.append(f'Pruefstatus muesste den Einstand nennen: {p["AH6"].value!r}')
    return m


def s17(wb):
    """Einen umgeschluesselten 50-%-Deal auf 60 % hochstufen."""
    ws = wb[P]
    for r in range(6, 90):
        if isinstance(ws[f'AI{r}'].value, (int, float)) and abs(ws[f'AI{r}'].value - 0.5) < 1e-9:
            ws[f'S{r}'] = 0.6
            break


def s17x(V):
    m = []
    p = V[P]
    ziel = next((r for r in range(6, 90)
                 if isinstance(p[f'AI{r}'].value, (int, float))
                 and abs(p[f'AI{r}'].value - 0.5) < 1e-9), None)
    if ziel is None:
        return ['keine umgeschluesselte 50-%-Zeile gefunden']
    if p[f'AJ{ziel}'].value != 0.5:
        m.append(f'Faktor nach Hochstufung erwartet 50 %, ist {p[f"AJ{ziel}"].value}')
    vol = p[f'I{ziel}'].value or 0
    if abs((p[f'Q{ziel}'].value or 0) - vol * 0.5) > 0.01:
        m.append(f'Gew.Wert erwartet {vol*0.5}, ist {p[f"Q{ziel}"].value}')
    return m

SZENARIEN = [
    ('neuer_deal', s1, s1x),
    ('won_wird_lost', s2, s2x),
    ('wahrsch_leer', s3, s3x),
    ('prozent_als_60', s4, s4x),
    ('bandgrenzen', s5, s5x),
    ('alles_geloescht', s6, s6x),
    ('volumen_text', s7, s7x),
    ('ohne_status', s8, s8x),
    ('gleiches_volumen', s9, s9x),
    ('letzte_zeile_860', s10, s10x),
    ('unter_druckbereich', s11, s11x),
    ('lost_mit_90_prozent', s12, s12x),
    ('won_belegt', s13, s13x),
    ('variante_ausschliessen', s14, s14x),
    ('mwst_schalter_netto', s15, s15x),
    ('einstand_unvollstaendig', s16, s16x),
    ('fuenfzig_auf_sechzig', s17, s17x),
]

for name, mut, ex in SZENARIEN:
    f, j = build(name, mut)
    if j.get('status') != 'success':
        results.append((name, 'RECALC', str(j)[:200], ['recalc fehlgeschlagen']))
        continue
    verify(f, name, ex)
    os.remove(f)

print('=' * 74)
ok = True
for name, model, devi, msgs in results:
    status = 'OK ' if not msgs else 'FEHLER'
    print(f'[{status}] {name:<18} {model} | {devi}')
    for m in msgs:
        ok = False
        print(f'        -> {m}')
print('=' * 74)
print('ALLE VERHALTENSTESTS BESTANDEN' if ok else 'ES GIBT BEFUNDE')
sys.exit(0 if ok else 1)
