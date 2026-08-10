#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""A) Statischer Formel-Audit: jede Formel parsen, jeden Bezug pruefen."""
import openpyxl, re, sys
from collections import Counter, defaultdict

F = 'CH_MiT_Strom_Customer_CEO_CFO_MASTER.xlsx'
wb = openpyxl.load_workbook(F)
val = openpyxl.load_workbook(F, data_only=True)
sheets = set(wb.sheetnames)
problems = []

# von LibreOffice/Excel ohne Prefix evaluierbare Funktionen
KNOWN = {
    'DATE', 'YEAR', 'MONTH', 'DAY', 'EOMONTH',
    'IF', 'IFERROR', 'AND', 'OR', 'NOT', 'SUM', 'SUMIF', 'SUMIFS', 'COUNT',
    'COUNTIF', 'COUNTIFS', 'INDEX', 'MATCH', 'LARGE', 'SMALL', 'LOOKUP',
    'VLOOKUP', 'TEXT', 'LEFT', 'RIGHT', 'MID', 'LEN', 'ROW', 'COLUMN',
    'TODAY', 'NOW', 'ROUND', 'ABS', 'ISNUMBER', 'ISTEXT', 'ISBLANK', 'N',
    'SUMPRODUCT', 'MAX', 'MIN', 'AVERAGE', 'CONCATENATE', 'TRIM', 'VALUE',
}
# post-2007 Funktionen, die zwingend _xlfn. brauchen
NEEDS_PREFIX = {'TEXTJOIN', 'CONCAT', 'IFS', 'SWITCH', 'MAXIFS', 'MINIFS'}
# in dieser Umgebung nicht evaluierbar
FORBIDDEN = {'XLOOKUP', 'XMATCH', 'SORT', 'FILTER', 'UNIQUE', 'SEQUENCE',
             'LET', 'LAMBDA', 'TEXTSPLIT', 'TOCOL', 'TOROW'}

fn_re = re.compile(r'(?<![A-Z0-9_.])([A-Z][A-Z0-9_.]{1,20})\s*\(')
# Blattbezug: 'Name'!  oder Name!
sheetref_re = re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_. ]*))!")
extref_re = re.compile(r'\[\d+\]')

funcs = Counter()
formula_cells = 0
per_sheet = Counter()

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not (isinstance(v, str) and v.startswith('=')):
                continue
            formula_cells += 1
            per_sheet[ws.title] += 1
            loc = f'{ws.title}!{c.coordinate}'

            # 1) externe Dateibezuege
            if extref_re.search(v):
                problems.append(('EXTERN', loc, v[:100]))

            # 2) Funktionsnamen
            for fn in fn_re.findall(v):
                base = fn.split('.')[-1] if fn.startswith('_xlfn.') else fn
                funcs[fn] += 1
                if fn in FORBIDDEN:
                    problems.append(('VERBOTEN', loc, fn))
                elif fn in NEEDS_PREFIX:
                    problems.append(('PREFIX-FEHLT', loc, fn))
                elif not fn.startswith('_xlfn.') and fn not in KNOWN:
                    problems.append(('UNBEKANNT', loc, fn))

            # 3) Blattbezuege - existieren die Blaetter?
            # Strings maskieren, damit "!" in Text nicht als Bezug zaehlt
            masked = re.sub(r'"[^"]*"', '""', v)
            for q, uq in sheetref_re.findall(masked):
                name = q or uq
                if name not in sheets:
                    problems.append(('BLATT-FEHLT', loc, name))
                # Blattname mit Leerzeichen/Sonderzeichen muss quotiert sein
                if uq and (' ' in uq or any(ch in uq for ch in '⚖️📊📄📑')):
                    problems.append(('NICHT-QUOTIERT', loc, uq))

            # 4) kaputte Fehlerwerte direkt in der Formel
            for err in ('#REF!', '#NAME?', '#DIV/0!', '#VALUE!', '#N/A', '#NULL!', '#NUM!'):
                if err in v:
                    problems.append(('FEHLER-IN-FORMEL', loc, err))

            # 5) LibreOffice schreibt unparsbare Formeln klein zurueck
            body = v[1:]
            if body and body[:1].islower() and not body.startswith('_xlfn'):
                problems.append(('KLEINSCHREIBUNG', loc, v[:80]))

print('=' * 70)
print(f'Formelzellen gesamt : {formula_cells}')
for s, n in per_sheet.most_common():
    print(f'   {s:<28} {n:>6}')
print()
print('Verwendete Funktionen:')
for fn, n in sorted(funcs.items()):
    print(f'   {fn:<16} {n:>6}')
print()
if problems:
    print(f'!!! {len(problems)} Befunde:')
    for kind, loc, det in problems[:60]:
        print(f'   [{kind}] {loc}: {det}')
else:
    print('OK - keine strukturellen Formelprobleme gefunden.')

# ---- Fehlerwerte im berechneten Ergebnis -------------------------------
ERRS = ('#REF!', '#NAME?', '#DIV/0!', '#VALUE!', '#N/A', '#NULL!', '#NUM!', 'Err:')
found = []
for ws in val.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and any(e in c.value for e in ERRS):
                found.append(f'{ws.title}!{c.coordinate} = {c.value}')
print()
print(f'Fehlerwerte in berechneten Zellen: {len(found)}')
for f in found[:40]:
    print('   ', f)

sys.exit(1 if (problems or found) else 0)
