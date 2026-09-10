#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Konsistenzkontrolle fuer den September-Stand (Spaltenlage OHNE die
geloeschte Spalte «Wahr. % bisher»: _GewFaktor=AI … _OffertSort=BC).

Die Kernzahlen sind hier KEINE festen Sollwerte mehr (die Pipeline
traegt neue Daten) — geprueft wird interne Konsistenz: jedes Blatt
muss dieselben Zahlen zeigen, die sich unabhaengig aus den
Pipeline-Rohdaten ergeben.

K2  Keine Fehlerwerte (data_only-Scan).
K3  Keine Ganzspaltenbezuege, kein Bereich ueber Zeile 864.
K4  Keine XVERWEIS/FILTER/EINDEUTIG/SORTIEREN/TEXTVERKETTEN/LET.
K7  Keine Datenzeile ohne Formeln in den Formelspalten.
R1  Bandtabelle unveraendert (Regel 1).
R2  Keine Marge in einem Berichtsblatt (Regel 2).
V   Verknuepfungspruefung: Dashboard/CEO/ExecPDF/⚖️/_data gegen das
    unabhaengige Python-Modell der Pipeline.
"""
import re
import sys
import openpyxl
from openpyxl.utils import get_column_letter

F = sys.argv[1] if len(sys.argv) > 1 else 'BURAK_MASTER_FIX.xlsx'
fehler = []
wb = openpyxl.load_workbook(F)
V = openpyxl.load_workbook(F, data_only=True)
p, vp = wb['MiT Strom Pipeline'], V['MiT Strom Pipeline']

# ------------------------------------------------ unabhaengiges Modell
AKTIV = {'WON', 'offered', 'to be offered', 'on hold', 'follow-up',
         'Evaluation', 'In evaluation', 'tbd'}
BAND = [(0, 0.4499999, 0.0), (0.45, 0.5999999, 0.3), (0.6, 0.8999999, 0.5),
        (0.9, 0.9999999, 0.9), (1.0, 1.0, 1.0)]

def faktor(s):
    if not isinstance(s, (int, float)):
        return 0.0
    s = s if s <= 1 else s / 100.0
    for von, bis, f in BAND:
        if von <= s <= bis:
            return f
    return 0.0

Z = []
for r in range(6, 861):
    b = vp[f'B{r}'].value
    if b in (None, ''):
        continue
    vol = vp[f'I{r}'].value
    if not isinstance(vol, (int, float)):
        # wie _I_Num (=IFERROR(I*1,0)): als Text erfasste Zahlen wandeln
        try:
            vol = float(str(vol).replace("'", "").replace('’', '')
                        .replace(' ', '').replace(',', '.'))
        except (TypeError, ValueError):
            vol = 0.0
    st = vp[f'R{r}'].value
    s = vp[f'S{r}'].value
    Z.append(dict(r=r, kunde=b, kt=vp[f'C{r}'].value, seg=vp[f'D{r}'].value,
                  vol=float(vol), st=st, fak=faktor(s),
                  aktiv=st in AKTIV, won=st == 'WON',
                  vertrag=vp[f'AB{r}'].value or '',
                  po=vp[f'AF{r}'].value, belegdat=vp[f'AG{r}'].value,
                  o=vp[f'O{r}'].value, pe=vp[f'P{r}'].value,
                  variante=vp[f'AD{r}'].value or ''))

n = len(Z)
won_n = sum(1 for z in Z if z['won'])
won = sum(z['vol'] for z in Z if z['won'])
akt_n = sum(1 for z in Z if z['aktiv'])
akt = sum(z['vol'] for z in Z if z['aktiv'])
gew = sum(z['vol'] * z['fak'] for z in Z if z['aktiv'])
gew_won = sum(z['vol'] * z['fak'] for z in Z if z['aktiv'] and z['won'])
gew_off = gew - gew_won
bereinigt = sum(z['vol'] for z in Z
                if z['aktiv'] and z['variante'] != 'Alternative – zählt nicht')
belegt_n = 0   # PO+Datum+Einstand: Einstand aus AQ lesen (Kostenspalten-Logik)
belegt = 0.0
for z in Z:
    if not z['won']:
        continue
    aq = vp[f'AQ{z["r"]}'].value    # _EinstandOK
    if z['po'] not in (None, '') and z['belegdat'] not in (None, '') and aq == 1:
        belegt_n += 1
        belegt += z['vol']
zeit_n = sum(1 for z in Z if z['aktiv'] and isinstance(z['o'], object)
             and hasattr(z['o'], 'year') and hasattr(z['pe'], 'year')
             and z['pe'] >= z['o'])

print(f'Modell: {n} Zeilen · WON {won_n}/{won:,.2f} · aktiv {akt_n}/{akt:,.2f} '
      f'· gewichtet {gew:,.2f} (WON {gew_won:,.2f} + Offerten {gew_off:,.2f}) '
      f'· belegt {belegt_n}/{belegt:,.2f} · Zeitraum {zeit_n}')


def soll(koord, wert, erwartet, toleranz=0.01):
    if isinstance(erwartet, str):
        if str(wert) != erwartet:
            fehler.append(f'V {koord}: «{wert}» statt «{erwartet}»')
        return
    if wert is None or not isinstance(wert, (int, float)) or abs(wert - erwartet) > toleranz:
        fehler.append(f'V {koord}: {wert} statt {erwartet:,.2f}')

# ------------------------------------------------ V: Verknuepfungen
dash, ceo = V['Dashboard'], V['CEO Report']
pdf, w, dt = V['📑 Executive PDF'], V['⚖️ Wahrscheinlichkeit'], V['_data']
d = V['📋 Definitionen & Klärung']

soll('Dashboard!A6', dash['A6'].value, won)
soll('Dashboard!A7', dash['A7'].value, won_n)
soll('Dashboard!I6', dash['I6'].value, akt)
soll('Dashboard!I7', dash['I7'].value, akt_n)
soll('CEO!A6', ceo['A6'].value, won)
soll('CEO!A7', ceo['A7'].value, won_n)
soll('⚖️!E22', w['E22'].value, akt_n)
soll('⚖️!F22', w['F22'].value, akt)
soll('⚖️!G22', w['G22'].value, gew)
soll('CEO!G215 (Zerlegung WON)', ceo['G215'].value, gew_won)
soll('CEO!G216 (Zerlegung Offerten)', ceo['G216'].value, gew_off)
soll('CEO!G217 (Summe Zerlegung)', ceo['G217'].value, gew)
soll('CEO!E264 (Weg 1)', ceo['E264'].value, gew)
soll('CEO!E265 (Weg 2)', ceo['E265'].value, gew)
soll('CEO!E266 (Abweichung)', ceo['E266'].value, 0.0)
soll('CEO!D220 (WON gemeldet n)', ceo['D220'].value, won_n)
soll('CEO!E220 (WON gemeldet)', ceo['E220'].value, won)
soll('CEO!D222 (belegt n)', ceo['D222'].value, belegt_n)
soll('CEO!E222 (belegt)', ceo['E222'].value, belegt)
soll('CEO!E224 (brutto)', ceo['E224'].value, akt)
soll('CEO!E225 (bereinigt)', ceo['E225'].value, bereinigt)
soll('CEO!E260 (Vertragsarten-Total)', ceo['E260'].value, won)
soll('CEO!D260', ceo['D260'].value, won_n)
soll('PDF!G14 (gew TOTAL)', pdf['G14'].value, gew)
soll('PDF!G15', pdf['G15'].value, gew_won)
soll('PDF!G16', pdf['G16'].value, gew_off)
soll('PDF!F21', pdf['F21'].value, won_n)
soll('PDF!G21', pdf['G21'].value, won)
soll('PDF!F22', pdf['F22'].value, belegt_n)

# Vertragsarten-Zeilen
for i, art in enumerate(['Einzelauftrag', 'Rahmenvertrag – Abruf bestätigt',
                         'Abrufbereitschaft', 'Option / Reservation', '']):
    r = 255 + i
    n_soll = sum(1 for z in Z if z['won'] and z['vertrag'] == art)
    v_soll = sum(z['vol'] for z in Z if z['won'] and z['vertrag'] == art)
    soll(f'CEO!D{r}', ceo[f'D{r}'].value, n_soll)
    soll(f'CEO!E{r}', ceo[f'E{r}'].value, v_soll)

# _data: Segment-Sortierung deckt die aktive Pipeline ab
segsum = sum(v for v in (dt[f'S{r}'].value for r in range(2, 34))
             if isinstance(v, (int, float)))
soll('_data Segmentsumme', segsum, akt)
kantsum = sum(v for v in (dt[f'T{r}'].value for r in range(14, 40))
              if isinstance(v, (int, float)))
soll('_data Kantonsumme', kantsum, akt)
soll('_data!Y2', dt['Y2'].value, gew_won)
soll('_data!Y3', dt['Y3'].value, gew_off)

# Status-Gruppen (Dashboard-Kacheln)
off_grp = sum(z['vol'] for z in Z if z['st'] in ('offered', 'to be offered', 'on hold'))
soll('Dashboard!D6 (Offerte-Box)', dash['D6'].value, off_grp)

# Erfassungsliste: Anzahl Eintraege = aktive ohne Zeitraum
offen_zeit = akt_n - zeit_n
liste = sum(1 for r in range(82, 132) if d[f'B{r}'].value not in (None, ''))
if liste != min(offen_zeit, 50):
    fehler.append(f'V Erfassungsliste: {liste} Einträge statt {min(offen_zeit,50)}')

# ------------------------------------------------ K2 Fehlerwerte
FW = ('#REF!', '#VALUE!', '#N/A', '#DIV/0!', '#NAME?', '#NUM!', '#NULL!')
k2 = [(ws.title, c.coordinate, c.value) for ws in V.worksheets
      for row in ws.iter_rows() for c in row
      if isinstance(c.value, str) and c.value in FW]
for t, k, v in k2[:8]:
    fehler.append(f'K2 {t}!{k}: {v}')
if len(k2) > 8:
    fehler.append(f'K2 … und {len(k2)-8} weitere')

# ------------------------------------------------ K3 + K4
GANZ = re.compile(r"(?<![A-Z0-9_.$])\$?[A-Z]{1,3}:\$?[A-Z]{1,3}(?![A-Z0-9_(])")
BER = re.compile(r"\$?[A-Z]{1,3}\$?(\d+):\$?[A-Z]{1,3}\$?(\d+)")
FN = re.compile(r"\b(XLOOKUP|_xlfn\.FILTER|UNIQUE|SORTBY|SORT|TEXTJOIN|LET)\s*\(")
k34 = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            f = c.value
            if not isinstance(f, str) or not f.startswith('='):
                continue
            if GANZ.search(f) or FN.search(f) or any(
                    int(a) > 864 or int(b) > 864 for a, b in BER.findall(f)):
                k34 += 1
                if k34 <= 5:
                    fehler.append(f'K3/K4 {ws.title}!{c.coordinate}: {f[:60]}')

# ------------------------------------------------ K7 Formel-Luecken
FORMELSPALTEN = ['G', 'Q', 'Y', 'Z', 'AH'] + \
    [get_column_letter(c) for c in range(35, 56) if get_column_letter(c) != 'AV']
for col in FORMELSPALTEN:
    luecken = [r for r in range(6, 861)
               if not (isinstance(p[f'{col}{r}'].value, str)
                       and p[f'{col}{r}'].value.startswith('='))]
    if luecken:
        fehler.append(f'K7 Pipeline!{col}: {len(luecken)} Zeilen ohne Formel '
                      f'(z.B. {luecken[:4]})')

# ------------------------------------------------ R1 Bandtabelle
wf = wb['⚖️ Wahrscheinlichkeit']
BS = [('0 – 44 %', 0, 0.4499999, 0), ('45 – 59 %', 0.45, 0.5999999, 0.3),
      ('60 – 89 %', 0.6, 0.8999999, 0.5), ('90 – 99 %', 0.9, 0.9999999, 0.9),
      ('100 %  ·  WON', 1, 1, 1)]
for i, (label, von, bis, fk) in enumerate(BS):
    r = 17 + i
    ist = (wf[f'A{r}'].value, wf[f'B{r}'].value, wf[f'C{r}'].value, wf[f'D{r}'].value)
    if ist != (label, von, bis, fk):
        fehler.append(f'R1 Band Zeile {r}: {ist}')
if wf['A22'].value != 'TOTAL aktiv':
    fehler.append('R1 Totalzeile fehlt')

# ------------------------------------------------ R2 Marge
MARGE = re.compile(r'\bMarge\b|\bDeckungsbeitrag\b|Margen[- ]?%', re.I)
OK = re.compile(r'nicht mehr|entfernt|keine Marge|ohne Marge|nicht ausgewiesen', re.I)
for name in ['📑 Executive PDF', '📄 Report', 'CEO Report', 'Dashboard', '📊 Diagramme']:
    for row in wb[name].iter_rows():
        for c in row:
            if isinstance(c.value, str) and MARGE.search(c.value) and not OK.search(c.value):
                fehler.append(f'R2 {name}!{c.coordinate}: {c.value[:60]}')

print(f'Befunde: {len(fehler)}')
for f in fehler[:40]:
    print('  ', f)
sys.exit(1 if fehler else 0)
