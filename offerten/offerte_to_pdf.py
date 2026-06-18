# -*- coding: utf-8 -*-
"""
Offerte -> PDF auf Knopfdruck (ohne Excel, ohne LibreOffice).

    python3 offerten/offerte_to_pdf.py [datei.xlsx] [ausgabe.pdf]

Standard:
    Eingabe : offerten/Offerten_Master_Vorlage.xlsx
    Ausgabe : offerten/<Offerte-Nr>.pdf

Liest die ausgefüllte Mappe (Blätter Offerte / Stammdaten / Artikel), löst
Artikelpreise auf, rechnet Zwischentotal, Rabatt, MwSt und Endbetrag und
rendert ein markenkonformes PDF (DIN A4) mit fpdf2.
"""
import os, sys, datetime
from openpyxl import load_workbook
from fpdf import FPDF

HERE = os.path.dirname(os.path.abspath(__file__))

# --- Farben (RGB) -----------------------------------------------------------
ORANGE = (226, 98, 10)
DARK   = (31, 41, 51)
GREY   = (123, 135, 148)
LINE   = (203, 210, 217)
ZEBRA  = (244, 246, 248)
TOTBG  = (252, 238, 227)


def chf(x):
    """18'454.74 – Schweizer Tausender-Apostroph."""
    s = f"{x:,.2f}".replace(",", "'")
    return s


def fmt_date(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%d.%m.%Y")
    return "" if v is None else str(v)


def read_workbook(src):
    """Liest alle für das PDF nötigen Werte und rechnet die Summen."""
    wb = load_workbook(src, data_only=False)
    sd, ar, of = wb["Stammdaten"], wb["Artikel"], wb["Offerte"]

    def sdv(row):           # Wert aus Stammdaten Spalte B
        return sd.cell(row=row, column=2).value

    stamm = dict(
        firma=sdv(2), strasse=sdv(3), plz=sdv(4), tel=sdv(5), mail=sdv(6),
        web=sdv(7), zusatz=sdv(8), agb=sdv(9),
        sb_name=sdv(12), sb_tel=sdv(13), sb_mail=sdv(14),
        rabatt=sdv(17) or 0, mwst=sdv(18) or 0, waehrung=sdv(19) or "CHF",
        einleitung=sdv(22), gruss=sdv(23), anmerkung=sdv(24),
    )

    # Artikel-Preisliste -> {beschreibung: (einheit, preis)}
    preis = {}
    for r in range(2, ar.max_row + 1):
        d = ar.cell(row=r, column=1).value
        if d:
            preis[str(d).strip()] = (ar.cell(row=r, column=2).value,
                                     ar.cell(row=r, column=3).value)

    # Kunde + Meta
    kunde = [of[f"A{r}"].value for r in (8, 9, 10, 11)]
    kunde = [k for k in kunde if k]
    nr        = of["G8"].value
    datum     = of["G9"].value
    mietbeginn= of["G10"].value
    mietende  = of["G11"].value
    mindest   = of["G13"].value
    zeitraum = ""
    if isinstance(mietbeginn, (datetime.date, datetime.datetime)) and \
       isinstance(mietende, (datetime.date, datetime.datetime)):
        zeitraum = f"{(mietende - mietbeginn).days + 1} Tage"

    # Positionen einlesen (Tabellenkopf in Zeile 21, Daten ab 22)
    positions = []
    for r in range(22, of.max_row + 1):
        desc = of.cell(row=r, column=3).value
        if not desc or str(desc).startswith("="):
            continue
        anz   = of.cell(row=r, column=2).value
        unit  = of.cell(row=r, column=4).value
        price = of.cell(row=r, column=5).value
        wochen= of.cell(row=r, column=6).value
        # Formeln/Leerwerte über die Artikelliste auflösen
        if unit is None or (isinstance(unit, str) and unit.startswith("=")):
            unit = preis.get(str(desc).strip(), ("", 0))[0]
        if price is None or (isinstance(price, str) and price.startswith("=")):
            price = preis.get(str(desc).strip(), ("", 0))[1]
        try:
            anz = float(anz) if anz not in (None, "") else 1
        except (TypeError, ValueError):
            anz = 1
        try:
            price = float(price) if price not in (None, "") else 0.0
        except (TypeError, ValueError):
            price = 0.0
        try:
            wochen = float(wochen) if wochen not in (None, "") else 1
        except (TypeError, ValueError):
            wochen = 1
        pos = anz * price * wochen
        positions.append(dict(anz=anz, desc=str(desc), unit=unit or "",
                              price=price, wochen=wochen, total=pos))

    subtotal = round(sum(p["total"] for p in positions), 2)
    rabatt_b = round(subtotal * float(stamm["rabatt"]), 2)
    netto    = round(subtotal - rabatt_b, 2)
    mwst_b   = round(netto * float(stamm["mwst"]), 2)
    endbetrag= round(netto + mwst_b, 2)

    return dict(
        stamm=stamm, kunde=kunde, nr=nr, datum=datum, mietbeginn=mietbeginn,
        mietende=mietende, mindest=mindest, zeitraum=zeitraum,
        positions=positions, subtotal=subtotal, rabatt_b=rabatt_b,
        netto=netto, mwst_b=mwst_b, endbetrag=endbetrag,
    )


class OffertePDF(FPDF):
    def __init__(self, data):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.data = data
        self.set_auto_page_break(True, margin=16)
        self.set_margins(15, 14, 15)

    def footer(self):
        s = self.data["stamm"]
        self.set_y(-14)
        self.set_draw_color(*LINE); self.set_line_width(0.2)
        self.line(15, self.get_y(), 195, self.get_y())
        self.set_y(-12)
        self.set_font("Helvetica", "", 7)
        self.set_text_color(*GREY)
        foot = " · ".join(filter(None, [s["firma"], s["strasse"], s["plz"],
                                        s["tel"], s["web"]]))
        self.cell(0, 4, _t(foot), align="C", new_x="LMARGIN", new_y="NEXT")
        if s["zusatz"]:
            self.cell(0, 4, _t(s["zusatz"]), align="C")
        self.set_font("Helvetica", "", 7)
        self.set_xy(170, -12)
        self.cell(25, 4, f"Seite {self.page_no()}", align="R")


def _t(s):
    """fpdf Helvetica = latin-1; Sonderzeichen absichern."""
    if s is None:
        return ""
    return str(s).encode("latin-1", "replace").decode("latin-1")


def build_pdf(data, dst):
    pdf = OffertePDF(data)
    s = data["stamm"]
    pdf.add_page()

    # ----- Kopf: Marke links / Sachbearbeiter rechts
    pdf.set_xy(15, 14)
    pdf.set_font("Helvetica", "B", 22); pdf.set_text_color(*ORANGE)
    pdf.cell(110, 9, "MOBIL IN TIME")
    pdf.set_xy(15, 24)
    pdf.set_font("Helvetica", "", 9); pdf.set_text_color(*DARK)
    pdf.cell(110, 4.5, _t(s["firma"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(*GREY)
    pdf.set_x(15); pdf.cell(110, 4.5, _t(f"{s['strasse']}  ·  {s['plz']}"),
                            new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(15); pdf.cell(110, 4.5,
                            _t(f"Tel. {s['tel']}  ·  {s['mail']}  ·  {s['web']}"))
    # Sachbearbeiter
    pdf.set_xy(120, 14)
    pdf.set_font("Helvetica", "B", 8); pdf.set_text_color(*GREY)
    pdf.cell(75, 4.5, "IHR KONTAKT", align="R", new_x="LEFT", new_y="NEXT")
    pdf.set_x(120); pdf.set_font("Helvetica", "B", 10); pdf.set_text_color(*DARK)
    pdf.cell(75, 5, _t(s["sb_name"]), align="R", new_x="LEFT", new_y="NEXT")
    pdf.set_x(120); pdf.set_font("Helvetica", "", 8); pdf.set_text_color(*GREY)
    pdf.cell(75, 4.5, _t(f"Tel. {s['sb_tel']}"), align="R", new_x="LEFT", new_y="NEXT")
    pdf.set_x(120); pdf.cell(75, 4.5, _t(s["sb_mail"]), align="R")

    # ----- Kunde links / Offerte-Meta rechts
    y0 = 44
    pdf.set_xy(15, y0)
    pdf.set_font("Helvetica", "B", 8); pdf.set_text_color(*GREY)
    pdf.cell(90, 4.5, "KUNDE", new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(15); pdf.set_font("Helvetica", "B", 10); pdf.set_text_color(*DARK)
    if data["kunde"]:
        pdf.cell(90, 5, _t(data["kunde"][0]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9); pdf.set_text_color(*DARK)
    for line in data["kunde"][1:]:
        pdf.set_x(15); pdf.cell(90, 4.5, _t(line), new_x="LMARGIN", new_y="NEXT")

    # Meta-Box rechts
    meta = [
        ("Offerte-Nr.", _t(data["nr"])),
        ("Datum",       fmt_date(data["datum"])),
        ("Mietbeginn",  fmt_date(data["mietbeginn"])),
        ("Mietende",    fmt_date(data["mietende"])),
        ("Mietzeitraum",_t(data["zeitraum"])),
        ("Mindestmiete",_t(data["mindest"])),
    ]
    pdf.set_xy(120, y0 - 2)
    pdf.set_font("Helvetica", "B", 13); pdf.set_text_color(*DARK)
    pdf.cell(75, 7, "OFFERTE", align="R", new_x="LEFT", new_y="NEXT")
    for lab, val in meta:
        pdf.set_x(120)
        pdf.set_font("Helvetica", "B", 8); pdf.set_text_color(*GREY)
        pdf.cell(33, 4.8, lab, align="R")
        pdf.set_font("Helvetica", "", 9); pdf.set_text_color(*DARK)
        pdf.cell(42, 4.8, val, align="R", new_x="LEFT", new_y="NEXT")

    # ----- Anrede + Einleitung
    pdf.ln(6)
    pdf.set_x(15)
    name = data["kunde"][0] if data["kunde"] else "Herr"
    parts = name.split()
    if parts and parts[0] == "Herr":
        anrede = f"Sehr geehrter Herr {parts[-1]}"
    elif parts and parts[0] == "Frau":
        anrede = f"Sehr geehrte Frau {parts[-1]}"
    else:
        anrede = f"Guten Tag {name}"
    pdf.set_font("Helvetica", "", 10); pdf.set_text_color(*DARK)
    pdf.multi_cell(180, 5, _t(anrede))
    pdf.ln(1)
    intro = (s["einleitung"] or "").replace("{AGB}", str(s["agb"] or ""))
    pdf.set_x(15)
    pdf.multi_cell(180, 5, _t(intro))
    pdf.ln(2)
    pdf.set_x(15)
    pdf.set_font("Helvetica", "B", 10)
    pdf.multi_cell(180, 5, "Preisaufstellung - Preise reflektieren Mengen")
    pdf.ln(1)

    # ----- Positionstabelle
    # Spalten: Pos | Anz | Beschreibung | Einh. | Preis/Wo. | Wo. | Position
    cw = [10, 12, 86, 14, 24, 11, 23]   # Summe = 180
    headers = ["Pos", "Anz", "Beschreibung", "Einh.", "Preis/Wo.", "Wo.", "Position"]
    aligns  = ["C", "C", "L", "C", "R", "C", "R"]
    pdf.set_x(15)
    pdf.set_fill_color(*DARK); pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 8.5)
    for w, h, a in zip(cw, headers, aligns):
        pdf.cell(w, 7, h, align=a, fill=True)
    pdf.ln(7)

    pdf.set_font("Helvetica", "", 8.5)
    for i, p in enumerate(data["positions"], start=1):
        zebra = (i % 2 == 0)
        # Höhe anhand der Beschreibung bestimmen
        desc = _t(p["desc"])
        pdf.set_text_color(*DARK)
        lines = pdf.multi_cell(cw[2], 4.6, desc, dry_run=True, output="LINES")
        nlines = max(1, len(lines))
        rh = max(6.5, nlines * 4.6 + 1.8)
        x0, y = 15, pdf.get_y()
        if y + rh > pdf.h - 18:
            pdf.add_page(); y = pdf.get_y()
        if zebra:
            pdf.set_fill_color(*ZEBRA)
            pdf.rect(x0, y, sum(cw), rh, style="F")
        pdf.set_draw_color(*LINE); pdf.set_line_width(0.15)
        pdf.line(x0, y + rh, x0 + sum(cw), y + rh)
        # Zellen
        cells = [
            (cw[0], str(i), "C"),
            (cw[1], f"{p['anz']:g}", "C"),
            (None, None, None),                       # Beschreibung separat
            (cw[3], _t(p["unit"]), "C"),
            (cw[4], chf(p["price"]), "R"),
            (cw[5], f"{p['wochen']:g}", "C"),
            (cw[6], chf(p["total"]), "R"),
        ]
        cx = x0
        for idx, (w, txt, a) in enumerate(cells):
            if idx == 2:
                # Beschreibung mehrzeilig, vertikal leicht eingerückt
                pdf.set_xy(cx, y + 1.4)
                pdf.multi_cell(cw[2], 4.6, desc, align="L")
                cx += cw[2]
                continue
            pdf.set_xy(cx, y + (rh - 4.6) / 2)
            pdf.cell(w, 4.6, txt, align=a)
            cx += w
        pdf.set_xy(x0, y + rh)

    # ----- Summenblock (rechtsbündig)
    def sumrow(label, value, bold=False, fill=None, big=False):
        y = pdf.get_y()
        if fill:
            pdf.set_fill_color(*fill)
            pdf.rect(15, y, 180, 8 if big else 6.5, style="F")
        pdf.set_xy(15, y + (1.8 if big else 1.1))
        pdf.set_font("Helvetica", "B" if bold else "", 11 if big else 9.5)
        pdf.set_text_color(*DARK)
        pdf.cell(150, 4.8, label, align="R")
        pdf.cell(30, 4.8, value, align="R")
        pdf.set_xy(15, y + (8 if big else 6.5))

    pdf.ln(1)
    rab_pct = f"{float(s['rabatt'])*100:.1f}".rstrip("0").rstrip(".")
    mwst_pct = f"{float(s['mwst'])*100:.1f}".rstrip("0").rstrip(".")
    sumrow("Zwischentotal (GESAMT)", chf(data["subtotal"]), bold=True)
    sumrow(f"Rabatt  {rab_pct}%", "-" + chf(data["rabatt_b"]))
    sumrow("Total nach Rabatt", chf(data["netto"]), bold=True)
    if float(s["mwst"]) > 0:
        sumrow(f"MwSt  {mwst_pct}%", chf(data["mwst_b"]))
    sumrow(f"Endbetrag ({s['waehrung']})", chf(data["endbetrag"]),
           bold=True, fill=TOTBG, big=True)

    # ----- Anmerkung + Gruss
    pdf.ln(4)
    pdf.set_x(15); pdf.set_font("Helvetica", "I", 7.5); pdf.set_text_color(*GREY)
    pdf.multi_cell(180, 3.8, _t("Anmerkung: " + (s["anmerkung"] or "")))
    pdf.ln(4)
    pdf.set_x(15); pdf.set_font("Helvetica", "", 10); pdf.set_text_color(*DARK)
    pdf.cell(0, 5, _t(s["gruss"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(15); pdf.set_font("Helvetica", "B", 10)
    pdf.cell(0, 5, _t(s["sb_name"]))

    pdf.output(dst)
    return dst


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else \
        os.path.join(HERE, "Offerten_Master_Vorlage.xlsx")
    if not os.path.exists(src):
        sys.exit(f"Nicht gefunden: {src}")
    data = read_workbook(src)
    nr = str(data["nr"] or "Offerte").strip().replace("/", "-").replace(" ", "_")
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, f"{nr}.pdf")
    build_pdf(data, dst)
    print("PDF erstellt:", dst)
    print(f"  Zwischentotal {chf(data['subtotal'])}  ·  "
          f"Rabatt {chf(data['rabatt_b'])}  ·  Endbetrag {chf(data['endbetrag'])} "
          f"{data['stamm']['waehrung']}")


if __name__ == "__main__":
    main()
