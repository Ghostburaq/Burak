# -*- coding: utf-8 -*-
"""
Baut die Offerten-Master-Vorlage (Excel, .xlsx) auf Knopfdruck.

    python3 offerten/build_template.py

Erzeugt:  offerten/Offerten_Master_Vorlage.xlsx

Aufbau der Mappe
----------------
1. "Offerte"     – das druckfertige Dokument (DIN A4, automatische Berechnung)
2. "Stammdaten"  – Firma, Sachbearbeiter, Standardwerte (per Formel in die Offerte gezogen)
3. "Artikel"     – Preisliste; speist Dropdown + automatischen Preis-Lookup

Alle Summen, Rabatt, MwSt und der Mietzeitraum rechnen sich selbst.
Die Beispieldaten reproduzieren die hochgeladene Mobil-in-Time-Offerte.
"""

import datetime as _dt
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, NamedStyle, Protection,
)
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- Marke / Farben
ORANGE   = "E2620A"   # Akzent (Mobil in Time)
DARK     = "1F2933"   # Dunkelgrau (Text / Kopf)
GREY     = "7B8794"   # Hellgrau (Labels)
LINE     = "CBD2D9"   # Tabellen-Linien
ZEBRA    = "F4F6F8"   # Zeilen-Hintergrund
HEADFILL = "1F2933"   # Tabellenkopf-Hintergrund
TOTFILL  = "FCEEE3"   # Summen-Hintergrund (helles Orange)

CHF_FMT  = "#'##0.00"          # Schweizer Format: 18'454.74
PCT_FMT  = "0.0%"
DATE_FMT = "DD.MM.YYYY"

# ---------------------------------------------------------------- Beispieldaten
# (Anz, Beschreibung, Einheit, Einzelpreis/Woche, Wochen)
# Einzelpreis = Positionspreis aus Original / Anzahl  -> reproduziert GESAMT exakt.
POSITIONS = [
    (1,  "6MVA Trafo 16,5kV YNyn0",                                    "Stk", 5409.82,            1),
    (1,  "3MVA Trafo 16,5kV YNyn0",                                    "Stk", 1588.15,            1),
    (2,  "25 Meter x Kabel 32A",                                       "Stk", 13.70/2,            1),
    (1,  "Load Bank 6000 kVA Resistive and Reactive : 400V 3-ph @ 50 Hz", "Stk", 4960.73,         1),
    (1,  "Load Bank 3300 kVA Resistive and Reactive : 400V 3-ph @ 50 Hz", "Stk", 3087.82,         1),
    (63, "10 Meter x Einzeladerkabel 240mm² (Set A)",                  "Stk", 282.47/63,          1),
    (32, "10 Meter x Einzeladerkabel 240mm² (Set B)",                  "Stk", 143.47/32,          1),
    (1,  "25 Meter x Kabel 125A",                                      "Stk", 12.57,              1),
    (1,  "25 Meter x Kabel 63A",                                       "Stk", 9.46,               1),
    (1,  "Umweltschutzpauschale",                                      "Psch", 1085.57,           1),
    (1,  "Befreiung von der Versicherungspflicht",                     "Psch", 1860.98,           1),
]
N_BLANK = 8  # leere Erfassungszeilen mit aktiven Formeln

# Artikel-Stammliste (Beschreibung, Einheit, Einzelpreis/Woche)
ARTIKEL = [
    ("6MVA Trafo 16,5kV YNyn0",                                    "Stk",  5409.82),
    ("3MVA Trafo 16,5kV YNyn0",                                    "Stk",  1588.15),
    ("1MVA Trafo 16,5kV YNyn0",                                    "Stk",   850.00),
    ("Load Bank 6000 kVA Resistive and Reactive : 400V 3-ph @ 50 Hz", "Stk", 4960.73),
    ("Load Bank 3300 kVA Resistive and Reactive : 400V 3-ph @ 50 Hz", "Stk", 3087.82),
    ("25 Meter x Kabel 32A",                                       "Stk",     6.85),
    ("25 Meter x Kabel 63A",                                       "Stk",     9.46),
    ("25 Meter x Kabel 125A",                                      "Stk",    12.57),
    ("10 Meter x Einzeladerkabel 240mm² (Set A)",                  "Stk",     4.483651),
    ("10 Meter x Einzeladerkabel 240mm² (Set B)",                  "Stk",     4.483438),
    ("10 Meter x Einzeladerkabel 240mm²",                          "Stk",     4.48),
    ("Umweltschutzpauschale",                                      "Psch", 1085.57),
    ("Befreiung von der Versicherungspflicht",                     "Psch", 1860.98),
    ("Lieferung / Abholung (pauschal)",                            "Psch",  450.00),
    ("Montage / Inbetriebnahme (pro Std.)",                        "Std",   135.00),
]

# ---------------------------------------------------------------- Hilfsfunktionen
thin = Side(style="thin", color=LINE)
med  = Side(style="medium", color=DARK)

def border(top=False, bottom=False, left=False, right=False, color=LINE):
    s = Side(style="thin", color=color)
    return Border(
        top=s if top else None, bottom=s if bottom else None,
        left=s if left else None, right=s if right else None,
    )

def set_cell(ws, ref, value=None, font=None, fill=None, align=None,
             fmt=None, bd=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    if font:  c.font = font
    if fill:  c.fill = fill
    if align: c.alignment = align
    if fmt:   c.number_format = fmt
    if bd:    c.border = bd
    return c

# ================================================================ Mappe
wb = Workbook()

# ----------------------------------------------------------------- STAMMDATEN
sd = wb.active
sd.title = "Stammdaten"
sd.sheet_view.showGridLines = False

sd_rows = [
    ("FIRMA (Absender)", ""),
    ("Firmenname",       "Mobil in Time AG"),
    ("Strasse",          "Mattenstrasse 3"),
    ("PLZ / Ort",        "8253 Diessenhofen"),
    ("Telefon",          "+41 44 806 13 00"),
    ("E-Mail",           "info@mobilintime.com"),
    ("Web",              "www.mobilintime.com"),
    ("Zusatz",           "Ein Unternehmen der Mobil in Time Gruppe"),
    ("AGB-Hinweis (URL)","www.mobilintime.com"),
    ("", ""),
    ("SACHBEARBEITER", ""),
    ("Name",   "Burak Ücöz"),
    ("Telefon","+41 76 202 01 70"),
    ("E-Mail", "uecoez@mobilintime.com"),
    ("", ""),
    ("STANDARDWERTE", ""),
    ("Rabatt (%)",  0.10),
    ("MwSt (%)",    0.00),
    ("Währung",     "CHF"),
    ("", ""),
    ("TEXTBAUSTEINE", ""),
    ("Einleitung",
     "Wir bedanken uns für Ihr Interesse an einer Zusammenarbeit mit uns. "
     "Aufgrund der von Ihnen gemachten Angaben und gemäss unseren Allgemeinen "
     "Geschäftsbedingungen, einsehbar auf {AGB}, bieten wir Ihnen folgende "
     "Leistungen freibleibend an."),
    ("Grussformel", "Mit freundlichen Grüssen"),
    ("Anmerkung",
     "Dieses Budgetangebot ist für beide Parteien unverbindlich. Alle "
     "Anmietungen erfolgen vorbehaltlich der Verfügbarkeit bei Aufgabe der "
     "Bestellung und zu den Preisen, die im offiziellen Angebot aufgeführt "
     "sind, sobald die zusätzlichen Detailinformationen vorliegen, die sich "
     "auf Ihr Equipment und Ihren Servicebedarf beziehen."),
]
sd.column_dimensions["A"].width = 22
sd.column_dimensions["B"].width = 95
hdr_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
hdr_fill = PatternFill("solid", fgColor=ORANGE)
for i, (k, v) in enumerate(sd_rows, start=1):
    a = sd.cell(row=i, column=1, value=k)
    b = sd.cell(row=i, column=2, value=v)
    if v == "" and k and k.isupper():           # Abschnittsüberschrift
        a.font = hdr_font; a.fill = hdr_fill
        b.fill = hdr_fill
        sd.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        a.alignment = Alignment(horizontal="left", vertical="center")
    else:
        a.font = Font(name="Calibri", bold=True, color=DARK)
        b.font = Font(name="Calibri", color=DARK)
        b.alignment = Alignment(wrap_text=True, vertical="top")
    if k in ("Rabatt (%)", "MwSt (%)"):
        b.number_format = PCT_FMT
sd.sheet_state = "visible"

# Schnellzugriff auf Stammdaten-Zellen (Spalte B)
def SD(label):
    for i, (k, _) in enumerate(sd_rows, start=1):
        if k == label:
            return f"Stammdaten!$B${i}"
    raise KeyError(label)

# ----------------------------------------------------------------- ARTIKEL
ar = wb.create_sheet("Artikel")
ar.sheet_view.showGridLines = False
ar.column_dimensions["A"].width = 52
ar.column_dimensions["B"].width = 10
ar.column_dimensions["C"].width = 18
heads = ["Beschreibung", "Einheit", "Einzelpreis/Woche (CHF)"]
for j, h in enumerate(heads, start=1):
    c = ar.cell(row=1, column=j, value=h)
    c.font = hdr_font; c.fill = PatternFill("solid", fgColor=DARK)
    c.alignment = Alignment(horizontal="left", vertical="center")
for i, (desc, unit, price) in enumerate(ARTIKEL, start=2):
    ar.cell(row=i, column=1, value=desc).font = Font(name="Calibri", color=DARK)
    ar.cell(row=i, column=2, value=unit).alignment = Alignment(horizontal="center")
    pc = ar.cell(row=i, column=3, value=price)
    pc.number_format = CHF_FMT
    pc.alignment = Alignment(horizontal="right")
ar_last = 1 + len(ARTIKEL)
ARTIKEL_RANGE = f"Artikel!$A$2:$C${ar_last}"
ARTIKEL_DESCS = f"Artikel!$A$2:$A${ar_last}"
ar.freeze_panes = "A2"

# ----------------------------------------------------------------- OFFERTE
of = wb.create_sheet("Offerte")
of.sheet_view.showGridLines = False
widths = {"A": 4, "B": 6.5, "C": 40, "D": 8, "E": 14, "F": 7, "G": 15,
          "H": 2, "I": 16}
for col, w in widths.items():
    of.column_dimensions[col].width = w
of.column_dimensions["I"].hidden = True   # Hilfsspalte (Listenpreis)

f_brand = Font(name="Calibri", bold=True, size=22, color=ORANGE)
f_small = Font(name="Calibri", size=9, color=GREY)
f_body  = Font(name="Calibri", size=10, color=DARK)
f_bold  = Font(name="Calibri", size=10, bold=True, color=DARK)
f_label = Font(name="Calibri", size=9, bold=True, color=GREY)
f_h2    = Font(name="Calibri", size=13, bold=True, color=DARK)
left    = Alignment(horizontal="left", vertical="center")
lefttop = Alignment(horizontal="left", vertical="top", wrap_text=True)
right   = Alignment(horizontal="right", vertical="center")
center  = Alignment(horizontal="center", vertical="center")

# --- Kopf: Marke (links) + Sachbearbeiter (rechts)
of.merge_cells("A1:D1"); set_cell(of, "A1", "MOBIL IN TIME", f_brand)
of.row_dimensions[1].height = 30
of.merge_cells("A2:D2"); set_cell(of, "A2", f"={SD('Firmenname')}", f_body, align=left)
of.merge_cells("A3:D3"); set_cell(of, "A3", f"={SD('Strasse')}&\"  ·  \"&{SD('PLZ / Ort')}", f_small, align=left)
of.merge_cells("A4:D4"); set_cell(of, "A4", f"=\"Tel. \"&{SD('Telefon')}&\"  ·  \"&{SD('E-Mail')}&\"  ·  \"&{SD('Web')}", f_small, align=left)

of.merge_cells("F1:G1"); set_cell(of, "F1", "IHR KONTAKT", f_label, align=right)
of.merge_cells("F2:G2"); set_cell(of, "F2", f"={SD('Name')}", f_bold, align=right)
of.merge_cells("F3:G3"); set_cell(of, "F3", f"=\"Tel. \"&Stammdaten!$B$13", f_small, align=right)
of.merge_cells("F4:G4"); set_cell(of, "F4", f"=Stammdaten!$B$14", f_small, align=right)

# --- Kunde (links) + Offerte-Meta (rechts) ab Zeile 7
of.merge_cells("A7:C7"); set_cell(of, "A7", "KUNDE", f_label, align=left)
of.merge_cells("A8:C8"); set_cell(of, "A8", "Herr Reto Gloor", f_bold, align=left)
of.merge_cells("A9:C9"); set_cell(of, "A9", "EK AG", f_body, align=left)
of.merge_cells("A10:C10");set_cell(of, "A10", "Tel. +41 62 767 80 62", f_small, align=left)
of.merge_cells("A11:C11");set_cell(of, "A11", "reto.gloor@ekag.ch", f_small, align=left)

set_cell(of, "F7", "OFFERTE", f_h2, align=right); of.merge_cells("F7:G7")
meta = [
    ("Offerte-Nr.", "Q663790202605072109", None),
    ("Datum",        _dt.date.today(), DATE_FMT),
    ("Mietbeginn",   _dt.date(2027, 2, 9), DATE_FMT),
    ("Mietende",     _dt.date(2027, 3, 1), DATE_FMT),
    ("Mietzeitraum", "=G11-G10+1&\" Tage\"", None),
    ("Mindestmiete", "7 Tage", None),
]
for i, (lab, val, fmt) in enumerate(meta):
    r = 8 + i
    set_cell(of, f"F{r}", lab, f_label, align=right)
    set_cell(of, f"G{r}", val, f_body, align=right, fmt=fmt)

# --- Anrede + Einleitung
set_cell(of, "A14", '="Sehr geehrter "&LEFT(A8,FIND(" ",A8)-1)&" "&MID(A8,FIND(" ",A8,FIND(" ",A8)+1)+1,99)', f_body, align=left)
of.merge_cells("A14:G14")
of.merge_cells("A15:G17")
set_cell(of, "A15", f'=SUBSTITUTE({SD("Einleitung")},"{{AGB}}",{SD("AGB-Hinweis (URL)")})', f_body, align=lefttop)
of.row_dimensions[15].height = 16
of.row_dimensions[16].height = 16
of.row_dimensions[17].height = 16

set_cell(of, "A19", "Preisaufstellung – Preise reflektieren Mengen", f_bold, align=left)
of.merge_cells("A19:G19")

# --- Tabellenkopf (Zeile 21)
HEAD_ROW = 21
headers = ["Pos", "Anz", "Beschreibung", "Einh.", "Preis/Wo.", "Wochen", "Position"]
aligns  = [center, center, left, center, right, center, right]
for j, (h, al) in enumerate(zip(headers, aligns), start=1):
    c = of.cell(row=HEAD_ROW, column=j, value=h)
    c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    c.fill = PatternFill("solid", fgColor=HEADFILL)
    c.alignment = al
    c.border = border(top=True, bottom=True, color=DARK)
of.row_dimensions[HEAD_ROW].height = 20

# --- Datenzeilen
first = HEAD_ROW + 1
total_rows = len(POSITIONS) + N_BLANK
last = first + total_rows - 1
for idx in range(total_rows):
    r = first + idx
    zebra = PatternFill("solid", fgColor=ZEBRA) if idx % 2 else None
    # Pos-Nummer (nur wenn Beschreibung vorhanden)
    set_cell(of, f"A{r}", f'=IF(C{r}="","",COUNTA($C${first}:C{r}))',
             f_body, zebra, center, bd=border(bottom=True))
    if idx < len(POSITIONS):
        anz, desc, unit, price, wochen = POSITIONS[idx]
        set_cell(of, f"B{r}", anz, f_body, zebra, center, bd=border(bottom=True))
        set_cell(of, f"C{r}", desc, f_body, zebra, left, bd=border(bottom=True))
        set_cell(of, f"D{r}", unit, f_body, zebra, center, bd=border(bottom=True))
        set_cell(of, f"E{r}", round(price, 6), f_body, zebra, right, CHF_FMT, border(bottom=True))
        set_cell(of, f"F{r}", wochen, f_body, zebra, center, bd=border(bottom=True))
    else:
        # leere Zeilen: Einzelpreis schlägt automatisch aus der Artikelliste vor
        set_cell(of, f"B{r}", None, f_body, zebra, center, bd=border(bottom=True))
        set_cell(of, f"C{r}", None, f_body, zebra, left, bd=border(bottom=True))
        set_cell(of, f"D{r}",
                 f'=IFERROR(VLOOKUP(C{r},{ARTIKEL_RANGE},2,FALSE),"")',
                 f_body, zebra, center, bd=border(bottom=True))
        set_cell(of, f"E{r}",
                 f'=IFERROR(VLOOKUP(C{r},{ARTIKEL_RANGE},3,FALSE),"")',
                 f_body, zebra, right, CHF_FMT, border(bottom=True))
        set_cell(of, f"F{r}", f'=IF(C{r}="","",1)', f_body, zebra, center, bd=border(bottom=True))
    # Position = Anz × Preis × Wochen
    set_cell(of, f"G{r}",
             f'=IF($C{r}="","",N($B{r})*N($E{r})*N($F{r}))',
             f_body, zebra, right, CHF_FMT, border(bottom=True))
    # Hilfsspalte I: Listenpreis aus Artikel
    set_cell(of, f"I{r}",
             f'=IFERROR(VLOOKUP(C{r},{ARTIKEL_RANGE},3,FALSE),"")',
             f_small, align=right, fmt=CHF_FMT)

# --- Summenblock
sub_r  = last + 1
rab_r  = last + 2
net_r  = last + 3
mwst_r = last + 4
end_r  = last + 5

def total_row(r, label, formula, bold=False, fill=None, fmt=CHF_FMT, box=False):
    of.merge_cells(f"A{r}:F{r}")
    lab = set_cell(of, f"A{r}", label, f_bold if bold else f_body, align=right)
    val = of.cell(row=r, column=7, value=formula)
    val.font = Font(name="Calibri", size=11 if bold else 10, bold=bold,
                    color=DARK)
    val.alignment = right
    val.number_format = fmt
    if fill:
        for col in range(1, 8):
            of.cell(row=r, column=col).fill = PatternFill("solid", fgColor=fill)
    if box:
        for col in range(1, 8):
            of.cell(row=r, column=col).border = border(top=True, bottom=True, color=ORANGE)

sum_formula = f"=SUM(G{first}:G{last})"
total_row(sub_r,  "Zwischentotal (GESAMT)", sum_formula, bold=True)
of.merge_cells(f"A{rab_r}:F{rab_r}")
set_cell(of, f"A{rab_r}", f'="Rabatt  "&TEXT({SD("Rabatt (%)")},"0.0%")', f_body, align=right)
set_cell(of, f"G{rab_r}", f"=-ROUND(G{sub_r}*{SD('Rabatt (%)')},2)", f_body, align=right, fmt=CHF_FMT)
total_row(net_r,  "Total nach Rabatt", f"=G{sub_r}+G{rab_r}", bold=True)
of.merge_cells(f"A{mwst_r}:F{mwst_r}")
set_cell(of, f"A{mwst_r}", f'="MwSt  "&TEXT({SD("MwSt (%)")},"0.0%")', f_body, align=right)
set_cell(of, f"G{mwst_r}", f"=ROUND(G{net_r}*{SD('MwSt (%)')},2)", f_body, align=right, fmt=CHF_FMT)
total_row(end_r,  "Endbetrag (CHF)", f"=G{net_r}+G{mwst_r}", bold=True,
          fill=TOTFILL, box=True)
of.row_dimensions[end_r].height = 22

# --- Anmerkung + Gruss + Footer
anm_r = end_r + 2
of.merge_cells(f"A{anm_r}:G{anm_r+2}")
set_cell(of, f"A{anm_r}", f'="Anmerkung: "&{SD("Anmerkung")}',
         Font(name="Calibri", size=8.5, italic=True, color=GREY), align=lefttop)

gr_r = anm_r + 4
set_cell(of, f"A{gr_r}", f"={SD('Grussformel')}", f_body, align=left)
of.merge_cells(f"A{gr_r}:G{gr_r}")
set_cell(of, f"A{gr_r+1}", f"={SD('Name')}", f_bold, align=left)
of.merge_cells(f"A{gr_r+1}:G{gr_r+1}")

foot_r = gr_r + 3
for col in range(1, 8):
    of.cell(row=foot_r, column=col).border = border(top=True, color=LINE)
of.merge_cells(f"A{foot_r}:G{foot_r}")
set_cell(of, f"A{foot_r}",
         f'={SD("Firmenname")}&" · "&{SD("Strasse")}&" · "&{SD("PLZ / Ort")}&" · "&{SD("Telefon")}&" · "&{SD("Web")}&"   |   "&{SD("Zusatz")}',
         Font(name="Calibri", size=8, color=GREY), align=center)

# --- Datenvalidierung: Beschreibung -> Dropdown aus Artikelliste
dv = DataValidation(type="list", formula1=f"={ARTIKEL_DESCS}", allow_blank=True)
dv.error = "Bitte aus der Artikelliste wählen oder eigenen Text eintippen."
dv.errorStyle = "warning"
dv.prompt = "Artikel wählen (oder frei eingeben)"
dv.promptTitle = "Beschreibung"
of.add_data_validation(dv)
dv.add(f"C{first}:C{last}")

# --- Seiteneinrichtung (DIN A4, eine Seite breit) + Druckbereich
of.print_area = f"A1:G{foot_r}"
of.page_setup.orientation = "portrait"
of.page_setup.paperSize = of.PAPERSIZE_A4
of.page_setup.fitToWidth = 1
of.page_setup.fitToHeight = 0
of.sheet_properties.pageSetUpPr.fitToPage = True
of.page_margins.left = 0.5
of.page_margins.right = 0.5
of.page_margins.top = 0.6
of.page_margins.bottom = 0.5
of.print_options.horizontalCentered = True

# Offerte als erstes Blatt
wb.move_sheet("Offerte", -(wb.sheetnames.index("Offerte")))
wb.active = wb.sheetnames.index("Offerte")

import os
out = os.path.join(os.path.dirname(__file__), "Offerten_Master_Vorlage.xlsx")
wb.save(out)
print("Gespeichert:", out)
print("Datenzeilen:", first, "bis", last, "| Summe-Zeile:", sub_r, "| Endbetrag:", end_r)
