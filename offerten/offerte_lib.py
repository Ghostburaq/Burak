# -*- coding: utf-8 -*-
"""
offerte_lib – gemeinsamer Baukasten für Offerten-Excel-Mappen.

Eine Konfiguration (Firma, Sachbearbeiter, Kunde, Positionen, Spalten …) ->
eine fertige, druckbare .xlsx mit drei Blättern (Offerte / Stammdaten /
Artikel). Layout, Marke und Formeln sind identisch zur EKAG-Vorlage; nur die
Daten unterscheiden sich. Wird genutzt von build_template.py und
build_offerte_kabel95.py.
"""
import datetime as _dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- Marke / Farben
ORANGE   = "E2620A"
DARK     = "1F2933"
GREY     = "7B8794"
LINE     = "CBD2D9"
ZEBRA    = "F4F6F8"
HEADFILL = "1F2933"
TOTFILL  = "FCEEE3"

CHF_FMT  = "#'##0.00"
PCT_FMT  = "0.0%"
DATE_FMT = "DD.MM.YYYY"


def _border(bottom=False, top=False, color=LINE):
    s = Side(style="thin", color=color)
    return Border(top=s if top else None, bottom=s if bottom else None)


def _set(ws, ref, value=None, font=None, fill=None, align=None, fmt=None, bd=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    if font:  c.font = font
    if fill:  c.fill = fill
    if align: c.alignment = align
    if fmt:   c.number_format = fmt
    if bd:    c.border = bd
    return c


# Standard-Spaltenköpfe (7 Spalten A..G). Leerer Kopf -> Spalte wird ausgeblendet.
HEADERS_RENTAL   = ["Pos", "Anz", "Beschreibung", "Einh.", "Preis/Wo.", "Wochen", "Position"]
HEADERS_PURCHASE = ["Pos", "Menge", "Beschreibung", "Einheit", "Einzelpreis", "", "Total"]


def build_offerte_workbook(out_path, *, company, sachbearbeiter, defaults, texts,
                           customer, meta, positions, artikel,
                           headers=None, subtitle="Preisaufstellung",
                           brand="MOBIL IN TIME", n_blank=8):
    """Erzeugt die Mappe und speichert sie unter out_path."""
    headers = headers or HEADERS_RENTAL
    factor_hidden = not str(headers[5]).strip()
    wb = Workbook()

    # ===================================================== STAMMDATEN
    sd = wb.active
    sd.title = "Stammdaten"
    sd.sheet_view.showGridLines = False
    sd_rows = [
        ("FIRMA (Absender)", ""),
        ("Firmenname",  company["name"]),
        ("Strasse",     company["strasse"]),
        ("PLZ / Ort",   company["plz"]),
        ("Telefon",     company["tel"]),
        ("E-Mail",      company["mail"]),
        ("Web",         company["web"]),
        ("Zusatz",      company.get("zusatz", "")),
        ("AGB-Hinweis (URL)", company.get("agb", "")),
        ("", ""),
        ("SACHBEARBEITER", ""),
        ("Name",    sachbearbeiter["name"]),
        ("Telefon", sachbearbeiter["tel"]),
        ("E-Mail",  sachbearbeiter["mail"]),
        ("", ""),
        ("STANDARDWERTE", ""),
        ("Rabatt (%)", defaults.get("rabatt", 0)),
        ("MwSt (%)",   defaults.get("mwst", 0)),
        ("Währung",    defaults.get("waehrung", "CHF")),
        ("", ""),
        ("TEXTBAUSTEINE", ""),
        ("Einleitung",  texts["einleitung"]),
        ("Grussformel", texts["gruss"]),
        ("Anmerkung",   texts["anmerkung"]),
    ]
    sd.column_dimensions["A"].width = 22
    sd.column_dimensions["B"].width = 95
    hf = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    hfill = PatternFill("solid", fgColor=ORANGE)
    for i, (k, v) in enumerate(sd_rows, start=1):
        a = sd.cell(row=i, column=1, value=k)
        b = sd.cell(row=i, column=2, value=v)
        if v == "" and k and k.isupper():
            a.font = hf; a.fill = hfill; b.fill = hfill
            sd.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
            a.alignment = Alignment(horizontal="left", vertical="center")
        else:
            a.font = Font(name="Calibri", bold=True, color=DARK)
            b.font = Font(name="Calibri", color=DARK)
            b.alignment = Alignment(wrap_text=True, vertical="top")
        if k in ("Rabatt (%)", "MwSt (%)"):
            b.number_format = PCT_FMT

    def SD(label):
        for i, (k, _) in enumerate(sd_rows, start=1):
            if k == label:
                return f"Stammdaten!$B${i}"
        raise KeyError(label)

    # ===================================================== ARTIKEL
    ar = wb.create_sheet("Artikel")
    ar.sheet_view.showGridLines = False
    ar.column_dimensions["A"].width = 56
    ar.column_dimensions["B"].width = 10
    ar.column_dimensions["C"].width = 18
    for j, h in enumerate(["Beschreibung", "Einheit", "Einzelpreis (CHF)"], start=1):
        c = ar.cell(row=1, column=j, value=h)
        c.font = hf; c.fill = PatternFill("solid", fgColor=DARK)
        c.alignment = Alignment(horizontal="left", vertical="center")
    for i, (desc, unit, price) in enumerate(artikel, start=2):
        ar.cell(row=i, column=1, value=desc).font = Font(name="Calibri", color=DARK)
        ar.cell(row=i, column=2, value=unit).alignment = Alignment(horizontal="center")
        pc = ar.cell(row=i, column=3, value=price)
        pc.number_format = CHF_FMT; pc.alignment = Alignment(horizontal="right")
    ar_last = 1 + len(artikel)
    A_RANGE = f"Artikel!$A$2:$C${ar_last}"
    A_DESCS = f"Artikel!$A$2:$A${ar_last}"
    ar.freeze_panes = "A2"

    # ===================================================== OFFERTE
    of = wb.create_sheet("Offerte")
    of.sheet_view.showGridLines = False
    widths = {"A": 4, "B": 7.5, "C": 40, "D": 9, "E": 14, "F": 7, "G": 15,
              "H": 2, "I": 16}
    for col, w in widths.items():
        of.column_dimensions[col].width = w
    of.column_dimensions["I"].hidden = True
    if factor_hidden:
        of.column_dimensions["F"].hidden = True

    f_brand = Font(name="Calibri", bold=True, size=22, color=ORANGE)
    f_small = Font(name="Calibri", size=9, color=GREY)
    f_body  = Font(name="Calibri", size=10, color=DARK)
    f_bold  = Font(name="Calibri", size=10, bold=True, color=DARK)
    f_label = Font(name="Calibri", size=9, bold=True, color=GREY)
    f_h2    = Font(name="Calibri", size=13, bold=True, color=DARK)
    L  = Alignment(horizontal="left", vertical="center")
    LT = Alignment(horizontal="left", vertical="top", wrap_text=True)
    R  = Alignment(horizontal="right", vertical="center")
    C  = Alignment(horizontal="center", vertical="center")

    # Kopf
    of.merge_cells("A1:D1"); _set(of, "A1", brand, f_brand); of.row_dimensions[1].height = 30
    of.merge_cells("A2:D2"); _set(of, "A2", f"={SD('Firmenname')}", f_body, align=L)
    of.merge_cells("A3:D3"); _set(of, "A3", f"={SD('Strasse')}&\"  ·  \"&{SD('PLZ / Ort')}", f_small, align=L)
    of.merge_cells("A4:D4"); _set(of, "A4", f"=\"Tel. \"&{SD('Telefon')}&\"  ·  \"&{SD('E-Mail')}&\"  ·  \"&{SD('Web')}", f_small, align=L)
    of.merge_cells("F1:G1"); _set(of, "F1", "IHR KONTAKT", f_label, align=R)
    of.merge_cells("F2:G2"); _set(of, "F2", f"={SD('Name')}", f_bold, align=R)
    of.merge_cells("F3:G3"); _set(of, "F3", f"=\"Tel. \"&Stammdaten!$B$13", f_small, align=R)
    of.merge_cells("F4:G4"); _set(of, "F4", "=Stammdaten!$B$14", f_small, align=R)

    # Kunde (max. 4 Zeilen: A8..A11)
    of.merge_cells("A7:C7"); _set(of, "A7", "KUNDE", f_label, align=L)
    for i in range(4):
        r = 8 + i
        of.merge_cells(f"A{r}:C{r}")
        val = customer[i] if i < len(customer) else ""
        font = f_bold if i == 0 else (f_body if i == 1 else f_small)
        _set(of, f"A{r}", val, font, align=L)

    # Offerte-Meta (max. 6 Zeilen: 8..13)
    _set(of, "F7", "OFFERTE", f_h2, align=R); of.merge_cells("F7:G7")
    for i, (lab, val, fmt) in enumerate(meta[:6]):
        r = 8 + i
        _set(of, f"F{r}", lab, f_label, align=R)
        _set(of, f"G{r}", val, f_body, align=R, fmt=fmt)

    # Anrede + Einleitung
    of.merge_cells("A14:G14")
    _set(of, "A14", '="Sehr geehrte"&IF(LEFT(A8,4)="Frau"," Frau ",IF(LEFT(A8,4)="Herr","r Herr "," Damen und Herren"))&IF(OR(LEFT(A8,4)="Frau",LEFT(A8,4)="Herr"),MID(A8,FIND(" ",A8&" ",FIND(" ",A8&" ")+1)+1,99),"")', f_body, align=L)
    of.merge_cells("A15:G17")
    _set(of, "A15", f'=SUBSTITUTE({SD("Einleitung")},"{{AGB}}",{SD("AGB-Hinweis (URL)")})', f_body, align=LT)
    for r in (15, 16, 17):
        of.row_dimensions[r].height = 16
    of.merge_cells("A19:G19"); _set(of, "A19", subtitle, f_bold, align=L)

    # Tabellenkopf (Zeile 21)
    HEAD_ROW = 21
    aligns = [C, C, L, C, R, C, R]
    for j, (h, al) in enumerate(zip(headers, aligns), start=1):
        c = of.cell(row=HEAD_ROW, column=j, value=h)
        c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=HEADFILL)
        c.alignment = al
        c.border = _border(top=True, bottom=True, color=DARK)
    of.row_dimensions[HEAD_ROW].height = 20

    # Datenzeilen
    first = HEAD_ROW + 1
    total_rows = len(positions) + n_blank
    last = first + total_rows - 1
    for idx in range(total_rows):
        r = first + idx
        zebra = PatternFill("solid", fgColor=ZEBRA) if idx % 2 else None
        _set(of, f"A{r}", f'=IF(C{r}="","",COUNTA($C${first}:C{r}))',
             f_body, zebra, C, bd=_border(bottom=True))
        if idx < len(positions):
            anz, desc, unit, price, factor = positions[idx]
            _set(of, f"B{r}", anz, f_body, zebra, C, bd=_border(bottom=True))
            _set(of, f"C{r}", desc, f_body, zebra, L, bd=_border(bottom=True))
            _set(of, f"D{r}", unit, f_body, zebra, C, bd=_border(bottom=True))
            _set(of, f"E{r}", round(price, 6), f_body, zebra, R, CHF_FMT, _border(bottom=True))
            _set(of, f"F{r}", factor, f_body, zebra, C, bd=_border(bottom=True))
        else:
            _set(of, f"B{r}", None, f_body, zebra, C, bd=_border(bottom=True))
            _set(of, f"C{r}", None, f_body, zebra, L, bd=_border(bottom=True))
            _set(of, f"D{r}", f'=IFERROR(VLOOKUP(C{r},{A_RANGE},2,FALSE),"")',
                 f_body, zebra, C, bd=_border(bottom=True))
            _set(of, f"E{r}", f'=IFERROR(VLOOKUP(C{r},{A_RANGE},3,FALSE),"")',
                 f_body, zebra, R, CHF_FMT, _border(bottom=True))
            _set(of, f"F{r}", f'=IF(C{r}="","",1)', f_body, zebra, C, bd=_border(bottom=True))
        _set(of, f"G{r}", f'=IF($C{r}="","",N($B{r})*N($E{r})*N($F{r}))',
             f_body, zebra, R, CHF_FMT, _border(bottom=True))
        _set(of, f"I{r}", f'=IFERROR(VLOOKUP(C{r},{A_RANGE},3,FALSE),"")',
             f_small, align=R, fmt=CHF_FMT)

    # Summenblock
    sub_r, rab_r, net_r, mwst_r, end_r = (last + i for i in range(1, 6))

    def total_row(r, label, formula, bold=False, fill=None, box=False):
        of.merge_cells(f"A{r}:F{r}")
        _set(of, f"A{r}", label, f_bold if bold else f_body, align=R)
        v = of.cell(row=r, column=7, value=formula)
        v.font = Font(name="Calibri", size=11 if bold else 10, bold=bold, color=DARK)
        v.alignment = R; v.number_format = CHF_FMT
        if fill:
            for col in range(1, 8):
                of.cell(row=r, column=col).fill = PatternFill("solid", fgColor=fill)
        if box:
            for col in range(1, 8):
                of.cell(row=r, column=col).border = _border(top=True, bottom=True, color=ORANGE)

    total_row(sub_r, "Zwischentotal (GESAMT)", f"=SUM(G{first}:G{last})", bold=True)
    of.merge_cells(f"A{rab_r}:F{rab_r}")
    _set(of, f"A{rab_r}", f'="Rabatt  "&TEXT({SD("Rabatt (%)")},"0.0%")', f_body, align=R)
    _set(of, f"G{rab_r}", f"=-ROUND(G{sub_r}*{SD('Rabatt (%)')},2)", f_body, align=R, fmt=CHF_FMT)
    total_row(net_r, "Total nach Rabatt", f"=G{sub_r}+G{rab_r}", bold=True)
    of.merge_cells(f"A{mwst_r}:F{mwst_r}")
    _set(of, f"A{mwst_r}", f'="MwSt  "&TEXT({SD("MwSt (%)")},"0.0%")', f_body, align=R)
    _set(of, f"G{mwst_r}", f"=ROUND(G{net_r}*{SD('MwSt (%)')},2)", f_body, align=R, fmt=CHF_FMT)
    total_row(end_r, f"Endbetrag", f"=G{net_r}+G{mwst_r}", bold=True, fill=TOTFILL, box=True)
    of.row_dimensions[end_r].height = 22

    # Anmerkung + Gruss + Footer
    anm_r = end_r + 2
    of.merge_cells(f"A{anm_r}:G{anm_r+2}")
    _set(of, f"A{anm_r}", f'="Anmerkung: "&{SD("Anmerkung")}',
         Font(name="Calibri", size=8.5, italic=True, color=GREY), align=LT)
    gr_r = anm_r + 4
    of.merge_cells(f"A{gr_r}:G{gr_r}")
    _set(of, f"A{gr_r}", f"={SD('Grussformel')}", f_body, align=L)
    of.merge_cells(f"A{gr_r+1}:G{gr_r+1}")
    _set(of, f"A{gr_r+1}", f"={SD('Name')}", f_bold, align=L)
    foot_r = gr_r + 3
    for col in range(1, 8):
        of.cell(row=foot_r, column=col).border = _border(top=True, color=LINE)
    of.merge_cells(f"A{foot_r}:G{foot_r}")
    _set(of, f"A{foot_r}",
         f'={SD("Firmenname")}&" · "&{SD("Strasse")}&" · "&{SD("PLZ / Ort")}&" · "&{SD("Telefon")}&" · "&{SD("Web")}&"   |   "&{SD("Zusatz")}',
         Font(name="Calibri", size=8, color=GREY), align=C)

    # Dropdown Beschreibung
    dv = DataValidation(type="list", formula1=f"={A_DESCS}", allow_blank=True)
    dv.errorStyle = "warning"
    dv.error = "Bitte aus der Artikelliste wählen oder eigenen Text eintippen."
    dv.promptTitle = "Beschreibung"; dv.prompt = "Artikel wählen (oder frei eingeben)"
    of.add_data_validation(dv); dv.add(f"C{first}:C{last}")

    # Seiteneinrichtung A4 + Druckbereich
    of.print_area = f"A1:G{foot_r}"
    of.page_setup.orientation = "portrait"
    of.page_setup.paperSize = of.PAPERSIZE_A4
    of.page_setup.fitToWidth = 1
    of.page_setup.fitToHeight = 0
    of.sheet_properties.pageSetUpPr.fitToPage = True
    of.page_margins.left = 0.5; of.page_margins.right = 0.5
    of.page_margins.top = 0.6;  of.page_margins.bottom = 0.5
    of.print_options.horizontalCentered = True

    wb.move_sheet("Offerte", -(wb.sheetnames.index("Offerte")))
    wb.active = wb.sheetnames.index("Offerte")
    wb.save(out_path)
    return dict(out=out_path, first=first, last=last, end=end_r)
