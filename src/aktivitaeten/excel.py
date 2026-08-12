"""Erzeugt die Excel-Mappe: Dashboard, Aktivitäten-Tabelle, Eingabeblatt.

Bewusst nur drei Blätter. Die Aktivitäten liegen als echte Excel-Tabelle
(ListObject) vor — sortier- und filterbar, mit strukturierten Bezügen. Das
Dashboard rechnet über Formeln auf diese Tabelle, aktualisiert sich also
mit, sobald in der Tabelle etwas geändert wird.
"""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from . import bericht
from .model import (BREITEN, EINGABE_SPALTEN, FELDNAMEN, KATEGORIEN, SICHTBAR,
                    STATUS_AUSWAHL, UEBERSCHRIFTEN, Aktivitaet)

# ---------------------------------------------------------------------------
# Farbwelt (an das bestehende Deck-Design angelehnt)
# ---------------------------------------------------------------------------
NAVY = "FF1B2430"
NAVY_HELL = "FF2C3A4B"
ROT = "FFE2001A"
BLAU = "FF1F6FB2"
HELLGRAU = "FFF4F6F8"
KARTE = "FFEDF1F5"
WEISS = "FFFFFFFF"
RAHMEN = "FFD8DEE4"
GRAU_TEXT = "FF5A6472"

DUENN = Side(style="thin", color=RAHMEN)
GITTER = Border(left=DUENN, right=DUENN, top=DUENN, bottom=DUENN)

GELD_FORMAT = '#,##0;[Red]-#,##0;"—"'
DATUM_FORMAT = "DD.MM.YYYY"
ZEIT_FORMAT = "HH:MM"
STD_FORMAT = '0.0;;"—"'

TABELLE_BLATT = "Aktivitäten"
DASHBOARD_BLATT = "Dashboard"
EINGABE_BLATT = "Eingabe"
TABELLE_NAME = "Aktivitaeten"
LEERE_EINGABEZEILEN = 40
MAX_FORMELZEILE = 2000        # Reserve, damit Formeln neue Zeilen mitnehmen


def _aufhellen(argb: str, anteil: float = 0.82) -> str:
    r, g, b = int(argb[2:4], 16), int(argb[4:6], 16), int(argb[6:8], 16)
    mische = lambda k: int(k + (255 - k) * anteil)
    return f"FF{mische(r):02X}{mische(g):02X}{mische(b):02X}"


KATEGORIE_FUELLUNG = {n: PatternFill("solid", fgColor=_aufhellen(f)) for n, f in KATEGORIEN.items()}
KATEGORIE_SCHRIFT = {n: Font(name="Calibri", size=10, bold=True, color=f)
                     for n, f in KATEGORIEN.items()}

# Spaltenbuchstaben der Aktivitäten-Tabelle, für die Dashboard-Formeln
SPALTE = {feld: get_column_letter(i + 1) for i, feld in enumerate(FELDNAMEN)}
TAB = f"'{TABELLE_BLATT}'"


def _bereich(feld: str) -> str:
    buchstabe = SPALTE[feld]
    return f"{TAB}!${buchstabe}$2:${buchstabe}${MAX_FORMELZEILE}"


# ---------------------------------------------------------------------------
# Blatt 1: Aktivitäten (echte Excel-Tabelle)
# ---------------------------------------------------------------------------

def _blatt_aktivitaeten(mappe: Workbook, eintraege: list[Aktivitaet]) -> None:
    blatt = mappe.active
    blatt.title = TABELLE_BLATT

    for spalte, text in enumerate(UEBERSCHRIFTEN, start=1):
        zelle = blatt.cell(row=1, column=spalte, value=text)
        zelle.font = Font(name="Calibri", size=11, bold=True, color=WEISS)
        zelle.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    blatt.row_dimensions[1].height = 28

    for index, eintrag in enumerate(eintraege):
        zeile = index + 2
        for spalte, feld in enumerate(FELDNAMEN, start=1):
            wert = getattr(eintrag, feld, "")
            if isinstance(wert, datetime):
                wert = wert.replace(tzinfo=None)
            zelle = blatt.cell(row=zeile, column=spalte,
                               value=wert if wert not in (None, "") else None)
            zelle.alignment = Alignment(
                vertical="top",
                wrap_text=feld in ("titel", "notizen", "naechster_schritt", "bedarf", "teilnehmer"))
            zelle.font = Font(name="Calibri", size=10)

            if feld == "nr":
                zelle.alignment = Alignment(horizontal="center", vertical="top")
                zelle.font = Font(name="Calibri", size=10, bold=True, color=GRAU_TEXT)
            elif feld == "datum":
                zelle.number_format = DATUM_FORMAT
            elif feld in ("von", "bis"):
                zelle.number_format = ZEIT_FORMAT
            elif feld == "zeit":
                zelle.alignment = Alignment(horizontal="center", vertical="top")
                zelle.font = Font(name="Calibri", size=10, color=GRAU_TEXT)
            elif feld == "dauer_h":
                zelle.number_format = STD_FORMAT
                zelle.alignment = Alignment(horizontal="center", vertical="top")
            elif feld in ("wert_chf", "potenzial_chf", "forecast_chf"):
                zelle.number_format = GELD_FORMAT
            elif feld == "wahrscheinlichkeit":
                zelle.number_format = '0%;;"—"'
            elif feld == "erfasst_am":
                zelle.number_format = "DD.MM.YYYY HH:MM"
            elif feld == "titel":
                zelle.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
            elif feld == "kategorie":
                zelle.fill = KATEGORIE_FUELLUNG.get(eintrag.kategorie,
                                                    PatternFill("solid", fgColor=HELLGRAU))
                zelle.font = KATEGORIE_SCHRIFT.get(eintrag.kategorie, Font(size=10))
                zelle.alignment = Alignment(horizontal="center", vertical="center")
            elif feld in ("id", "quelle"):
                zelle.font = Font(name="Consolas", size=8, color="FF9AA2AB")
            elif feld in ("meeting_link", "website") and wert:
                zelle.hyperlink = wert if str(wert).startswith("http") else f"https://{wert}"
                zelle.font = Font(name="Calibri", size=9, color=BLAU, underline="single")
        blatt.row_dimensions[zeile].height = 32

    # Echte Excel-Tabelle: Filterknöpfe, Zebrastreifen, strukturierte Bezüge
    letzte_spalte = get_column_letter(len(FELDNAMEN))
    letzte_zeile = max(len(eintraege) + 1, 2)
    tabelle = Table(displayName=TABELLE_NAME, ref=f"A1:{letzte_spalte}{letzte_zeile}")
    tabelle.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False,
        showFirstColumn=False, showLastColumn=False)
    blatt.add_table(tabelle)

    # Breiten; Nebenspalten eingeklappt statt gelöscht
    for index, (breite, sichtbar) in enumerate(zip(BREITEN, SICHTBAR), start=1):
        masse = blatt.column_dimensions[get_column_letter(index)]
        masse.width = breite
        if not sichtbar:
            masse.hidden = True
            masse.outlineLevel = 1
    blatt.sheet_properties.outlinePr.summaryRight = True

    blatt.freeze_panes = "C2"
    blatt.sheet_view.showGridLines = False


# ---------------------------------------------------------------------------
# Blatt 2: Dashboard
# ---------------------------------------------------------------------------

def _titelbalken(blatt: Worksheet, eintraege: list[Aktivitaet]) -> None:
    blatt.merge_cells("A1:N1")
    zelle = blatt.cell(row=1, column=1, value="AKTIVITÄTENTRACKER")
    zelle.font = Font(name="Calibri", size=22, bold=True, color=WEISS)
    zelle.fill = PatternFill("solid", fgColor=NAVY)
    zelle.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    blatt.row_dimensions[1].height = 44

    tage = sorted({a.datum for a in eintraege if a.datum})
    zeitraum = (f"{tage[0]:%d.%m.%Y} – {tage[-1]:%d.%m.%Y}" if tage else "noch keine Daten")
    blatt.merge_cells("A2:N2")
    unter = blatt.cell(row=2, column=1,
                       value=f"Zeitraum {zeitraum}   ·   Stand {datetime.now():%d.%m.%Y %H:%M}"
                             f"   ·   {len(eintraege)} Einträge")
    unter.font = Font(name="Calibri", size=10, color=WEISS)
    unter.fill = PatternFill("solid", fgColor=NAVY_HELL)
    unter.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    blatt.row_dimensions[2].height = 20


def _kpi_karten(blatt: Worksheet, eintraege: list[Aktivitaet]) -> None:
    tage = {a.datum for a in eintraege if a.datum}
    karten = [
        ("Aktivitäten", f"=COUNTA({_bereich('titel')})", "0"),
        ("Tage erfasst", len(tage), "0"),
        ("Stunden", f"=SUM({_bereich('dauer_h')})", "0.0"),
        ("Kundenkontakte", "=" + "+".join(
            f'COUNTIF({_bereich("kategorie")},"{k}")'
            for k in ("Akquise", "Kundentermin", "Beratung")), "0"),
        ("Potenzial CHF", f"=SUM({_bereich('potenzial_chf')})", "#,##0"),
        ("Offene Schritte", f'=COUNTIF({_bereich("naechster_schritt")},"<>")', "0"),
    ]
    for index, (titel, wert, format_) in enumerate(karten):
        spalte = 1 + index * 2                       # je Karte zwei Spalten
        buchstabe = get_column_letter(spalte)
        rechts = get_column_letter(spalte + 1)
        blatt.merge_cells(f"{buchstabe}4:{rechts}4")
        blatt.merge_cells(f"{buchstabe}5:{rechts}5")

        kopf = blatt.cell(row=4, column=spalte, value=titel.upper())
        kopf.font = Font(name="Calibri", size=9, bold=True, color=GRAU_TEXT)
        kopf.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        zahl = blatt.cell(row=5, column=spalte, value=wert)
        zahl.font = Font(name="Calibri", size=20, bold=True,
                         color=ROT if "Potenzial" in titel else NAVY)
        zahl.number_format = format_
        zahl.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        for zeile in (4, 5):
            for s in (spalte, spalte + 1):
                blatt.cell(row=zeile, column=s).fill = PatternFill("solid", fgColor=KARTE)
    blatt.row_dimensions[4].height = 16
    blatt.row_dimensions[5].height = 30


def _abschnitt(blatt: Worksheet, zeile: int, spalte: int, breite: int, text: str) -> None:
    blatt.merge_cells(start_row=zeile, start_column=spalte,
                      end_row=zeile, end_column=spalte + breite - 1)
    zelle = blatt.cell(row=zeile, column=spalte, value=text.upper())
    zelle.font = Font(name="Calibri", size=11, bold=True, color=WEISS)
    zelle.fill = PatternFill("solid", fgColor=NAVY)
    zelle.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for s in range(spalte, spalte + breite):
        blatt.cell(row=zeile, column=s).fill = PatternFill("solid", fgColor=NAVY)
    blatt.row_dimensions[zeile].height = 22


def _minitabelle(blatt: Worksheet, zeile: int, spalte: int, kopf: list[str],
                 zeilen: list[list], formate: dict[int, str] | None = None) -> int:
    """Kleiner Datenblock; gibt die erste freie Zeile darunter zurück."""
    for versatz, text in enumerate(kopf):
        zelle = blatt.cell(row=zeile, column=spalte + versatz, value=text)
        zelle.font = Font(name="Calibri", size=9, bold=True, color=GRAU_TEXT)
        zelle.border = Border(bottom=Side(style="medium", color=NAVY))
        zelle.alignment = Alignment(horizontal="left", vertical="center")
    for index, werte in enumerate(zeilen):
        aktuell = zeile + 1 + index
        for versatz, wert in enumerate(werte):
            zelle = blatt.cell(row=aktuell, column=spalte + versatz, value=wert)
            zelle.font = Font(name="Calibri", size=10)
            zelle.border = Border(bottom=DUENN)
            zelle.alignment = Alignment(vertical="center")
            if formate and versatz in formate:
                zelle.number_format = formate[versatz]
        blatt.row_dimensions[aktuell].height = 18
    return zeile + len(zeilen) + 2


def _blatt_dashboard(mappe: Workbook, eintraege: list[Aktivitaet], quellen: dict) -> None:
    blatt = mappe.create_sheet(DASHBOARD_BLATT, index=0)
    blatt.sheet_view.showGridLines = False
    for spalte, breite in zip("ABCDEFGHIJKLMN",
                              [13, 11, 13, 11, 13, 11, 4, 13, 13, 13, 13, 13, 13, 13]):
        blatt.column_dimensions[spalte].width = breite

    _titelbalken(blatt, eintraege)
    _kpi_karten(blatt, eintraege)

    # -- Auswertungsblöcke als Datengrundlage der Diagramme ----------------
    # Die Werte sind Formeln auf die Aktivitäten-Tabelle: ändert sich dort
    # etwas, rechnen Block und Diagramm sofort mit.
    _abschnitt(blatt, 7, 1, 5, "Stunden und Aktivitäten je Kategorie")
    kategorien = [k for k in KATEGORIEN if any(a.kategorie == k for a in eintraege)]
    kategorie_kopf = 8
    kategorie_erste = kategorie_kopf + 1
    kategorie_zeilen = [
        [name,
         f'=SUMIF({_bereich("kategorie")},$A{kategorie_erste + i},{_bereich("dauer_h")})',
         f'=COUNTIF({_bereich("kategorie")},$A{kategorie_erste + i})']
        for i, name in enumerate(kategorien)
    ]
    naechste = _minitabelle(blatt, kategorie_kopf, 1, ["Kategorie", "Stunden", "Anzahl"],
                            kategorie_zeilen, {1: STD_FORMAT, 2: "0"})
    kategorie_letzte = kategorie_erste + len(kategorien) - 1

    tage = bericht.nach_tagen(eintraege)
    _abschnitt(blatt, naechste, 1, 5, "Aktivitäten je Tag")
    tag_kopf = naechste + 1
    tag_erste = tag_kopf + 1
    tag_zeilen = [
        [tag,
         f'=COUNTIF({_bereich("datum")},$A{tag_erste + i})',
         f'=SUMIF({_bereich("datum")},$A{tag_erste + i},{_bereich("dauer_h")})']
        for i, (tag, _) in enumerate(tage)
    ]
    naechste = _minitabelle(blatt, tag_kopf, 1, ["Datum", "Anzahl", "Stunden"],
                            tag_zeilen, {0: DATUM_FORMAT, 1: "0", 2: STD_FORMAT})
    tag_letzte = tag_erste + len(tage) - 1

    # -- Diagramme ---------------------------------------------------------
    def _diagramm(art: str, titel: str, spalte_werte: int, erste: int, letzte: int,
                  farbe: str, anker: str) -> None:
        if letzte < erste:
            return
        figur = BarChart()
        figur.type = art
        figur.title = titel
        figur.legend = None
        figur.height, figur.width = 7.5, 12
        figur.gapWidth = 60
        figur.add_data(Reference(blatt, min_col=spalte_werte, min_row=erste, max_row=letzte),
                       titles_from_data=False)
        figur.set_categories(Reference(blatt, min_col=1, min_row=erste, max_row=letzte))
        figur.series[0].graphicalProperties = GraphicalProperties(solidFill=farbe[2:])
        blatt.add_chart(figur, anker)

    # nebeneinander rechts neben den Datenblöcken — darunter bleibt alles frei
    _diagramm("bar", "Stunden je Kategorie", 2, kategorie_erste, kategorie_letzte, BLAU, "H7")
    _diagramm("col", "Aktivitäten je Tag", 2, tag_erste, tag_letzte, ROT, "P7")

    # -- Zusammenfassung in Worten ----------------------------------------
    # Die Diagramme belegen rechts die Zeilen 7 bis 22; der Fliesstext beginnt
    # darunter, damit nichts überlappt und keine leere Fläche entsteht.
    zeile = max(naechste, 24)
    _abschnitt(blatt, zeile, 1, 14, "Zusammenfassung")
    zeile += 1
    blatt.merge_cells(start_row=zeile, start_column=1, end_row=zeile, end_column=14)
    gesamt = blatt.cell(row=zeile, column=1, value=bericht.gesamttext(eintraege))
    gesamt.font = Font(name="Calibri", size=11, color=NAVY)
    gesamt.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
    gesamt.fill = PatternFill("solid", fgColor=KARTE)
    for s in range(1, 15):
        blatt.cell(row=zeile, column=s).fill = PatternFill("solid", fgColor=KARTE)
    blatt.row_dimensions[zeile].height = 46
    zeile += 2

    _abschnitt(blatt, zeile, 1, 14, "Tagebuch — was an welchem Tag lief")
    zeile += 1
    for tag, gruppe in tage:
        stunden = sum(a.dauer_h or 0 for a in gruppe)
        kopf = blatt.cell(row=zeile, column=1,
                          value=f"{tag:%d.%m.%Y} · {gruppe[0].wochentag}")
        kopf.font = Font(name="Calibri", size=11, bold=True, color=WEISS)
        blatt.merge_cells(start_row=zeile, start_column=1, end_row=zeile, end_column=3)
        rechts = blatt.cell(row=zeile, column=4,
                            value=f"{len(gruppe)} Aktivitäten · "
                                  f"{stunden:.1f} h".replace(".", ","))
        rechts.font = Font(name="Calibri", size=10, color=WEISS)
        blatt.merge_cells(start_row=zeile, start_column=4, end_row=zeile, end_column=14)
        for s in range(1, 15):
            blatt.cell(row=zeile, column=s).fill = PatternFill("solid", fgColor=NAVY_HELL)
        for zelle in (kopf, rechts):
            zelle.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        blatt.row_dimensions[zeile].height = 20
        zeile += 1

        blatt.merge_cells(start_row=zeile, start_column=1, end_row=zeile, end_column=14)
        text = blatt.cell(row=zeile, column=1, value=bericht.tagestext(tag, gruppe))
        text.font = Font(name="Calibri", size=10, color=NAVY)
        text.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        blatt.row_dimensions[zeile].height = max(34, 15 * (1 + len(
            bericht.tagestext(tag, gruppe)) // 150))
        zeile += 1

        for eintrag in gruppe:
            zeitangabe = eintrag.zeit or "—"
            zusatz = bericht.chf(eintrag.potenzial_chf or eintrag.wert_chf)
            # Positionsnummer wie in der Tabelle — damit die Zeile dort
            # sofort wiederzufinden ist
            zeile_text = (f"   Pos. {eintrag.nr:>3}   {zeitangabe}   {eintrag.titel}"
                          + (f"   ·   {eintrag.firma}" if eintrag.firma else "")
                          + (f"   ·   Potenzial {zusatz}" if zusatz else "")
                          + (f"   ·   nächster Schritt: {eintrag.naechster_schritt}"
                             if eintrag.naechster_schritt else ""))
            blatt.merge_cells(start_row=zeile, start_column=1, end_row=zeile, end_column=14)
            punkt = blatt.cell(row=zeile, column=1, value=zeile_text)
            punkt.font = Font(name="Calibri", size=10, color=GRAU_TEXT)
            punkt.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
            farbe = KATEGORIE_FUELLUNG.get(eintrag.kategorie)
            if farbe:
                blatt.cell(row=zeile, column=1).border = Border(
                    left=Side(style="thick", color=KATEGORIEN[eintrag.kategorie]))
            blatt.row_dimensions[zeile].height = 17
            zeile += 1
        zeile += 1

    # -- Pipeline und offene Punkte ---------------------------------------
    pipeline = sorted([a for a in eintraege if (a.potenzial_chf or a.wert_chf)],
                      key=lambda a: -(a.potenzial_chf or a.wert_chf or 0))
    if pipeline:
        _abschnitt(blatt, zeile, 1, 7, "Pipeline")
        zeile = _minitabelle(
            blatt, zeile + 1, 1,
            ["Pos.", "Firma", "Potenzial CHF", "Status", "Nächster Schritt", "Kontakt", "Datum"],
            [[a.nr, a.firma or a.titel, a.potenzial_chf or a.wert_chf, a.status,
              a.naechster_schritt, a.kontakt or a.email, a.datum] for a in pipeline],
            {2: GELD_FORMAT, 6: DATUM_FORMAT})

    offen = [a for a in eintraege if a.naechster_schritt or a.follow_up]
    if offen:
        _abschnitt(blatt, zeile, 1, 7, "Offene nächste Schritte")
        zeile = _minitabelle(
            blatt, zeile + 1, 1,
            ["Pos.", "Firma / Vorgang", "Was", "Bis wann", "Kontakt", "Kategorie", "Datum"],
            [[a.nr, a.firma or a.titel, a.naechster_schritt or f"Follow-up: {a.follow_up}",
              a.follow_up, a.kontakt or a.email, a.kategorie, a.datum] for a in offen],
            {6: DATUM_FORMAT})

    # -- Herkunft ----------------------------------------------------------
    doppelt = sum(max(m.get("importe", 1) - 1, 0) for m in quellen.values())
    _abschnitt(blatt, zeile, 1, 6, "Datenherkunft")
    zeile += 1
    blatt.merge_cells(start_row=zeile, start_column=1, end_row=zeile, end_column=14)
    quelle_text = (f"{len(quellen)} Quelldateien eingelesen "
                   f"({', '.join(sorted({m.get('typ', '?') for m in quellen.values()}))})"
                   f" · {doppelt} doppelte Einwürfe erkannt und übersprungen"
                   f" · manuelle Korrekturen in der Tabelle bleiben bei jedem Lauf erhalten.")
    hinweis = blatt.cell(row=zeile, column=1, value=quelle_text)
    hinweis.font = Font(name="Calibri", size=9, color=GRAU_TEXT)
    hinweis.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    blatt.row_dimensions[zeile].height = 28


# ---------------------------------------------------------------------------
# Blatt 3: Eingabe
# ---------------------------------------------------------------------------

def _blatt_eingabe(mappe: Workbook) -> None:
    blatt = mappe.create_sheet(EINGABE_BLATT)
    feldnamen = [s[0] for s in EINGABE_SPALTEN]
    kopf = [s[1] for s in EINGABE_SPALTEN]
    letzte = get_column_letter(len(kopf))

    blatt.merge_cells(f"A1:{letzte}1")
    hinweis = blatt.cell(row=1, column=1, value=(
        "Hier Aktivitäten von Hand eintippen — eine Zeile pro Vorgang. "
        "Pflicht sind nur Datum und Titel. Beim nächsten Lauf von "
        "«python3 src/aktivitaeten/cli.py» wandern die Zeilen ins Blatt "
        "«Aktivitäten» und dieses Blatt ist wieder leer. ID-Spalte leer lassen."))
    hinweis.font = Font(name="Calibri", size=10, bold=True, color=WEISS)
    hinweis.fill = PatternFill("solid", fgColor=ROT)
    hinweis.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    blatt.row_dimensions[1].height = 30

    for spalte, (text, breite) in enumerate(zip(kopf, [s[2] for s in EINGABE_SPALTEN]), start=1):
        zelle = blatt.cell(row=2, column=spalte, value=text)
        zelle.font = Font(name="Calibri", size=11, bold=True, color=WEISS)
        zelle.fill = PatternFill("solid", fgColor=NAVY)
        zelle.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        zelle.border = GITTER
        blatt.column_dimensions[get_column_letter(spalte)].width = breite
    blatt.row_dimensions[2].height = 28

    erste, letzte_zeile = 3, 3 + LEERE_EINGABEZEILEN - 1
    kategorie_auswahl = DataValidation(
        type="list", allow_blank=True, formula1='"' + ",".join(KATEGORIEN) + '"',
        prompt="Leer lassen genügt — dann ordnet das Tool selbst zu.", promptTitle="Kategorie")
    status_auswahl = DataValidation(
        type="list", allow_blank=True, formula1='"' + ",".join(STATUS_AUSWAHL) + '"',
        prompt="Auswählen oder frei eintippen", promptTitle="Status")
    blatt.add_data_validation(kategorie_auswahl)
    blatt.add_data_validation(status_auswahl)

    for zeile in range(erste, letzte_zeile + 1):
        for spalte, feld in enumerate(feldnamen, start=1):
            zelle = blatt.cell(row=zeile, column=spalte)
            zelle.border = GITTER
            zelle.font = Font(name="Calibri", size=10)
            zelle.alignment = Alignment(vertical="top",
                                        wrap_text=feld in ("titel", "notizen",
                                                           "naechster_schritt"))
            if zeile % 2 == 0:
                zelle.fill = PatternFill("solid", fgColor=HELLGRAU)
            if feld == "datum":
                zelle.number_format = DATUM_FORMAT
            elif feld in ("von", "bis"):
                zelle.number_format = ZEIT_FORMAT
                zelle.alignment = Alignment(horizontal="center", vertical="top")
            elif feld == "potenzial_chf":
                zelle.number_format = GELD_FORMAT
            elif feld == "id":
                zelle.font = Font(name="Consolas", size=8, color="FFB0B6BD")
        blatt.row_dimensions[zeile].height = 22

    kategorie_auswahl.add(f"{get_column_letter(feldnamen.index('kategorie') + 1)}{erste}:"
                          f"{get_column_letter(feldnamen.index('kategorie') + 1)}{letzte_zeile}")
    status_auswahl.add(f"{get_column_letter(feldnamen.index('status') + 1)}{erste}:"
                       f"{get_column_letter(feldnamen.index('status') + 1)}{letzte_zeile}")

    blatt.freeze_panes = "B3"
    blatt.sheet_view.showGridLines = False


# ---------------------------------------------------------------------------

def schreiben(ziel: Path, eintraege: list[Aktivitaet], quellen: dict[str, dict]) -> Path:
    mappe = Workbook()
    mappe.properties.title = "Aktivitätentracker"
    mappe.properties.creator = "Aktivitaeten-Pipeline"

    _blatt_aktivitaeten(mappe, eintraege)
    _blatt_dashboard(mappe, eintraege, quellen)   # wird als erstes Blatt eingehängt
    _blatt_eingabe(mappe)
    mappe.active = 0                              # Dashboard beim Öffnen zeigen

    ziel.parent.mkdir(parents=True, exist_ok=True)
    mappe.save(ziel)
    return ziel
