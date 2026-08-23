#!/usr/bin/env python3
"""Gestaltungssystem des kabuu-Berichts: Farben, Masse, Typografie, XML-Werkzeug.

Alle Gestaltungsentscheide stehen hier an einer Stelle. ``build_report.py``
setzt den Bericht ausschliesslich aus diesen Bausteinen zusammen, damit
Abstände, Linien und Schriftgrade im ganzen Dokument identisch sind — und
sich an einer Stelle ändern lassen.
"""

from __future__ import annotations

import re

from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

# ---------------------------------------------------------------- Farbwelt --
# Die Marke ist monochrom: Graphit und Stahl. Farbe trägt ausschliesslich
# Bedeutung — bestätigt, Vorbehalt, intern — und wird sonst nicht eingesetzt.
INK = "1B2430"  # Überschriften, Tabellenkopf
BODY = "23282E"  # Fliesstext
STEEL = "6B7480"  # Beschriftungen, Sekundärtext
STEEL_LT = "97A0AB"  # Kapitelziffern, Wasserzeichen
HAIR = "D8DDE3"  # Haarlinien in Tabellen
RULE = "B6BEC7"  # kräftigere Trennlinien
BAND = "F5F7F9"  # Zebrastreifen
PANEL = "F0F3F6"  # Formel- und Infoflächen

ACCENT = {
    "success": {"line": "2E7D32", "tint": "F1F7F1"},
    "warning": {"line": "B9770E", "tint": "FDF6E9"},
    "danger": {"line": "C0392B", "tint": "FDF0EE"},
    "info": {"line": "2C4A6B", "tint": "F1F4F8"},
}

# ------------------------------------------------------------------ Masse --
# A4 mit grosszügigen, aber nicht verspielten Rändern. Der Satzspiegel ist
# 16,6 cm breit; jede Tabelle rechnet gegen genau diesen Wert.
PAGE_W, PAGE_H = 11906, 16838  # Twips (A4 hoch)
MARGIN_X, MARGIN_TOP, MARGIN_BOT = 1247, 1418, 1191
CONTENT_W = PAGE_W - 2 * MARGIN_X  # 9412 Twips = 16,6 cm

FONT = "Arial"
MONO = "Consolas"
LANG = "de-CH"

# Schriftgrade in Halbpunkten (Word-Einheit für w:sz)
SZ_BODY = 20  # 10,0 pt
SZ_SMALL = 18  # 9,0 pt
SZ_TABLE = 18
SZ_TABLE_HEAD = 17
SZ_CAPTION = 17
SZ_H1 = 34  # 17,0 pt
SZ_H2 = 25  # 12,5 pt
SZ_H3 = 21  # 10,5 pt
SZ_MONO = 17

# Abstände in Twentieths of a Point (20 = 1 pt)
SP_BODY_AFTER = 120  # 6 pt
LINE_BODY = 288  # 1,20-facher Zeilenabstand
LINE_TIGHT = 264


# ------------------------------------------------------- XML-Kleinwerkzeug --
def el(tag: str, **attrs) -> OxmlElement:
    """Ein w:-Element mit w:-Attributen, kurz geschrieben."""
    node = OxmlElement(tag if ":" in tag else "w:" + tag)
    for key, value in attrs.items():
        node.set(qn("w:" + key), str(value))
    return node


def _sub(parent, tag: str, **attrs) -> OxmlElement:
    node = el(tag, **attrs)
    parent.append(node)
    return node


def _pPr(p):
    return p._p.get_or_add_pPr()


def _first(parent, tag: str):
    """Vorhandenes Kind zurückgeben oder neu anlegen — Reihenfolge egal."""
    found = parent.find(qn("w:" + tag))
    if found is None:
        found = _sub(parent, tag)
    return found


# --------------------------------------------------------------- Textläufe --
def run(p, text: str, *, size=SZ_BODY, color=BODY, bold=False, italic=False,
        font=FONT, caps=False, spacing=0, underline=False):
    """Textlauf mit den Auszeichnungen, die dieses Dokument kennt."""
    r = p.add_run()
    rpr = r._r.get_or_add_rPr()
    _sub(rpr, "rFonts", ascii=font, hAnsi=font, cs=font, eastAsia=font)
    if bold:
        _sub(rpr, "b")
        _sub(rpr, "bCs")
    if italic:
        _sub(rpr, "i")
    if underline:
        _sub(rpr, "u", val="single")
    if caps:
        _sub(rpr, "caps")
    if spacing:
        _sub(rpr, "spacing", val=spacing)
    _sub(rpr, "color", val=color)
    _sub(rpr, "sz", val=size)
    _sub(rpr, "szCs", val=size)
    _sub(rpr, "lang", val=LANG)
    # \n und \t sind in OOXML eigene Elemente, kein Text — sonst erscheinen sie
    # als Leerzeichen und jeder Tabulatorstopp bleibt wirkungslos.
    for i, line in enumerate(text.split("\n")):
        if i:
            r._r.append(el("br"))
        for j, chunk in enumerate(line.split("\t")):
            if j:
                r._r.append(el("tab"))
            if not chunk:
                continue
            node = el("t")
            node.set(qn("xml:space"), "preserve")
            node.text = chunk
            r._r.append(node)
    return r


def para(container, *, space_before=0, space_after=SP_BODY_AFTER, line=LINE_BODY,
         align=None, keep_next=False, keep_lines=True, page_break=False,
         indent_left=0, indent_right=0, hanging=0, contextual=False, style=None,
         hyphenate=True):
    """Absatz mit vollständig gesetztem Abstands- und Umbruchverhalten.

    ``keep_lines`` ist Vorgabe: Absätze sollen nicht mit einer Restzeile auf
    der Folgeseite landen. ``keep_next`` bindet einen Absatz an den nächsten,
    etwa Überschrift an Text oder Bild an Legende. ``hyphenate=False`` gilt
    für kurze, ausgezeichnete Zeilen — Überschriften, Legenden, Verzeichnis —
    in denen eine Worttrennung nur stört.
    """
    p = container.add_paragraph(style=style)
    pr = _pPr(p)
    if page_break:
        _sub(pr, "pageBreakBefore")
    if keep_next:
        _sub(pr, "keepNext")
    if keep_lines:
        _sub(pr, "keepLines")
    _sub(pr, "widowControl", val="true")
    if not hyphenate:
        _sub(pr, "suppressAutoHyphens", val="true")
    if contextual:
        _sub(pr, "contextualSpacing")
    if indent_left or indent_right or hanging:
        attrs = {}
        if indent_left:
            attrs["left"] = indent_left
        if indent_right:
            attrs["right"] = indent_right
        if hanging:
            attrs["hanging"] = hanging
        _sub(pr, "ind", **attrs)
    _sub(pr, "spacing", before=space_before, after=space_after, line=line, lineRule="auto")
    if align is not None:
        p.alignment = align
    return p


def borders(p, *, top=None, bottom=None, left=None, right=None):
    """Absatzrahmen. Werte als ``(Farbe, Stärke in Achtelpunkten, Abstand)``."""
    pr = _pPr(p)
    bdr = _sub(pr, "pBdr")
    for name, spec in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        if spec:
            color, sz, space = spec
            _sub(bdr, name, val="single", color=color, sz=sz, space=space)
    return p


def shade(p, fill: str):
    _sub(_pPr(p), "shd", val="clear", color="auto", fill=fill)
    return p


def tabs(p, stops: list[tuple[int, str, str]]):
    """Tabulatoren als ``(Position, Ausrichtung, Füllzeichen)``."""
    node = _sub(_pPr(p), "tabs")
    for pos, align, leader in stops:
        _sub(node, "tab", val=align, pos=pos, leader=leader)
    return p


# --------------------------------------------------------------- Tabellen --
def table_frame(table, widths: list[int]):
    """Feste Spaltenbreiten — sonst rechnet Word die Tabelle selbst um."""
    tbl = table._tbl
    pr = tbl.tblPr
    for tag in ("tblW", "tblLayout", "tblBorders", "tblCellMar", "tblInd"):
        old = pr.find(qn("w:" + tag))
        if old is not None:
            pr.remove(old)
    _sub(pr, "tblW", type="dxa", w=sum(widths))
    _sub(pr, "tblLayout", type="fixed")
    _sub(pr, "tblInd", type="dxa", w=0)
    margins = _sub(pr, "tblCellMar")
    _sub(margins, "top", type="dxa", w=85)
    _sub(margins, "bottom", type="dxa", w=85)
    _sub(margins, "left", type="dxa", w=110)
    _sub(margins, "right", type="dxa", w=110)

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        _sub(grid, "gridCol", w=w)
    pr.addnext(grid)
    return table


def cell_style(cell, *, width: int, fill: str | None = None, top=None, bottom=None,
               left=None, right=None, valign="top"):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("tcW", "shd", "tcBorders", "vAlign"):
        old = tcPr.find(qn("w:" + tag))
        if old is not None:
            tcPr.remove(old)
    _sub(tcPr, "tcW", type="dxa", w=width)
    if fill:
        _sub(tcPr, "shd", val="clear", color="auto", fill=fill)
    if any((top, bottom, left, right)):
        bdr = _sub(tcPr, "tcBorders")
        for name, spec in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            if spec:
                color, sz = spec
                _sub(bdr, name, val="single", color=color, sz=sz)
            else:
                _sub(bdr, name, val="none", sz=0)
    _sub(tcPr, "vAlign", val=valign)
    cell.vertical_alignment = {
        "top": WD_ALIGN_VERTICAL.TOP,
        "center": WD_ALIGN_VERTICAL.CENTER,
    }[valign]
    return cell


def cell_margins(cell, *, top=85, bottom=85, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcMar"))
    if old is not None:
        tcPr.remove(old)
    mar = _sub(tcPr, "tcMar")
    _sub(mar, "top", type="dxa", w=top)
    _sub(mar, "bottom", type="dxa", w=bottom)
    _sub(mar, "left", type="dxa", w=left)
    _sub(mar, "right", type="dxa", w=right)
    return cell


def row_rules(row, *, header=False, cant_split=True):
    """Kopfzeilen werden auf jeder Folgeseite wiederholt."""
    trPr = row._tr.get_or_add_trPr()
    if cant_split:
        _sub(trPr, "cantSplit")
    if header:
        _sub(trPr, "tblHeader")
    return row


def clear_cell(cell):
    """Die von python-docx angelegte Leerzeile entfernen."""
    for p in list(cell.paragraphs):
        p._p.getparent().remove(p._p)
    return cell


# ------------------------------------------------------------ Feldbefehle --
def field(p, instruction: str, cached: str = "", *, size=SZ_SMALL, color=STEEL,
          bold=False, font=FONT):
    """Word-Feld mit zwischengespeichertem Ergebnis.

    Das Ergebnis ist das, was Word und jeder PDF-Export anzeigen, bevor die
    Felder aktualisiert werden. Es wird deshalb beim Bauen mit dem richtigen
    Wert gefüllt und nicht leer gelassen.
    """
    def _mk(kind):
        r = p.add_run()
        rpr = r._r.get_or_add_rPr()
        _sub(rpr, "rFonts", ascii=font, hAnsi=font, cs=font)
        if bold:
            _sub(rpr, "b")
        _sub(rpr, "color", val=color)
        _sub(rpr, "sz", val=size)
        _sub(rpr, "szCs", val=size)
        return r

    begin = _mk("begin")
    begin._r.append(el("fldChar", fldCharType="begin"))
    instr = _mk("instr")
    node = el("instrText")
    node.set(qn("xml:space"), "preserve")
    node.text = instruction
    instr._r.append(node)
    sep = _mk("separate")
    sep._r.append(el("fldChar", fldCharType="separate"))
    if cached:
        result = _mk("result")
        t = el("t")
        t.set(qn("xml:space"), "preserve")
        t.text = cached
        result._r.append(t)
    end = _mk("end")
    end._r.append(el("fldChar", fldCharType="end"))
    return p


def bookmark(p, name: str, bid: int):
    p._p.append(el("bookmarkStart", id=bid, name=name))
    p._p.append(el("bookmarkEnd", id=bid))
    return p


def hyperlink_to(p, anchor: str):
    """Absatzinhalt in einen internen Verweis umhängen (für das Verzeichnis)."""
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    for r in list(p._p.findall(qn("w:r"))):
        p._p.remove(r)
        link.append(r)
    p._p.append(link)
    return p


# ------------------------------------------------------------- Typografie --
NBSP = "\u00a0"  # geschütztes Leerzeichen
NBHYPH = "\u2011"  # geschützter Bindestrich

_UNITS = (
    r"kVA|kWh|kW|kHz|kΩ|mΩ|MΩ|µF|µH|mH|mA|ms|min|Hz|VA|°C|Ohm|Grad|Std|"
    r"cm|mm|km|nF|pF|kA|kV|%|V|A|W|s|h|m|Ω"
)
_PREFIX = r"Kapitel|Abbildung|Abschnitt|Bild|Tabelle|Anhang|Seite|Nr\.?|Rang|Typ"
_STANDARD = r"EN|IEC|DIN|VDE|NIN|ISO|SN"


# Vierstellige und längere Zahlen bekommen die Tausendertrennung nur, wenn
# ihnen eine Einheit folgt oder sie die untere Grenze eines Bereichs sind.
# Normnummern (EN 50160), Postleitzahlen, Serien- und Typennummern bleiben
# damit unberührt — dort wäre die Trennung schlicht falsch.
_THOUSANDS = re.compile(
    rf"(?<![\d\u2019.,\-])(\d{{4,}})"
    rf"(?={NBSP}(?:{_UNITS})\b|(?: bis | und )\d{{3,}}{NBSP}(?:{_UNITS})\b)"
)
_STANDARD_BEFORE = re.compile(rf"(?:{_STANDARD})[ {NBSP}]$")


def _swiss_thousands(digits: str) -> str:
    out, rest = digits[-3:], digits[:-3]
    while rest:
        out, rest = rest[-3:] + "\u2019" + out, rest[:-3]
    return out


def typo(text: str, *, insert: bool = True) -> str:
    """Schweizer Zahlensatz und geschützte Abstände.

    Zahl und Einheit, Normbezeichnung und Nummer, Verweiswort und Ziffer
    gehören zusammen und dürfen nicht über den Zeilenumbruch getrennt werden.
    Tausendertrennung als Hochkomma, wie in der Schweiz üblich.

    ``insert=False`` lässt die beiden Regeln weg, die Zeichen hinzufügen —
    die nachträgliche Tausendertrennung und das Multiplikationszeichen. Für
    die Rechenblöcke ist das nötig: dort stehen die Spalten über Leerzeichen
    untereinander, und ein zusätzliches Zeichen verschiebt die ganze Zeile.
    """
    # Tausendertrennung des Originals auf Hochkomma umstellen: 86'600 -> 86\u2019600
    for _ in range(3):
        text = re.sub(r"(?<=\d)\'(?=\d{3}\b)", "\u2019", text)
    # Zahl + Einheit
    text = re.sub(rf"(\d(?:[.,]\d+)?) ({_UNITS})(?![\wäöüÄÖÜß])", rf"\1{NBSP}\2", text)
    # Normbezeichnung + Nummer, Verweiswort + Nummer
    text = re.sub(rf"\b({_STANDARD}) (\d)", rf"\1{NBSP}\2", text)
    text = re.sub(rf"\b({_PREFIX}) (\d+|[A-Z]\b)", rf"\1{NBSP}\2", text)
    # Datum und Uhrzeit einer Zeitangabe zusammenhalten
    text = re.sub(r"(\d{2}\.\d{2}\.(?:\d{4})?) (\d{2}:\d{2})", rf"\1{NBSP}\2", text)
    # Messpunktbezeichner nicht umbrechen
    text = re.sub(r"\bMP-(\d)", rf"MP{NBHYPH}\1", text)
    if not insert:
        return text
    # Multiplikationszeichen an die Zahl binden
    text = re.sub(r"(\d)\s*×\s*", rf"\1{NBSP}× ", text)
    # Tausendertrennung für Messwerte, die sie im Original noch nicht hatten
    text = _THOUSANDS.sub(
        lambda m: m.group(1)
        if _STANDARD_BEFORE.search(text[max(0, m.start() - 8):m.start()])
        else _swiss_thousands(m.group(1)),
        text,
    )
    return text


# "rund 125 A" ist ein Zahlenwert, "rund das 6-Fache" eine Formulierung.
_NUMERIC = re.compile(
    rf"^[−–\-+]?[\d’'.,]+(?:\s*(?:{_UNITS}))?$|^—$|^rund\s+\d|^[+\-]?\d",
    re.IGNORECASE,
)


def looks_numeric(text: str) -> bool:
    text = text.strip()
    if not text:
        return False
    return bool(_NUMERIC.match(text))


# ------------------------------------------------- Schemakonforme Reihenfolge --
# OOXML schreibt die Reihenfolge der Kinder von pPr, rPr und den Tabellen-
# Eigenschaften fest vor. Word verweigert eine Datei, die davon abweicht,
# während LibreOffice sie klaglos öffnet — der Fehler fiele also erst beim
# Kunden auf. Deshalb wird das fertige Dokument einmal durchsortiert, statt
# an jeder Aufrufstelle auf die Reihenfolge zu achten.
_ORDER = {
    "pPr": (
        "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr "
        "suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap "
        "overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
        "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc "
        "textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr "
        "sectPr pPrChange"
    ),
    "rPr": (
        "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss "
        "imprint noProof snapToGrid vanish webHidden color spacing w kern position sz "
        "szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang "
        "eastAsianLayout specVanish oMath rPrChange"
    ),
    "tcPr": (
        "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection "
        "tcFitText vAlign hideMark headers tcPrChange"
    ),
    "trPr": (
        "cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader "
        "tblCellSpacing jc hidden ins del trPrChange"
    ),
    "tblPr": (
        "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize "
        "tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook "
        "tblCaption tblDescription tblPrChange"
    ),
    # Auch die Einstellungsdatei hat eine feste Folge; neu eingefügte Optionen
    # wie die Silbentrennung landen sonst hinter Elementen, die laut Schema
    # zuletzt kommen müssen.
    "settings": (
        "writeProtection view zoom removePersonalInformation removeDateAndTime "
        "doNotDisplayPageBoundaries displayBackgroundShape printPostScriptOverText "
        "printFractionalCharacterWidth printFormsData embedTrueTypeFonts "
        "embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins "
        "alignBordersAndEdges bordersDoNotSurroundHeader bordersDoNotSurroundFooter "
        "gutterAtTop hideSpellingErrors hideGrammaticalErrors activeWritingStyle "
        "proofState formsDesign attachedTemplate linkStyles stylePaneFormatFilter "
        "stylePaneSortMethod documentType mailMerge revisionView trackChanges "
        "doNotTrackMoves doNotTrackFormatting documentProtection autoFormatOverride "
        "styleLockTheme styleLockQFSet defaultTabStop autoHyphenation "
        "consecutiveHyphenLimit hyphenationZone doNotHyphenateCaps showEnvelope "
        "summaryLength clickAndTypeStyle defaultTableStyle evenAndOddHeaders "
        "bookFoldRevPrinting bookFoldPrinting bookFoldPrintingSheets "
        "drawingGridHorizontalSpacing drawingGridVerticalSpacing "
        "displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery "
        "doNotUseMarginsForDrawingGridOrigin drawingGridHorizontalOrigin "
        "drawingGridVerticalOrigin doNotShadeFormData noPunctuationKerning "
        "characterSpacingControl printTwoOnOne strictFirstAndLastChars "
        "noLineBreaksAfter noLineBreaksBefore savePreviewPicture "
        "doNotValidateAgainstSchema saveInvalidXml ignoreMixedContent "
        "alwaysShowPlaceholderText doNotDemarcateInvalidXml saveXmlDataOnly "
        "useXSLTWhenSaving saveThroughXslt showXMLTags alwaysMergeEmptyNamespace "
        "updateFields hdrShapeDefaults footnotePr endnotePr compat docVars rsids "
        "mathPr "  # steht laut Schema zwischen rsids und attachedSchema (m:-Namensraum)
        "attachedSchema themeFontLang clrSchemeMapping doNotIncludeSubdocsInStats "
        "doNotAutoCompressPictures forceUpgrade captions readModeInkLockDown "
        "smartTagType schemaLibrary shapeDefaults doNotEmbedSmartTags decimalSymbol "
        "listSeparator"
    ),
    "pBdr": "top left bottom right between bar",
    "tcBorders": "top start left bottom end right insideH insideV tl2br tr2bl",
    "tblBorders": "top start left bottom end right insideH insideV",
    "tblCellMar": "top start left bottom end right",
    "tcMar": "top start left bottom end right",
}
_RANK = {
    parent: {qn("w:" + name): i for i, name in enumerate(children.split())}
    for parent, children in _ORDER.items()
}
# mathPr ist das einzige Kind der Einstellungen aus dem Mathematik-Namensraum.
_RANK["settings"][qn("m:mathPr")] = _RANK["settings"].pop(qn("w:mathPr"))


def normalize(*roots):
    """Kinder aller Eigenschaftselemente in die vom Schema verlangte Folge bringen.

    Gilt für den Textkörper wie für Kopf- und Fusszeilen — jede dieser Teile
    ist ein eigenes XML-Dokument und muss einzeln übergeben werden.
    """
    for root in roots:
        for parent_name, rank in _RANK.items():
            nodes = list(root.iter(qn("w:" + parent_name)))
            if root.tag == qn("w:" + parent_name):
                nodes.append(root)
            for node in nodes:
                children = list(node)
                children.sort(key=lambda c: rank.get(c.tag, len(rank)))
                for child in children:
                    node.append(child)


ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

__all__ = [name for name in dir() if not name.startswith("_")]
