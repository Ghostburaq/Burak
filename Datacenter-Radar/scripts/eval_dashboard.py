#!/usr/bin/env python3
"""
Rechnet die Dashboard-Formeln in Python nach.

Ersatz fuer den LibreOffice-Recalc aus CLAUDE.md Abschnitt 5, falls
soffice nicht laeuft. Deckt genau die Funktionen ab, die auf 01_Dashboard
vorkommen: COUNTA, COUNT, COUNTIF und SUM, dazu Summen und Differenzen
solcher Aufrufe.

Aufruf:  python3 scripts/eval_dashboard.py
"""

import os
import re
import sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

HERE = os.path.dirname(os.path.abspath(__file__))
TARGET = os.path.join(os.path.dirname(HERE), "MiT_Datacenter_Radar_CH_V1.0.xlsx")

wb = load_workbook(TARGET)
radar = wb["02_Projekt_Radar"]
dash = wb["01_Dashboard"]

AUFRUF = re.compile(
    r"(COUNTA|COUNTIF|COUNT|SUM)\('02_Projekt_Radar'!"
    r"([A-Z]+)(\d+):([A-Z]+)(\d+)(?:,\"([^\"]*)\")?\)"
)


def bereich(sp_von, z_von, sp_bis, z_bis):
    werte = []
    for z in range(int(z_von), int(z_bis) + 1):
        for s in range(column_index_from_string(sp_von),
                       column_index_from_string(sp_bis) + 1):
            werte.append(radar.cell(row=z, column=s).value)
    return werte


def loese_auf(m):
    fn, sp1, z1, sp2, z2, kriterium = m.groups()
    werte = bereich(sp1, z1, sp2, z2)
    if fn == "COUNTA":
        return str(sum(1 for w in werte if w not in (None, "")))
    if fn == "COUNT":
        return str(sum(1 for w in werte if isinstance(w, (int, float))))
    if fn == "COUNTIF":
        return str(sum(1 for w in werte if str(w).strip() == kriterium))
    if fn == "SUM":
        return str(sum(w for w in werte if isinstance(w, (int, float))))
    raise ValueError(fn)


fehler = 0
print("01_Dashboard, nachgerechnet\n")
for zeile in dash.iter_rows():
    for zelle in zeile:
        if not (isinstance(zelle.value, str) and zelle.value.startswith("=")):
            continue
        label = dash.cell(row=zelle.row, column=zelle.column - 1).value
        formel = zelle.value[1:]
        ausdruck = AUFRUF.sub(loese_auf, formel)
        if re.search(r"[A-Za-z']", ausdruck):
            print(f"  NICHT AUFLOESBAR  {label}: {zelle.value}")
            fehler += 1
            continue
        try:
            wert = eval(ausdruck)  # nur Zahlen, Plus und Minus nach der Ersetzung
        except Exception as e:                                   # noqa: BLE001
            print(f"  FEHLER  {label}: {e}")
            fehler += 1
            continue
        if label:
            print(f"  {label:<40} {wert}")

if fehler:
    print(f"\n{fehler} Formeln konnten nicht nachgerechnet werden.")
    sys.exit(1)
print("\nAlle Dashboard-Formeln loesen sauber auf.")
