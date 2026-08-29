#!/usr/bin/env python3
"""
Prueft MiT_Datacenter_Radar_CH_V1.0.xlsx nach einem schreibenden Lauf.

Prueft: sieben Blaetter, Dropdowns, Formeln in R und S, gelbe Spalten,
Quellenpflicht, Dubletten ueber Betreiber plus Standort, Dropdown-Werte.
Laeuft ohne data_only, damit Formeln sichtbar bleiben.

Aufruf:  python3 scripts/check_radar.py
Exit 0 = sauber, Exit 1 = Befunde.
"""

import os
import sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
TARGET = os.path.join(os.path.dirname(HERE), "MiT_Datacenter_Radar_CH_V1.0.xlsx")

BLAETTER = ["00_Anleitung", "01_Dashboard", "02_Projekt_Radar", "03_Kontakte",
            "04_Akquise_Tracker", "05_Quellen_Montag", "06_Changelog"]

PHASEN = {"Planung", "Bewilligung", "Baustart", "Rohbau", "Fit-out", "Commissioning",
          "Betrieb", "Ausbau", "Verzoegert", "Gestoppt"}
STATUS = {"Kalt", "Angeschrieben", "Erstkontakt", "Erstgespraech", "Qualifiziert",
          "Offerte", "Gewonnen", "Verloren", "Ruhend"}
PRIO = {"A", "B", "C"}
KANAL = {"Betreiber", "GU/TU", "Fachplaner", "Ausschreibung", "Elektro-Installateur"}

GELB = "00FFF2CC"
befunde = []
hinweise = []


def fehler(text):
    befunde.append(text)


wb = load_workbook(TARGET)          # bewusst ohne data_only, sonst sind Formeln weg

# 1 Blattstruktur
if wb.sheetnames != BLAETTER:
    fehler(f"Blattstruktur weicht ab: {wb.sheetnames}")

ws = wb["02_Projekt_Radar"]

# 2 Kopfzeile
erwartet = ["Ref", "Projekt / Campus", "Betreiber", "Standort", "Kt.", "Kapazitaet_MW",
            "Phase", "IBN_Ziel", "GU_TU", "Fachplaner_Elektro", "Kanal_Prio", "MiT_Chance",
            "Ansprache_Fenster", "Prio", "Status_Kontakt", "Letzter_Kontakt", "Wiedervorlage",
            "Tage_bis_WV", "Ampel", "Quelle_URL", "Quelle_Stand", "Naechster_Schritt",
            "Bemerkung"]
ist = [ws.cell(row=4, column=i).value for i in range(1, 24)]
if ist != erwartet:
    fehler(f"Kopfzeile 02_Projekt_Radar weicht ab: {ist}")

# 3 Zeilen durchgehen
projekte, gesehen = 0, {}
for zeile in range(5, ws.max_row + 1):
    projekt = ws.cell(row=zeile, column=2).value
    if not projekt:
        continue
    projekte += 1

    r = ws.cell(row=zeile, column=18).value
    s = ws.cell(row=zeile, column=19).value
    if r != f'=IF(Q{zeile}="","",Q{zeile}-TODAY())':
        fehler(f"Zeile {zeile}: Formel in R fehlt oder weicht ab -> {r!r}")
    if not (isinstance(s, str) and s.startswith(f'=IF(Q{zeile}="","-"')):
        fehler(f"Zeile {zeile}: Formel in S fehlt oder weicht ab -> {s!r}")

    for spalte, name in ((15, "O"), (16, "P"), (17, "Q"), (22, "V")):
        fill = ws.cell(row=zeile, column=spalte).fill
        if fill is None or fill.fgColor is None or fill.fgColor.rgb != GELB:
            fehler(f"Zeile {zeile}: Spalte {name} ist nicht gelb (FFF2CC)")

    if not ws.cell(row=zeile, column=20).value:
        fehler(f"Zeile {zeile} ({projekt}): Quelle_URL in T fehlt")
    if not ws.cell(row=zeile, column=21).value:
        fehler(f"Zeile {zeile} ({projekt}): Quelle_Stand in U fehlt")

    for spalte, name, erlaubt in ((7, "Phase", PHASEN), (14, "Prio", PRIO),
                                  (11, "Kanal_Prio", KANAL), (15, "Status_Kontakt", STATUS)):
        wert = ws.cell(row=zeile, column=spalte).value
        if wert and wert not in erlaubt:
            fehler(f"Zeile {zeile}: {name} hat den unzulaessigen Wert {wert!r}")
        if not wert and name in ("Phase", "Kanal_Prio"):
            hinweise.append(f"Zeile {zeile} ({projekt}): {name} nicht belegt")

    schluessel = (str(ws.cell(row=zeile, column=3).value).strip().lower(),
                  str(ws.cell(row=zeile, column=4).value).strip().lower())
    if schluessel in gesehen:
        vorher_zeile, vorher_name = gesehen[schluessel]
        hinweise.append(f"Zeile {zeile} ({projekt}): gleicher Betreiber und Standort wie "
                        f"Zeile {vorher_zeile} ({vorher_name}). Pruefen ob Dublette oder "
                        f"zwei Bauetappen auf einem Campus.")
    else:
        gesehen[schluessel] = (zeile, projekt)

# 4 Dropdowns vorhanden
spalten_mit_dv = set()
for dv in ws.data_validations.dataValidation:
    for bereich in str(dv.sqref).split():
        spalten_mit_dv.add(bereich.split("!")[-1][0])
for erwartete_spalte in ("G", "K", "N", "O"):
    if erwartete_spalte not in spalten_mit_dv:
        fehler(f"Dropdown auf Spalte {erwartete_spalte} fehlt")

# 5 Dashboard zeigt auf den richtigen Bereich
dash = wb["01_Dashboard"]
formeln = [c.value for row in dash.iter_rows() for c in row
           if isinstance(c.value, str) and c.value.startswith("=")]
if not formeln:
    fehler("01_Dashboard enthaelt keine Formeln")
letzte_datenzeile = max([z for z in range(5, ws.max_row + 1)
                         if ws.cell(row=z, column=2).value] or [4])
if letzte_datenzeile > 45:
    fehler(f"Radar reicht bis Zeile {letzte_datenzeile}, Dashboard-Formeln zeigen nur bis 45. "
           f"Bereiche in allen Dashboard-Formeln mitwachsen lassen.")
for f in formeln:
    if any(verboten in f for verboten in ("ISOWEEKNUM", "XLOOKUP", "FILTER(", "UNIQUE(", "SORT(")):
        fehler(f"Verbotene Funktion in Dashboard-Formel: {f}")

# 6 Changelog
cl = wb["06_Changelog"]
cl_zeilen = sum(1 for z in range(5, cl.max_row + 1) if cl.cell(row=z, column=3).value)
if cl_zeilen == 0:
    fehler("06_Changelog ist leer")

# 7 Quellenblatt
q = wb["05_Quellen_Montag"]
# Nur Zeilen mit laufender Nummer in Spalte A zaehlen, damit Fussnoten nicht mitgezaehlt werden
q_daten = [z for z in range(5, q.max_row + 1)
           if isinstance(q.cell(row=z, column=1).value, int)]
q_zeilen = len(q_daten)
ohne_pruefdatum = [q.cell(row=z, column=2).value for z in q_daten
                   if not q.cell(row=z, column=6).value]

print(f"Blaetter            {len(wb.sheetnames)} / 7")
print(f"Projekte im Radar   {projekte}")
print(f"Changelog-Zeilen    {cl_zeilen}")
print(f"Quellen             {q_zeilen}, davon ohne Pruefdatum {len(ohne_pruefdatum)}")
print(f"Kontaktzeilen       {sum(1 for z in range(5, wb['03_Kontakte'].max_row + 1) if wb['03_Kontakte'].cell(row=z, column=1).value)}")

if hinweise:
    print("\nOFFENE PUNKTE")
    for h in hinweise:
        print(f"  {h}")
if ohne_pruefdatum:
    print("\nQUELLEN OHNE PRUEFDATUM")
    for name in ohne_pruefdatum:
        print(f"  {name}")

if befunde:
    print("\nBEFUNDE")
    for b in befunde:
        print(f"  {b}")
    sys.exit(1)

print("\nStruktur, Formeln, Dropdowns und Quellenpflicht sind intakt.")
