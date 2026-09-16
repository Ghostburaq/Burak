# -*- coding: utf-8 -*-
"""Baut die RAMS HVO Refuelling ZRH12 im MiT-Corporate-Layout (Vorlage MiT_Vorlage.docx)."""
import copy, os, shutil, zipfile, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

TPL = "/root/.claude/uploads/4884a130-1809-50b5-a3c8-dff0a6f0182b/ed9c7009-MiT_Vorlage.docx"
OUT = "/home/user/Burak/RAMS_HVO_Refuelling_ZRH12_Rev3_MiT.docx"

RED    = "E00036"
INK    = "1A1A1A"
GREY   = "5F5F5F"
LGREY  = "F4F4F5"
ZEBRA  = "FAFAFA"
LINE   = "DCDCDC"
HEAD   = "Alata"
BODY   = "IBM Plex Sans"

RISK = [(4, "3B8F4A", "Low"), (9, "E0A100", "Medium"), (14, "E2601A", "High"), (25, "C00028", "Unacceptable")]
def risk_band(v):
    for hi, col, name in RISK:
        if v <= hi:
            return col, name
    return RISK[-1][1], RISK[-1][2]

doc = Document(TPL)
body = doc.element.body

# ---------------------------------------------------------------- Titelseite
def set_textbox(marker, lines, size_pt):
    """Ersetzt den Platzhaltertext in allen Textbox-Kopien (wps + VML-Fallback)."""
    for t in body.iter(qn('w:t')):
        if (t.text or '').strip() != marker:
            continue
        run = t.getparent()
        rPr = run.find(qn('w:rPr'))
        t.text = lines[0]
        anchor = t
        for ln in lines[1:]:
            br = parse_xml('<w:br %s/>' % nsdecls('w'))
            anchor.addnext(br); anchor = br
            nt = parse_xml('<w:t %s xml:space="preserve"></w:t>' % nsdecls('w'))
            nt.text = ln
            anchor.addnext(nt); anchor = nt
        if rPr is not None:
            for tag, val in (('w:sz', str(int(size_pt*2))), ('w:szCs', str(int(size_pt*2)))):
                el = rPr.find(qn(tag))
                if el is None:
                    el = parse_xml('<%s %s w:val="%s"/>' % (tag, nsdecls('w'), val)); rPr.append(el)
                else:
                    el.set(qn('w:val'), val)

def _ln(el):
    return el.tag.split('}')[-1] if isinstance(el.tag, str) else ''

def move_box(marker, y_emu, cx=None, cy=None):
    """Verschiebt und skaliert die Textbox, die den Marker enthaelt."""
    for t in body.iter(qn('w:t')):
        if (t.text or '').strip() != marker:
            continue
        anchor = t
        while anchor is not None and _ln(anchor) != 'anchor':
            anchor = anchor.getparent()
        if anchor is None:
            continue
        pv = anchor.find(qn('wp:positionV'))
        if pv is not None:
            off = pv.find(qn('wp:posOffset'))
            if off is not None:
                off.text = str(y_emu)
        ext = anchor.find(qn('wp:extent'))
        if ext is not None:
            if cx: ext.set('cx', str(cx))
            if cy: ext.set('cy', str(cy))
        for e in anchor.iter(qn('a:ext')):
            if cx: e.set('cx', str(cx))
            if cy: e.set('cy', str(cy))
        return

# Platzhalterbild ("Bitte hier Bild einfuegen") der Vorlage entfernen
for dr in list(body.iter(qn('w:drawing'))):
    if any(d.get('name', '').startswith('Grafik') for d in dr.iter(qn('wp:docPr'))):
        run = dr.getparent()
        run.getparent().remove(run)

set_textbox("Titel", ["Risk Assessment", "& Method Statement"], 20)
set_textbox("Subtitel", ["Refuelling of the auxiliary generator with HVO fuel",
                         "Vantage ZRH12, Winterthur",
                         "",
                         "Mobil in Time AG \u2013 An Aggreko Company"], 11.5)
def resize_shape(name, y_emu=None, cx=None, cy=None):
    """Passt eine benannte Form (z. B. den roten Titelblock) an."""
    for dp in body.iter(qn('wp:docPr')):
        if dp.get('name') != name:
            continue
        anchor = dp.getparent()
        while anchor is not None and _ln(anchor) != 'anchor':
            anchor = anchor.getparent()
        if anchor is None:
            continue
        if y_emu is not None:
            pv = anchor.find(qn('wp:positionV'))
            if pv is not None:
                off = pv.find(qn('wp:posOffset'))
                if off is not None:
                    off.text = str(y_emu)
        ext = anchor.find(qn('wp:extent'))
        if ext is not None:
            if cx: ext.set('cx', str(cx))
            if cy: ext.set('cy', str(cy))
        for e in anchor.iter(qn('a:ext')):
            if cx: e.set('cx', str(cx))
            if cy: e.set('cy', str(cy))
        return

resize_shape("Rechteck 7", y_emu=426720, cx=3400425, cy=3400000)
move_box("Risk Assessment", 980000, 3200000, 1240000)
move_box("Refuelling of the auxiliary generator with HVO fuel", 2280000, 3200000, 1450000)

# Info-Leiste unterhalb des roten Blocks (verankerte Textbox)
def strip_line(text, bold=False, size=9, color=GREY, space_before=0, border=False):
    b = ('<w:pBdr><w:top w:val="single" w:sz="12" w:space="6" w:color="%s"/></w:pBdr>' % RED) if border else ''
    return ('<w:p><w:pPr>%s<w:spacing w:before="%d" w:after="40" w:line="240" w:lineRule="auto"/>'
            '<w:rPr><w:rFonts w:ascii="%s" w:hAnsi="%s"/><w:sz w:val="%d"/></w:rPr></w:pPr>'
            '<w:r><w:rPr><w:rFonts w:ascii="%s" w:hAnsi="%s"/>%s<w:color w:val="%s"/>'
            '<w:sz w:val="%d"/><w:szCs w:val="%d"/></w:rPr><w:t xml:space="preserve">%s</w:t></w:r></w:p>'
            % (b, space_before, BODY, BODY, size*2, BODY, BODY,
               '<w:b/>' if bold else '', color, size*2, size*2, text))

INFO_LINES = (
    strip_line(u"RAMS-MIT-ZRH12-004   \u00b7   Revision 3   \u00b7   Date of issue  _________________",
               bold=True, size=10, color=INK, border=True, space_before=0)
    + strip_line(u"Issued for submission to DPR Construction CH AG", size=9)
    + strip_line(u"Site: Vantage ZRH12, Fabrikstrasse 10, 8404 Winterthur", size=9)
    + strip_line(u"Prepared by Burak \u00dcc\u00f6z, Project Manager Switzerland", size=9)
)

NS_BOX = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')

INFO_BOX = ('<w:r %s><w:rPr><w:noProof/></w:rPr><w:drawing>'
 '<wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="251700000" '
 'behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
 '<wp:simplePos x="0" y="0"/>'
 '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
 '<wp:positionV relativeFrom="paragraph"><wp:posOffset>4060000</wp:posOffset></wp:positionV>'
 '<wp:extent cx="5760000" cy="900000"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
 '<wp:docPr id="901" name="Infoleiste"/><wp:cNvGraphicFramePr/>'
 '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
 '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="5760000" cy="900000"/></a:xfrm>'
 '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>'
 '<wps:txbx><w:txbxContent>%s</w:txbxContent></wps:txbx>'
 '<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"><a:noAutofit/></wps:bodyPr>'
 '</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'
 % (NS_BOX, INFO_LINES))

title_p = body[0]
title_p.append(parse_xml(INFO_BOX))

# ------------------------------------------------- Vorlagen-Platzhalter raus
sectPr = body.find(qn('w:sectPr'))
for el in list(body):
    if el is title_p or el is sectPr:
        continue
    body.remove(el)

# --------------------------------------------------------------- Bausteine
def _p():
    p = parse_xml('<w:p %s/>' % nsdecls('w'))
    body.insert(list(body).index(sectPr), p)
    from docx.text.paragraph import Paragraph
    return Paragraph(p, doc)

def rpr(run, size=11, color=INK, bold=False, font=BODY, italic=False, caps=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    r = run._element.get_or_add_rPr()
    for tag in ('w:rFonts',):
        el = r.find(qn(tag))
        if el is not None:
            el.set(qn('w:cs'), font); el.set(qn('w:eastAsia'), font)
    if caps:
        r.append(parse_xml('<w:caps %s w:val="1"/>' % nsdecls('w')))
        r.append(parse_xml('<w:spacing %s w:val="30"/>' % nsdecls('w')))
    return run

def spacing(p, before=0, after=0, line=None, keep_next=False, keep_lines=True):
    pPr = p._p.get_or_add_pPr()
    sp = pPr.find(qn('w:spacing'))
    if sp is None:
        sp = parse_xml('<w:spacing %s/>' % nsdecls('w')); pPr.append(sp)
    sp.set(qn('w:before'), str(before)); sp.set(qn('w:after'), str(after))
    if line:
        sp.set(qn('w:line'), str(line)); sp.set(qn('w:lineRule'), 'auto')
    if keep_next:
        pPr.append(parse_xml('<w:keepNext %s/>' % nsdecls('w')))
    if keep_lines:
        pPr.append(parse_xml('<w:keepLines %s/>' % nsdecls('w')))

def h1(num, text, pagebreak=False):
    p = _p()
    if pagebreak:
        p.paragraph_format.page_break_before = True
    p.style = doc.styles['Überschrift']
    if num:
        rpr(p.add_run(num + "   "), size=18, color=RED, bold=False, font=HEAD)
    rpr(p.add_run(text), size=18, color=RED, bold=False, font=HEAD)
    spacing(p, before=360 if not pagebreak else 0, after=60, keep_next=True)
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml('<w:pBdr %s><w:bottom w:val="single" w:sz="8" w:space="4" w:color="%s"/></w:pBdr>'
                         % (nsdecls('w'), RED)))
    return p

def h2(text):
    p = _p()
    p.style = doc.styles['Unterüberschrift']
    rpr(p.add_run(text), size=12.5, color=RED, bold=False, font=HEAD)
    spacing(p, before=280, after=80, keep_next=True)
    return p

def h3(text):
    p = _p()
    p._p.get_or_add_pPr().append(parse_xml('<w:ind %s w:left="0" w:firstLine="0"/>' % nsdecls('w')))
    rpr(p.add_run(text), size=10, color=INK, bold=True, font=BODY)
    spacing(p, before=240, after=60, keep_next=True)
    return p

def para(text, size=10.5, color=INK, before=0, after=110, bold=False, italic=False, justify=True):
    p = _p()
    p.style = doc.styles['Fliesstext']
    rpr(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    spacing(p, before=before, after=after, line=250, keep_lines=False)
    return p

def bullet(text, size=10.5, after=50):
    p = _p()
    p.style = doc.styles['Fliesstext']
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml('<w:ind %s w:left="340" w:hanging="340"/>' % nsdecls('w')))
    pPr.append(parse_xml('<w:tabs %s><w:tab w:val="left" w:pos="340"/></w:tabs>' % nsdecls('w')))
    rpr(p.add_run("▪"), size=size, color=RED, bold=True)
    p.add_run("\t")
    rpr(p.add_run(text), size=size, color=INK)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    spacing(p, after=after, line=250, keep_lines=False)
    return p

def note(text, size=9.5):
    """Hervorgehobener Hinweis: linke rote Linie, hellgrauer Grund."""
    p = _p()
    p.style = doc.styles['Fliesstext']
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml('<w:pBdr %s><w:left w:val="single" w:sz="18" w:space="8" w:color="%s"/></w:pBdr>'
                         % (nsdecls('w'), RED)))
    pPr.append(parse_xml('<w:shd %s w:val="clear" w:color="auto" w:fill="%s"/>' % (nsdecls('w'), LGREY)))
    pPr.append(parse_xml('<w:ind %s w:left="170" w:right="113"/>' % nsdecls('w')))
    rpr(p.add_run(text), size=size, color=INK)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    spacing(p, before=120, after=160, line=250, keep_lines=False)
    return p

# ------------------------------------------------------------ Tabellenhilfen
def new_table(rows, cols, widths):
    """Tabelle mit fixen Spaltenbreiten; tblPr wird in Schema-Reihenfolge gesetzt."""
    t = doc.add_table(rows=rows, cols=cols)
    sectPr.addprevious(t._tbl)
    tbl = t._tbl
    total = sum(widths)
    tbl.replace(tbl.tblPr, parse_xml(
        '<w:tblPr %s>'
        '<w:tblW w:w="%d" w:type="dxa"/>'
        '<w:tblInd w:w="0" w:type="dxa"/>'
        '<w:tblBorders>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
        '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
        '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
        '<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '</w:tblBorders>'
        '<w:tblLayout w:type="fixed"/>'
        '<w:tblCellMar>'
        '<w:top w:w="70" w:type="dxa"/><w:left w:w="100" w:type="dxa"/>'
        '<w:bottom w:w="70" w:type="dxa"/><w:right w:w="100" w:type="dxa"/>'
        '</w:tblCellMar>'
        '<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="1" '
        'w:lastColumn="0" w:noHBand="0" w:noVBand="1"/>'
        '</w:tblPr>' % (nsdecls('w'), total, LINE, LINE, LINE)))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    grid = tbl.find(qn('w:tblGrid'))
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(w))
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            tcPr = c._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn('w:tcW')):
                tcPr.remove(old)
            tcPr.insert(0, parse_xml('<w:tcW %s w:w="%d" w:type="dxa"/>' % (nsdecls('w'), w)))
    return t

def table_borders(t, inside_h=None, inside_v=None, top=None, bottom=None):
    """Ueberschreibt einzelne Rahmenlinien einer Tabelle."""
    bd = t._tbl.tblPr.find(qn('w:tblBorders'))
    for tag, spec in (('w:top', top), ('w:bottom', bottom),
                      ('w:insideH', inside_h), ('w:insideV', inside_v)):
        if spec is None:
            continue
        el = bd.find(qn(tag))
        val, sz, color = spec
        el.set(qn('w:val'), val); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), color)

def shade(cell, color):
    cell._tc.get_or_add_tcPr().append(
        parse_xml('<w:shd %s w:val="clear" w:color="auto" w:fill="%s"/>' % (nsdecls('w'), color)))

def cell_margins(cell, left=None, right=None, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:tcMar')):
        tcPr.remove(old)
    parts = ''
    for tag, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        if val is not None:
            parts += '<w:%s w:w="%d" w:type="dxa"/>' % (tag, val)
    tcPr.append(parse_xml('<w:tcMar %s>%s</w:tcMar>' % (nsdecls('w'), parts)))

def vcenter(cell):
    cell._tc.get_or_add_tcPr().append(parse_xml('<w:vAlign %s w:val="center"/>' % nsdecls('w')))

def cell_text(cell, text, size=9.5, color=INK, bold=False, align='left', font=BODY,
              caps=False, space_after=0, line=240):
    cell.text = ''
    p = cell.paragraphs[0]
    first = True
    for chunk in str(text).split('\n'):
        if not first:
            p = cell.add_paragraph()
        first = False
        rpr(p.add_run(chunk), size=size, color=color, bold=bold, font=font, caps=caps)
        p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                       'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml('<w:spacing %s w:before="10" w:after="%d" w:line="%d" w:lineRule="auto"/>'
                             % (nsdecls('w'), space_after, line)))
    return cell

def header_row(t, labels, sizes=9, align=None):
    align = align or ['left'] * len(labels)
    for i, lab in enumerate(labels):
        c = t.rows[0].cells[i]
        shade(c, RED); vcenter(c)
        cell_text(c, lab, size=sizes, color='FFFFFF', bold=True, align=align[i])
    t.rows[0]._tr.get_or_add_trPr().append(parse_xml('<w:tblHeader %s/>' % nsdecls('w')))
    return t

def row_height(row, h):
    row._tr.get_or_add_trPr().append(
        parse_xml('<w:trHeight %s w:val="%d" w:hRule="atLeast"/>' % (nsdecls('w'), h)))

def keep_row_together(t):
    for row in t.rows:
        row._tr.get_or_add_trPr().append(parse_xml('<w:cantSplit %s/>' % nsdecls('w')))
        for c in row.cells:
            for p in c.paragraphs:
                p._p.get_or_add_pPr().append(parse_xml('<w:keepLines %s/>' % nsdecls('w')))

def kv_table(data, widths=(2600, 6471), header=("Field", "Detail"), size=9.5):
    t = new_table(len(data) + 1, 2, list(widths))
    header_row(t, list(header))
    for i, (k, v) in enumerate(data, start=1):
        r = t.rows[i]
        shade(r.cells[0], LGREY)
        cell_text(r.cells[0], k, size=size, bold=True, color=INK)
        cell_text(r.cells[1], v, size=size, color=INK)
        vcenter(r.cells[0])
    keep_row_together(t)
    spacer()
    return t

def spacer(h=140):
    p = _p()
    p.style = doc.styles['Fliesstext']
    rpr(p.add_run(""), size=2)
    spacing(p, after=h, line=240)
    return p

def chip_cell(cell, value):
    """Farbiger Risiko-Chip in einer Zelle."""
    col, name = risk_band(value)
    shade(cell, col); vcenter(cell)
    cell.text = ''
    p = cell.paragraphs[0]
    rpr(p.add_run(str(value)), size=11, color='FFFFFF', bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.get_or_add_pPr().append(parse_xml('<w:spacing %s w:before="30" w:after="0" w:line="220" w:lineRule="auto"/>' % nsdecls('w')))
    return cell

def steps(items, start=1):
    """Nummerierte Ablaufschritte mit rotem Nummernfeld."""
    t = new_table(len(items), 2, [660, 8411])
    table_borders(t, inside_h=('single', 12, 'FFFFFF'),
                  top=('none', 0, 'auto'), bottom=('none', 0, 'auto'))
    for i, txt in enumerate(items):
        r = t.rows[i]
        num = r.cells[0]
        shade(num, RED)
        cell_text(num, str(start + i), size=10, color='FFFFFF', bold=True, align='center')
        cell_margins(num, left=30, right=30, top=90, bottom=120)
        cell_text(r.cells[1], txt, size=10, color=INK, align='justify', space_after=130)
        cell_margins(r.cells[1], left=200, bottom=140)
    keep_row_together(t)
    spacer(60)
    return t

W = 9071  # nutzbare Breite in twips

# ================================================================ INHALT
# ---- Seitenumbruch nach Titelseite
brk = _p()
brk.paragraph_format.page_break_before = True
rpr(brk.add_run(""), size=2)
spacing(brk, after=0)

# ---------------------------------------------------------------- Contents
h1("", "Contents")
toc = [("1", "Document Control", "8", "Risk Assessment"),
       ("2", "Scope of Works", "9", "Personal Protective Equipment"),
       ("3", "Reference Documents and Legal Framework", "10", "Environmental Controls and Waste"),
       ("4", "Equipment, Substance Data and Certification", "11", "Emergency Arrangements"),
       ("5", "Personnel, Competence and Supervision", "12", "Monitoring and Review"),
       ("6", "Timing, Delivery Booking and Interfaces", "13", "Briefing Record and Declaration"),
       ("7", "Method Statement: Sequence of Operations", "", "")]
t = new_table(len(toc), 4, [520, 4015, 520, 4016])
for i, (a, b, c, d) in enumerate(toc):
    r = t.rows[i]
    cell_text(r.cells[0], a, size=10, color=RED, bold=True, align='center')
    cell_text(r.cells[1], b, size=10, color=INK)
    cell_text(r.cells[2], c, size=10, color=RED, bold=True, align='center')
    cell_text(r.cells[3], d, size=10, color=INK)
    for cc in r.cells:
        vcenter(cc)
    row_height(r, 300)
keep_row_together(t)
spacer(200)

# ---------------------------------------------------------------- 1
h1("1", "Document Control")
kv_table([
 ("Document reference", "RAMS-MIT-ZRH12-004"),
 ("Revision", "Rev. 3"),
 ("Date of issue", "____________________"),
 ("Valid from", "Date of issue until superseded by a later revision"),
 ("Project", "Vantage ZRH12, Fabrikstrasse 10, 8404 Winterthur"),
 ("Principal Contractor", "DPR Construction CH AG"),
 ("Trade Partner", "Mobil in Time AG – An Aggreko Company"),
 ("Activity", "Delivery of HVO fuel by road tanker and transfer into the 3,000 litre double-walled tank supplying the 100 kVA auxiliary generator"),
 ("Location", "External trailer position within the load bank set-up area"),
 ("Frequency", "Initial fill, then periodic top-up over the rental period, typically once the tank level falls below 30 per cent"),
 ("Prepared by", "Burak Ücöz, Project Manager Switzerland"),
 ("Applies to", "MiT personnel, the fuel supplier's tanker driver, and any person within the demarcated working area"),
])

h2("Revision history")
rev = [("0", "First issue", "Initial submission for the refuelling operation"),
       ("1", "Legal framework, hierarchy of controls, competence requirements and rating legend added",
             "Alignment with the site EHS plan and Swiss regulatory references"),
       ("2", "Working hours and delivery booking, supervision, interfaces with other trades, certification, spill kit inventory, emergency contacts, monitoring and review added",
             "Full alignment with the Supplemental Conditions and the required method statement contents"),
       ("3", "Contents page and issue date added; editorial review throughout",
             "Final issue for submission to DPR Construction")]
t = new_table(len(rev) + 1, 3, [700, 4320, 4051])
header_row(t, ["Rev", "Change", "Reason"], align=['center', 'left', 'left'])
for i, (a, b, c) in enumerate(rev, start=1):
    r = t.rows[i]
    shade(r.cells[0], LGREY); vcenter(r.cells[0])
    cell_text(r.cells[0], a, size=10, bold=True, color=RED, align='center')
    cell_text(r.cells[1], b, size=9.5)
    cell_text(r.cells[2], c, size=9.5, color=GREY)
keep_row_together(t)
spacer()

h2("Review")
for b in ["Before every delivery as part of the pre-start evaluation.",
          "On any change of location, equipment, personnel, fuel product or site conditions.",
          "After any incident, near miss or spillage, before the next delivery takes place."]:
    bullet(b)

# ---------------------------------------------------------------- 2
h1("2", "Scope of Works")
para("This RAMS covers the delivery of HVO fuel by road tanker and its transfer into the double-walled bunded tank that supplies the 100 kVA auxiliary generator. The generator provides auxiliary power to the fans and the control system of the 6 MVA load bank and is positioned externally on the trailer.")
para("The activity is limited to fuelling the generator set and its associated tank. No other plant or equipment on site is fuelled by this operation.")
h3("Explicitly excluded")
for b in ["Any work on the electrical installation of the generator, the load bank or the customer's installation",
          "Fuelling of third-party plant or equipment",
          "Decanting into drums, cans or intermediate containers",
          "Any activity inside the building",
          "Hot works of any kind"]:
    bullet(b)

# ---------------------------------------------------------------- 3
h1("3", "Reference Documents and Legal Framework")
kv_table([
 ("Vantage ZRH12 Environmental Health & Safety Plan, Rev. 4", "Governing site EHS document, including the hierarchy of controls and the permit regime"),
 ("DPR Supplemental Conditions", "Working hours, delivery notification and booking, vehicle movements, PPE, site rules"),
 ("SUVA guidance and EKAS directives", "Swiss occupational safety framework"),
 ("Bauarbeitenverordnung (BauAV)", "Swiss construction site safety ordinance"),
 ("Gewässerschutzgesetz and Gewässerschutzverordnung", "Protection of ground and surface water from substances hazardous to water"),
 ("VKF fire protection directives", "Handling and storage of flammable liquids"),
 ("SDR / ADR", "Carriage of dangerous goods by road; driver qualification and vehicle equipment"),
 ("VeVA (Ordinance on Movements of Waste)", "Disposal of contaminated absorbent material"),
 ("EN 15940", "Product standard for paraffinic diesel fuel (HVO)"),
], widths=(3400, 5671), header=("Reference", "Relevance"))
note("This assessment applies the hierarchy of controls set out in the site EHS plan: elimination and substitution where reasonably practicable, then engineering controls, then administrative controls, with personal protective equipment as the final layer rather than the primary measure.")

# ---------------------------------------------------------------- 4
h1("4", "Equipment, Substance Data and Certification")
kv_table([
 ("Generator", "100 kVA auxiliary set, Stage V emissions compliant"),
 ("Fuel tank", "3,000 litre double-walled combination tank with integrated secondary containment and overfill protection"),
 ("Fuel", "HVO, hydrotreated vegetable oil to EN 15940 (paraffinic diesel fuel)"),
 ("Flash point", "Above 55 °C. Flammable liquid, category 3 under CLP. The fire risk profile is comparable to conventional diesel and not to volatile fuels such as petrol"),
 ("Transport classification", "UN 1202, carried under SDR/ADR"),
 ("Environmental classification", "Substance hazardous to water; release to ground, drainage or surface water must be prevented"),
 ("Delivery vehicle", "Road tanker of the appointed fuel supplier, equipped in accordance with SDR/ADR"),
 ("Transfer method", "Pumped transfer via the tanker's own hose and metered nozzle, attended throughout"),
 ("Typical duration", "30 to 45 minutes per delivery, including positioning and clearing the area"),
], widths=(2800, 6271), header=("Item", "Detail"))

h2("Certification held and available for inspection")
for b in ["SDR/ADR approval certificate for the tanker and current vehicle inspection",
          "SDR/ADR training certificate for the driver",
          "Current inspection tag on the fire extinguisher provided at the delivery point",
          "Earthing and bonding cable with visual inspection before each use",
          "Safety data sheet for HVO, held at the point of use",
          "Tank documentation and overfill protection specification"]:
    bullet(b)

# ---------------------------------------------------------------- 5
h1("5", "Personnel, Competence and Supervision")
kv_table([
 ("MiT Project Manager", "Notifies DPR in advance of each delivery and books the delivery slot; confirms route and standing position; verifies that the controls in this RAMS are in place before transfer begins; site induction completed; communicates in English and German"),
 ("Tanker driver", "Holds a valid SDR/ADR qualification; operates pump and nozzle; remains at the delivery point throughout; site induction completed and valid site access badge held"),
 ("MiT attendant", "Second competent person present throughout; monitors the tank level; acts as banksman during all vehicle manoeuvres; initiates emergency stop where required; trained in the use of the spill kit"),
 ("DPR site management", "Grants site access; confirms route and standing position; issues any permit required under the site permit regime"),
], widths=(2400, 6671), header=("Role", "Competence and responsibility"))
note("A minimum of two persons is present for the entire operation. The transfer is never left unattended at any point.")

h2("Supervision and language")
bullet("An English-speaking supervisor from Mobil in Time is present on site for the duration of the operation, in accordance with the Supplemental Conditions.")
bullet("Where the driver does not command sufficient English or German, instructions and the content of this RAMS are translated by MiT before the operation begins, and communication is maintained in a language the driver understands.")

h2("Briefing")
bullet("All persons involved are briefed on this RAMS before the first delivery and confirm their understanding by signature in section 13.")
bullet("A pre-start evaluation is carried out before every subsequent delivery, covering current site conditions, the standing position and any change since the previous operation.")
bullet("Under the site EHS plan, refuelling is not listed among the activities requiring a permit to work. MiT will nevertheless apply for a permit or attend a toolbox talk on request by DPR.")

# ---------------------------------------------------------------- 6
h1("6", "Timing, Delivery Booking and Interfaces")
h2("6.1  Working hours")
bullet("The operation is carried out within the site working hours of Monday to Friday, 07:00 to 17:30.")
bullet("Should a delivery outside these hours become necessary, it is notified to DPR at least three working days in advance and carried out only with a DPR supervisor present, in accordance with the Supplemental Conditions.")
h2("6.2  Delivery booking")
bullet("Each delivery is booked through the site booking system with the required notice. The tanker does not attend site before its allocated slot and does not wait on approach roads.")
bullet("Mobil in Time currently has no access to the site booking platform. Until access is provided, deliveries are notified to DPR site management by email and confirmed before the tanker is dispatched.")
h2("6.3  Interfaces with other trades and the adjacent facility")
bullet("The route and standing position are agreed with DPR so that access routes, laydown areas and the working areas of other trades are not obstructed.")
bullet("The operation is coordinated through the daily site briefing so that adjacent trades are aware of the activity and the demarcated area.")
bullet("Particular attention is paid to the proximity of the live ZRH11 facility and of the railway line: no obstruction of emergency routes, and no activity that could affect the operation of the adjacent facility.")
bullet("Should another trade require access through the demarcated area during transfer, the transfer is paused until safe passage is arranged.")

# ---------------------------------------------------------------- 7
h1("7", "Method Statement: Sequence of Operations")
h2("7.1  Before arrival")
steps([
 "The delivery is notified to DPR site management and the slot confirmed. The tanker and driver are registered for site access in accordance with the site access procedure.",
 "The approved access route and the standing position of the tanker are confirmed with DPR. The route is kept clear of pedestrian walkways as far as reasonably practicable.",
 "Ground conditions and weather are assessed. The operation is postponed in the event of an electrical storm or where ground conditions are unsuitable for the tanker.",
 "A pre-start evaluation is held with all persons involved, covering the sequence, the controls and the emergency arrangements.",
], start=1)
h2("7.2  Arrival and positioning")
steps([
 "The tanker enters site at walking pace and is escorted. The MiT attendant acts as banksman for all reversing and manoeuvring movements.",
 "The tanker is positioned on level, firm ground at the designated point. The parking brake is applied and the wheels are chocked. Where the transfer pump is driven by the vehicle's power take-off, the engine runs at idle for the pump only; no movement of the vehicle takes place while the hose is connected.",
 "The working area is demarcated with rigid barriers. Barrier tape is not used, in accordance with the site rules. No-smoking and no-naked-flame signage is displayed, and mobile telephones are not used within the demarcated area.",
 "The spill kit and a fire extinguisher of at least 6 kg ABC dry powder are positioned at the delivery point before any connection is made. The tanker's own extinguishers remain accessible.",
], start=5)
h2("7.3  Preparation")
steps([
 "The generator is shut down and allowed to cool before the transfer begins. Refuelling of a running set does not take place.",
 "The tank level is verified and the delivery volume calculated so that a filling level of 95 per cent is not exceeded.",
 "The tank shell and the secondary containment are inspected visually for damage or accumulated liquid. Any water present in the containment is removed and disposed of correctly before filling begins.",
 "An earthing and bonding connection is established between tanker and tank before the hose is connected, to dissipate any static charge.",
 "A drip tray is placed beneath the coupling point.",
], start=9)
h2("7.4  Transfer")
steps([
 "The hose is connected and the coupling is checked for secure engagement before the pump is started.",
 "Transfer begins at reduced flow. The connection is observed for leakage, and only once flow is confirmed stable and leak-free is the normal rate applied.",
 "Both operatives remain at the delivery point for the entire transfer. The tank level is monitored continuously and the emergency stop is kept within reach.",
 "The transfer is stopped at the predetermined volume and in every case before the 95 per cent level is reached.",
], start=14)
h2("7.5  Completion")
steps([
 "The pump is stopped, the hose is drained back into the tanker and the coupling is disconnected over the drip tray.",
 "The earthing connection is removed only after the hose has been drained and disconnected. The tank cap is closed and secured.",
 "The delivery point is inspected for spillage. Any residue is absorbed with material from the spill kit and collected for disposal as hazardous waste.",
 "Barriers are removed, the delivery note is signed and the tanker is escorted off site.",
 "The generator is restarted and checked for correct operation. The delivery is recorded in the project fuel log and reported to DPR as part of routine site reporting.",
], start=18)

# ---------------------------------------------------------------- 8
h1("8", "Risk Assessment")
h2("8.1  Rating method")
t = new_table(6, 2, [4535, 4536])
header_row(t, ["Likelihood (L)", "Severity (S)"])
LS = [("1  Very unlikely", "1  Negligible, no injury"),
      ("2  Unlikely", "2  Minor, first aid"),
      ("3  Possible", "3  Moderate, lost time"),
      ("4  Likely", "4  Major, serious injury"),
      ("5  Almost certain", "5  Fatality or major environmental damage")]
for i, (a, b) in enumerate(LS, start=1):
    r = t.rows[i]
    cell_text(r.cells[0], a, size=9.5)
    cell_text(r.cells[1], b, size=9.5)
    if i % 2 == 0:
        shade(r.cells[0], ZEBRA); shade(r.cells[1], ZEBRA)
    row_height(r, 260)
keep_row_together(t)
spacer()

t = new_table(5, 3, [1500, 2200, 5371])
header_row(t, ["Rating L × S", "Band", "Action"])
BANDS = [("1 to 4", "Low", "3B8F4A", "Acceptable, maintain controls"),
         ("5 to 9", "Medium", "E0A100", "Acceptable with all listed controls applied"),
         ("10 to 14", "High", "E2601A", "Additional controls required before work proceeds"),
         ("15 to 25", "Unacceptable", "C00028", "Work must not proceed in this form")]
for i, (rng, band, col, action) in enumerate(BANDS, start=1):
    r = t.rows[i]
    cell_text(r.cells[0], rng, size=9.5, bold=True, align='center'); vcenter(r.cells[0])
    shade(r.cells[1], col); vcenter(r.cells[1])
    cell_text(r.cells[1], band, size=9.5, color='FFFFFF', bold=True, align='center')
    cell_text(r.cells[2], action, size=9.5); vcenter(r.cells[2])
    row_height(r, 300)
keep_row_together(t)
spacer()
note("No element of this activity is carried out where the residual rating exceeds 9. Initial ratings are stated without controls in place; residual ratings assume every listed control is applied.")

h2("8.2  Hazard assessment")
HAZ = [
 ("Release of fuel to ground, drainage or surface water", "Operatives, environment", 3, 4,
  "Double-walled tank with integrated containment; drip tray at the coupling; containment inspected and emptied before filling; spill kit at the delivery point; operatives trained in its use; immediate notification of DPR in the event of any release", 1, 3),
 ("Overfilling of the tank", "Operatives, environment", 2, 4,
  "Delivery volume calculated in advance from the measured level; maximum fill 95 per cent; overfill protection fitted; level monitored continuously; transfer attended throughout", 1, 2),
 ("Ignition of fuel or vapour, fire", "All site personnel", 2, 5,
  "HVO flash point above 55 °C; generator shut down and cooled before transfer; no hot works in or adjacent to the area; smoking, naked flames and mobile telephone use excluded from the demarcated area; ABC extinguisher at the delivery point; tanker extinguishers accessible", 1, 3),
 ("Static discharge during transfer", "Operatives", 2, 4,
  "Earthing and bonding established between tanker and tank before connection; connection removed only after the hose is drained and disconnected; flow started at reduced rate", 1, 2),
 ("Vehicle movement striking persons, plant or structures", "Operatives, site personnel", 3, 5,
  "Escorted entry at walking pace; banksman for all reversing and manoeuvring; rigid barriers around the working area; pedestrian routes kept clear; parking brake applied and wheels chocked; no vehicle movement while the hose is connected", 1, 4),
 ("Hose or coupling failure under pressure", "Operatives", 2, 3,
  "Coupling checked for secure engagement before start; transfer started at reduced flow and observed; operatives positioned clear of the coupling; emergency stop within reach throughout", 1, 2),
 ("Skin or eye contact with fuel", "Operatives", 3, 2,
  "Fuel-resistant gloves and safety glasses worn for connection and disconnection; eyewash available at the delivery point; safety data sheet at the point of use; washing facilities used after the operation", 1, 2),
 ("Slip on spilled fuel", "Operatives", 2, 3,
  "Drip tray in place throughout; immediate absorption of any residue; area inspected and confirmed clear before barriers are removed", 1, 2),
 ("Work in proximity to energised equipment", "Operatives", 2, 5,
  "Generator shut down before transfer; electrical enclosures remain closed and are not accessed during fuelling; electrical work strictly separated in time from the fuelling activity", 1, 2),
 ("Inhalation of vapour", "Operatives", 2, 2,
  "Operation carried out in the open air; operatives positioned upwind where practicable; low volatility of the product", 1, 2),
 ("Adverse weather or unsuitable ground", "Operatives, environment", 2, 3,
  "Operation postponed during electrical storms; ground assessed for bearing capacity and level before positioning; no operation on unconsolidated ground", 1, 2),
 ("Unauthorised persons entering the working area", "Third parties", 3, 3,
  "Rigid barriers and signage; two attendants present; the area is not left unattended at any time during the transfer", 1, 3),
 ("Interference with access routes or other trades", "Site personnel", 2, 3,
  "Route and standing position agreed with DPR in advance; activity announced at the daily briefing; transfer paused if access through the area is required; emergency routes kept clear at all times", 1, 2),
]
HW = [2000, 1400, 700, 4271, 700]
t = new_table(len(HAZ) + 1, 5, HW)
header_row(t, ["Hazard", "Persons at risk", "Init.", "Control measures", "Resid."],
           sizes=8.5, align=['left', 'left', 'center', 'left', 'center'])
for i, (hz, pers, l1, s1, ctrl, l2, s2) in enumerate(HAZ, start=1):
    r = t.rows[i]
    init, res = l1 * s1, l2 * s2
    cell_text(r.cells[0], hz, size=8.5, bold=True)
    cell_text(r.cells[1], pers, size=8.5, color=GREY)
    for idx, (val, ll, ss) in ((2, (init, l1, s1)), (4, (res, l2, s2))):
        chip_cell(r.cells[idx], val)
        q = r.cells[idx].add_paragraph()
        rpr(q.add_run("%d \u00d7 %d" % (ll, ss)), size=6.5, color='FFFFFF')
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q._p.get_or_add_pPr().append(parse_xml(
            '<w:spacing %s w:before="0" w:after="40" w:line="180" w:lineRule="auto"/>' % nsdecls('w')))
        cell_margins(r.cells[idx], left=20, right=20, top=60, bottom=40)
    cell_text(r.cells[3], ctrl, size=8.5, align='left')
    if i % 2 == 0:
        for idx in (0, 1, 3):
            shade(r.cells[idx], ZEBRA)
keep_row_together(t)
spacer()

# ---------------------------------------------------------------- 9
h1("9", "Personal Protective Equipment")
note("PPE is the final layer of control and does not replace the engineering and administrative measures set out above.")
for b in ["Safety footwear, safety helmet, high-visibility clothing and long trousers at all times on site",
          "Fuel-resistant gloves (nitrile) for connection, disconnection and any handling of fuel or contaminated material",
          "Safety glasses for connection and disconnection",
          "Any additional PPE required by the site-specific rules of DPR Construction"]:
    bullet(b)

# ---------------------------------------------------------------- 10
h1("10", "Environmental Controls and Waste")
for b in ["Double-walled tank with integrated secondary containment maintained throughout the rental period",
          "Drip tray beneath every coupling point during transfer",
          "Spill kit held at the delivery point, comprising absorbent pads and granulate, a drain cover, sealable waste bags, gloves and a scoop; contents checked before each delivery and replenished after any use",
          "Safety data sheet for HVO held at the point of use and provided to DPR on request",
          "No discharge of any liquid from the containment or drip trays to ground, drainage or surface water under any circumstances",
          "Contaminated absorbent material collected in sealed containers and disposed of as hazardous waste through an approved route in accordance with VeVA; disposal documentation retained and available to DPR",
          "Deliveries recorded in the project fuel log and reported to DPR as part of routine site reporting"]:
    bullet(b)

# ---------------------------------------------------------------- 11
h1("11", "Emergency Arrangements")
h2("11.1  Spillage")
steps([
 "Stop the transfer immediately at the pump emergency stop.",
 "Contain the spill with absorbent material and place the drain cover to prevent migration towards drainage or unsealed ground.",
 "Notify DPR site management without delay.",
 "Recover all contaminated material into sealed containers for disposal as hazardous waste.",
 "Record the incident and review this RAMS before the next delivery takes place.",
])
h2("11.2  Fire")
steps([
 "Stop the transfer, raise the alarm and evacuate to the site assembly point.",
 "Tackle only a small incipient fire with the extinguisher provided, and only where it is safe to do so.",
 "Follow the DPR site emergency procedure and the instructions of site management.",
])
h2("11.3  Injury or exposure")
steps([
 "Skin contact: wash the affected area with soap and water. Eye contact: rinse thoroughly with the eyewash and seek medical advice.",
 "Report every injury to DPR site management immediately, however minor, in accordance with the site incident reporting procedure.",
 "First aid facilities and the site first aider are as advised in the site induction.",
])
h2("11.4  Emergency contacts")
CONT = [("Fire and hazardous materials (Switzerland)", "118"),
        ("Ambulance (Switzerland)", "144"),
        ("Police (Switzerland)", "117"),
        ("Chemical emergency advice, Tox Info Suisse", "145"),
        ("MiT Project Manager, Burak Ücöz", "+41 76 202 01 70"),
        ("DPR site management", "As advised at induction"),
        ("Site first aider and assembly point", "As advised at induction")]
t = new_table(len(CONT) + 1, 2, [5871, 3200])
header_row(t, ["Contact", "Number"], align=['left', 'right'])
for i, (a, b) in enumerate(CONT, start=1):
    r = t.rows[i]
    cell_text(r.cells[0], a, size=9.5); vcenter(r.cells[0])
    emph = b.startswith('+') or b.isdigit()
    cell_text(r.cells[1], b, size=10.5 if emph else 9.5, bold=emph,
              color=RED if emph else GREY, align='right'); vcenter(r.cells[1])
    row_height(r, 300)
    if i % 2 == 0:
        shade(r.cells[0], ZEBRA); shade(r.cells[1], ZEBRA)
keep_row_together(t)
spacer()

# ---------------------------------------------------------------- 12
h1("12", "Monitoring and Review")
for b in ["Compliance with this RAMS is monitored by the MiT supervisor throughout every delivery.",
          "Work is stopped immediately if any control cannot be applied as described, and does not resume until the situation is resolved or this document is revised.",
          "Every delivery is recorded in the fuel log, including date, volume, personnel present and any observation.",
          "This RAMS is reviewed before each delivery, on any change of conditions and after any incident or near miss.",
          "DPR may inspect the operation at any time and may request any of the certification listed in section 4."]:
    bullet(b)

# ---------------------------------------------------------------- 13
h1("13", "Briefing Record and Declaration", pagebreak=True)
para("The persons named below have been briefed on the contents of this RAMS before the start of work and confirm that they have understood it. Any deviation from this method statement requires the work to be stopped and this document to be reviewed before the activity continues.")
spacer(80)
SIGN = [("Burak Ücöz", "Project Manager, Mobil in Time AG"),
        ("", "Tanker driver, fuel supplier"),
        ("", "Attendant, Mobil in Time AG"),
        ("", ""),
        ("", "Reviewed for DPR Construction")]
t = new_table(len(SIGN) + 1, 4, [2400, 3400, 1371, 1900])
header_row(t, ["Name", "Function", "Date", "Signature"])
for i, (n, f) in enumerate(SIGN, start=1):
    r = t.rows[i]
    cell_text(r.cells[0], n, size=9.5, bold=bool(n))
    cell_text(r.cells[1], f, size=9.5, color=GREY)
    cell_text(r.cells[2], "", size=9.5)
    cell_text(r.cells[3], "", size=9.5)
    shade(r.cells[2], LGREY)
    shade(r.cells[3], LGREY)
    for c in r.cells:
        vcenter(c)
    row_height(r, 620)
    r._tr.get_or_add_trPr().append(parse_xml('<w:cantSplit %s/>' % nsdecls('w')))
keep_row_together(t)
spacer(100)
para("End of document  ·  RAMS-MIT-ZRH12-004, Rev. 3", size=8.5, color=GREY, justify=False)


# ------------------------------------------------ XML-Reihenfolge normalisieren
ORDER = {
 'pPr': ['pStyle','keepNext','keepLines','pageBreakBefore','framePr','widowControl','numPr',
         'suppressLineNumbers','pBdr','shd','tabs','suppressAutoHyphens','kinsoku','wordWrap',
         'overflowPunct','topLinePunct','autoSpaceDE','autoSpaceDN','bidi','adjustRightInd',
         'snapToGrid','spacing','ind','contextualSpacing','mirrorIndents','suppressOverlap','jc',
         'textDirection','textAlignment','textboxTightWrap','outlineLvl','divId','cnfStyle','rPr',
         'sectPr','pPrChange'],
 'rPr': ['rStyle','rFonts','b','bCs','i','iCs','caps','smallCaps','strike','dstrike','outline',
         'shadow','emboss','imprint','noProof','snapToGrid','vanish','webHidden','color','spacing',
         'w','kern','position','sz','szCs','highlight','u','effect','bdr','shd','fitText',
         'vertAlign','rtl','cs','em','lang','eastAsianLayout','specVanish','oMath'],
 'tcPr': ['cnfStyle','tcW','gridSpan','hMerge','vMerge','tcBorders','shd','noWrap','tcMar',
          'textDirection','tcFitText','vAlign','hideMark'],
 'trPr': ['cnfStyle','divId','gridBefore','gridAfter','wBefore','wAfter','cantSplit','trHeight',
          'tblHeader','tblCellSpacing','jc','hidden'],
 'tblPr': ['tblStyle','tblpPr','tblOverlap','bidiVisual','tblStyleRowBandSize','tblStyleColBandSize',
           'tblW','tblJc','tblCellSpacing','tblInd','tblBorders','shd','tblLayout','tblCellMar',
           'tblLook','tblCaption','tblDescription'],
}

def normalize(root):
    """Sortiert Formatierungselemente in die vom OOXML-Schema geforderte Reihenfolge."""
    for el in root.iter():
        tag = el.tag.split('}')[-1] if isinstance(el.tag, str) else ''
        order = ORDER.get(tag)
        if not order:
            continue
        kids = list(el)
        def key(k):
            name = k.tag.split('}')[-1] if isinstance(k.tag, str) else ''
            return order.index(name) if name in order else len(order)
        for k in sorted(kids, key=key):
            el.append(k)

# ------------------------------------------------------- Dokument-Eigenschaften
cp = doc.core_properties
cp.title = "Risk Assessment & Method Statement \u2013 HVO Refuelling, Vantage ZRH12"
cp.subject = "RAMS-MIT-ZRH12-004, Rev. 3"
cp.author = "Burak \u00dcc\u00f6z, Mobil in Time AG"
cp.last_modified_by = "Burak \u00dcc\u00f6z"
cp.category = "Risk Assessment & Method Statement"
cp.comments = "Prepared for DPR Construction CH AG, Vantage ZRH12, Winterthur"
cp.keywords = "RAMS, HVO, refuelling, ZRH12, Mobil in Time"

def clear_cell_indents(root):
    """Das Format "Standard" der Vorlage hat ind left=567. In Tabellenzellen
    frisst das die Spaltenbreite, deshalb dort Einzug hart auf null setzen."""
    for tc in root.iter(qn('w:tc')):
        for par in tc.iter(qn('w:p')):
            pPr = par.find(qn('w:pPr'))
            if pPr is None:
                pPr = parse_xml('<w:pPr %s/>' % nsdecls('w'))
                par.insert(0, pPr)
            for old_ind in pPr.findall(qn('w:ind')):
                pPr.remove(old_ind)
            pPr.append(parse_xml('<w:ind %s w:left="0" w:right="0" w:firstLine="0"/>' % nsdecls('w')))

clear_cell_indents(doc.element.body)
normalize(doc.element.body)

doc.save(OUT)

# ------------------------------------------- Footer nachbearbeiten (EN + E-Mail)
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in ('word/footer1.xml', 'word/footer2.xml'):
            s = data.decode('utf-8')
            s = s.replace('<w:t xml:space="preserve">Seite </w:t>', '<w:t xml:space="preserve">Page </w:t>')
            s = s.replace('<w:t xml:space="preserve"> von </w:t>', '<w:t xml:space="preserve"> of </w:t>')
            # Reihenfolge wichtig: zuerst die alte Web-Zeile leeren,
            # dann die fremde E-Mail-Adresse durch die Website ersetzen.
            s = s.replace('<w:t>www.mobilintime.com</w:t>', '<w:t xml:space="preserve"> </w:t>')
            s = s.replace('<w:t>kw@zb-heizungen.ch</w:t>', '<w:t>www.mobilintime.com</w:t>')
            s = s.replace('<w:t>Ein Unternehmen der</w:t>', '<w:t>An Aggreko Company</w:t>')
            s = s.replace('<w:t>Mobil in Time Gruppe</w:t>', '<w:t xml:space="preserve"> </w:t>')
            data = s.encode('utf-8')
        zout.writestr(item, data)
os.replace(tmp, OUT)
print("OK ->", OUT, os.path.getsize(OUT), "bytes")
