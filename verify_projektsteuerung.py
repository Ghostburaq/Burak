# -*- coding: utf-8 -*-
"""Funktions- und Integritaets-Pruefer fuer Projektsteuerung_MiT-PWR-WIN-2026-001.xlsx
Prueft JEDE Formel: Funktions-Whitelist, Tabellen-/Blattbezuege, Bereichsgrenzen,
Klammern/Anfuehrungszeichen, plus Kalender-, XML-, Chart- und Reload-Checks."""
import zipfile, re, html, sys, datetime
import xml.dom.minidom as MD
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

F = "/home/user/Burak/Projektsteuerung_MiT-PWR-WIN-2026-001.xlsx"
fails = []
def check(name, ok, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" - {detail}" if detail else ""))
    if not ok: fails.append(name)

wb = load_workbook(F)
z = zipfile.ZipFile(F)

# ---------------------------------------------------------------- Formeln einsammeln
formulas = []   # (blatt, zelle/kontext, formel)
for sn in wb.sheetnames:
    ws = wb[sn]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                formulas.append((sn, c.coordinate, c.value))
    for rng, rules in ws.conditional_formatting._cf_rules.items():
        for rule in rules:
            for fm in (rule.formula or []):
                formulas.append((sn, f"CF@{rng.sqref}", "=" + str(fm)))
print(f"== {len(formulas)} Formeln gefunden (Zellen + bedingte Formatierung) ==\n")

print("1) FUNKTIONS-WHITELIST (jede verwendete Funktion: geht / geht nicht)")
WHITELIST = {  # Basis-Funktionen, in jeder Excel-Version ab 2010 inkl. Browser verfuegbar
    "COUNTIF","COUNTIFS","COUNTA","SUMPRODUCT","TODAY","TEXT","IF","IFERROR",
    "ISNUMBER","SEARCH","INT","SMALL","INDEX","MATCH","ROW","REPT","ROUND",
    "WEEKNUM","DATE","AND","OR","WORKDAY","SUM",
}
used = {}
for sn, coord, f in formulas:
    body = re.sub(r'"[^"]*"', '""', f)  # Strings ausblenden
    for fn in re.findall(r'([A-Z][A-Z0-9]*)\s*\(', body):
        used.setdefault(fn, 0); used[fn] += 1
for fn in sorted(used):
    ok = fn in WHITELIST
    check(f"Funktion {fn} ({used[fn]}x)", ok, "Basis-Funktion, ueberall verfuegbar" if ok else "NICHT in Whitelist!")

print("\n2) TABELLEN-BEZUEGE (Tabelle[Spalte] existiert?)")
tables = {}
for sn in wb.sheetnames:
    for tname, tref in wb[sn].tables.items():
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", str(wb[sn].tables[tname].ref))
        hdr_row = int(m.group(2)); c0 = column_index_from_string(m.group(1)); c1 = column_index_from_string(m.group(3))
        cols = [wb[sn].cell(row=hdr_row, column=c).value for c in range(c0, c1+1)]
        tables[tname] = (sn, cols)
tab_refs = set()
for sn, coord, f in formulas:
    for tname, colname in re.findall(r"([A-Za-z_][A-Za-z0-9_]*)\[([^\]]+)\]", f):
        tab_refs.add((tname, colname))
for tname, colname in sorted(tab_refs):
    ok = tname in tables and colname in tables[tname][1]
    where = tables.get(tname, ("?", []))[0]
    check(f"{tname}[{colname}]", ok, f"Tabelle auf Blatt '{where}'" if ok else "Bezug kaputt!")

print("\n3) BLATT-BEZUEGE (Sheet!Zelle existiert?)")
sheet_refs = set()
for sn, coord, f in formulas:
    for ref_sheet in re.findall(r"([A-Za-z][A-Za-z0-9 &._-]*)!\$?[A-Z]+\$?\d+", f):
        sheet_refs.add(ref_sheet.strip())
for rs in sorted(sheet_refs):
    check(f"Blattbezug '{rs}'", rs in wb.sheetnames)

print("\n4) BEREICHS-GRENZEN (referenzierte Zeilen liegen in den Daten?)")
ms = wb["Meilensteine"]
ms_last = ms.max_row
rng_errs = []
for sn, coord, f in formulas:
    for sheet, col, row in re.findall(r"([A-Za-z][A-Za-z0-9 &._-]*)!\$?([A-Z]+)\$?(\d+)", f):
        sheet = sheet.strip()
        if sheet in wb.sheetnames:
            if int(row) > wb[sheet].max_row + 25:
                rng_errs.append(f"{sn}!{coord}: {sheet}!{col}{row}")
check("Alle Blatt-Bezuege innerhalb der Datenbereiche", not rng_errs, "; ".join(rng_errs[:5]))
helper_ref = f"$L$2:$L${ms_last}"
dash_f = [f for _,_,f in formulas if "Meilensteine!$L$" in f]
check(f"Hilfsspalte Meilensteine L2:L{ms_last} (25 Meilensteine)",
      all(helper_ref in f for f in dash_f) and len(dash_f) >= 15, f"{len(dash_f)} Formeln")

print("\n5) SYNTAX (Klammern & Anfuehrungszeichen balanciert, keine Semikolons, keine dt. Namen)")
syn = []
for sn, coord, f in formulas:
    if f.count("(") != f.count(")"): syn.append(f"{sn}!{coord}: Klammern")
    if f.count('"') % 2: syn.append(f"{sn}!{coord}: Quotes")
    if ";" in re.sub(r'"[^"]*"', '', f): syn.append(f"{sn}!{coord}: Semikolon")
check("Syntax aller Formeln", not syn, "; ".join(syn[:5]))
german = ["ZÄHLENWENN","HEUTE(","WENNFEHLER","KKLEINSTE","VERGLEICH","SUMMENPRODUKT","ANZAHL2","_xludf","AGGREGATE("]
ger = [f"{sn}!{coord}" for sn, coord, f in formulas if any(g in f for g in german)]
check("Keine deutschen Funktionsnamen / _xludf / AGGREGATE", not ger, "; ".join(ger[:5]))

print("\n6) KALENDER 2026 (Wochentage, ISO-KW, Gantt-Kopf)")
WD = ["Mo","Di","Mi","Do","Fr","SA","SO"]
weekend = []
for sheet, dcol in [("Meilensteine",4),("Aufgaben",5),("Lieferungen & Logistik",4)]:
    for r in range(2, wb[sheet].max_row+1):
        d = wb[sheet].cell(row=r,column=dcol).value
        if hasattr(d,"weekday") and d.weekday() >= 5:
            weekend.append(f"{sheet} Z{r}: {d:%d.%m.} {WD[d.weekday()]}")
check("Keine Termine an Wochenenden", not weekend, "; ".join(weekend))
t = wb["Terminplan"]; kw_err = 0; n_weeks = 0
for col in range(5, 60):
    lbl = t.cell(row=2,column=col).value; dat = t.cell(row=3,column=col).value
    if dat is None: break
    n_weeks += 1
    if lbl != f"KW{dat.isocalendar()[1]}" or dat.weekday() != 0: kw_err += 1
check(f"Gantt-Kopf: {n_weeks} Wochen, alle Montage mit korrekter ISO-KW", kw_err == 0 and n_weeks == 25)
n_ms = ms.max_row - 1
diamonds = sum(1 for sn,coord,f in formulas if sn=="Terminplan" and '"◆"' in f)
check(f"Meilenstein-Rauten: {n_ms} Meilensteine x {n_weeks} Wochen = {n_ms*n_weeks} Formeln",
      diamonds == n_ms*n_weeks, f"gefunden: {diamonds}")

print("\n7) DATEI-INTEGRITAET")
xml_err = 0
for n in z.namelist():
    if n.endswith(".xml"):
        try: MD.parseString(z.read(n))
        except Exception: xml_err += 1
check(f"XML wohlgeformt ({sum(1 for n in z.namelist() if n.endswith('.xml'))} Teile)", xml_err == 0)
esz = [n for n in z.namelist() if n.endswith(".xml") and b'\xc3\x9f' in z.read(n)]
check("Kein scharfes S (Schweizer Schreibweise)", not esz)
dv_max = 0
for n in z.namelist():
    if n.startswith("xl/worksheets/"):
        for f1 in re.findall(r"<formula1>(.*?)</formula1>", z.read(n).decode()):
            dv_max = max(dv_max, len(html.unescape(f1)))
check(f"Dropdown-Listen max. {dv_max} Zeichen (Limit 255)", dv_max <= 255)
charts = [n for n in z.namelist() if "charts/chart" in n]
chart_refs_ok = True
for n in charts:
    d = z.read(n).decode()
    for ref in re.findall(r"<f>(.*?)</f>", d):
        shref = ref.split("!")[0].strip("'")
        if shref not in wb.sheetnames: chart_refs_ok = False
check(f"Diagramme ({len(charts)}) referenzieren existierende Blaetter", chart_refs_ok)
for n in sorted(charts):
    d = z.read(n).decode()
    if "barChart" in d: kind = "Balken"
    elif "doughnutChart" in d: kind = "Doughnut"
    elif "lineChart" in d: kind = "Linie"
    else: kind = "?"
    check(f"{kind}: Kategorien als Text (strRef) -> Namen sichtbar", "<strRef>" in d.split("<cat>")[1][:40] if "<cat>" in d else False)
    if kind == "Linie":
        check("Linie: 2 Serien (Aufgaben offen, Meilensteine erledigt)", d.count("<ser>") == 2)
        check("Linie: Achsen sichtbar", d.count('<delete val="0" />') == 2)
        continue
    if kind == "Balken":
        check("Balken: Datenbeschriftung (Werte) aktiv", '<showVal val="1"' in d)
        check("Balken: beide Achsen sichtbar (delete=0)", d.count('<delete val="0" />') == 2)
        check("Balken: Wertachse unten (axPos b)", '<axPos val="b" />' in d)
        check("Balken: Achsenbeschriftung aktiviert (tickLblPos)", d.count('<tickLblPos val="nextTo" />') == 2)
    else:
        check("Doughnut: Prozent-Beschriftung aktiv", '<showPercent val="1"' in d)
        check("Doughnut: ausgeblendete Hilfsdaten werden geplottet (plotVisOnly=0)", '<plotVisOnly val="0" />' in d)
import warnings
warnings.simplefilter("error")
try:
    wb2 = load_workbook(F); reload_ok = True
except Exception as e:
    reload_ok = False
check(f"Strict-Reload ohne Warnungen ({len(wb.sheetnames)} Blaetter)", reload_ok)

print("\n8) INHALTS-STICHPROBEN (V5-Daten)")
kws = wb["Kommerziell"]
check("EUR 310'328.50 (P-640380-3)", kws["D2"].value == 310328.50)
check("EUR 280'511.84 (P-650395-2)", kws["D3"].value == 280511.84)
check("Total-Formel =D2+D3", kws["D4"].value == "=D2+D3")
check("CHF-Spalte Platzhalter", kws["G2"].value == "[Innendienst/ARM]")
kt = wb["Kontakte"]
namen = [kt.cell(row=r,column=2).value for r in range(2, kt.max_row+1)]
for n in ["Roxana Minor","Martin Wendsche","Tanya Hoffmann","WORKcontrol Suisse AG","Olli"]:
    check(f"Kontakt '{n}'", n in namen)
check("WORKcontrol Telefon", any(kt.cell(row=r,column=5).value == "+41 44 512 88 11" for r in range(2, kt.max_row+1)))
mt = wb["Termine & Meetings"]
check("9 Meetings aus Supplemental Conditions", mt.max_row - 1 == 9)
rk = wb["Stopp-Punkte & Risiken"]
check("12 Stopp-Punkte/Risiken mit Wahrscheinlichkeit", rk.max_row - 1 == 12 and rk["E1"].value == "Wahrscheinlichkeit")
ms_names = [ms.cell(row=r,column=2).value for r in range(2, ms.max_row+1)]
check("Meilenstein 'Abruf 3-MW-Block'", "Abruf 3-MW-Block" in ms_names)
check("Meilenstein Closeout 30.11. (15 Tage vor Leistungsende)",
      any("Closeout" in str(n) and ms.cell(row=i+2,column=4).value == datetime.datetime(2026,11,30)
          for i,n in enumerate(ms_names)))
d = wb["Dashboard"]
check("Naechste 5 Meilensteine (5 Zeilen)", sum(1 for _,c,f in formulas if _=="Dashboard" and "SMALL" in f and "INT(" in f) == 5)
# Countdown 'Tage bis naechster Meilenstein' muss auf die ERSTE Zeile des
# 'Naechste Meilensteine'-Blocks zeigen (dynamisch ermittelt, nicht hartkodiert)
first_nx_row = min(int(c[1:]) for s,c,f in formulas
                   if s=="Dashboard" and c.startswith("B") and "SMALL" in f and "INT(" in f)
cd_ref = next((f for s,c,f in formulas if s=="Dashboard" and "-TODAY()" in f and f.startswith('=IFERROR(B')), "")
m = re.search(r"B(\d+)-TODAY", cd_ref)
check(f"Countdown referenziert 1. Meilenstein-Zeile (B{first_nx_row})",
      m is not None and int(m.group(1)) == first_nx_row, f"Formel: {cd_ref}")

print("\n" + "="*60)
if fails:
    print(f"ERGEBNIS: {len(fails)} PRUEFUNG(EN) FEHLGESCHLAGEN:"); [print("  -", x) for x in fails]
    sys.exit(1)
print("ERGEBNIS: ALLE PRUEFUNGEN BESTANDEN - jede Funktion GEHT.")
