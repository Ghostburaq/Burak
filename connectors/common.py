"""Gemeinsame Helfer für alle CRM-Konnektoren."""
from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

import openpyxl


@dataclass
class Row:
    nr: int | None
    prio: str | None
    segment: str | None
    firma: str
    ort: str | None
    plz: str | int | None
    kanton: str | None
    ansprechpartner: str | None
    funktion: str | None
    email: str | None
    telefon: str | None
    website: str | None
    status: str | None
    naechster_schritt: str | None
    bedarf: str | None
    hauptprodukt: str | None
    followup: object | None
    notizen: str | None
    internes: str | None
    wahrscheinlichkeit: float | int | None = None
    wert_chf: float | int | None = None
    letzte_aktivitaet: object | None = None

    def as_dict(self) -> dict:
        return self.__dict__.copy()


def load_rows(path: str, sheet: str | None = None,
              only_prio: str | None = None,
              only_status: str | None = None,
              only_segment: str | None = None) -> list[Row]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet is None:
        for s in ("📋 Kunden-Datenbank", "Kunden-Datenbank"):
            if s in wb.sheetnames:
                sheet = s
                break
        else:
            sheet = wb.sheetnames[0]
    ws = wb[sheet]
    rows: list[Row] = []
    for r in range(4, ws.max_row + 1):
        firma = ws.cell(row=r, column=4).value
        if not firma:
            continue
        vals = [ws.cell(row=r, column=c).value for c in range(1, 23)]
        while len(vals) < 22:
            vals.append(None)
        row = Row(*vals)
        if only_prio and row.prio != only_prio:
            continue
        if only_status and row.status != only_status:
            continue
        if only_segment and row.segment != only_segment:
            continue
        rows.append(row)
    return rows


def parse_args_base(description: str):
    ap = argparse.ArgumentParser(description=description)
    ap.add_argument("file", help="Pfad zur MiT-CRM xlsx-Datei")
    ap.add_argument("--prio", help="Nur Prio (A/B/C)")
    ap.add_argument("--status", help="Nur Status")
    ap.add_argument("--segment", help="Nur Segment")
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--dry-run", action="store_true",
                    help="Keine API-Calls, nur Vorschau")
    return ap
