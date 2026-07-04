#!/usr/bin/env python3
"""Struktur-Test: Formel-Syntax, Cross-Sheet-Referenzen, Validierungen."""
import re, sys
import openpyxl
from openpyxl.formula import Tokenizer

FILE = sys.argv[1] if len(sys.argv) > 1 else "MiT_GESAMTMAPPE_2026.xlsx"
wb = openpyxl.load_workbook(FILE, data_only=False)
sheets = set(wb.sheetnames)
errors = []
n_formulas = 0
ref_re = re.compile(r"'([^']+)'!")

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                n_formulas += 1
                f = c.value
                # Syntax-Check per Tokenizer
                try:
                    Tokenizer(f)
                except Exception as e:
                    errors.append(f"{ws.title}!{c.coordinate}: Tokenizer: {e} :: {f[:80]}")
                # Cross-Sheet-Referenzen prüfen
                for ref in ref_re.findall(f):
                    if ref not in sheets:
                        errors.append(f"{ws.title}!{c.coordinate}: unbekanntes Sheet '{ref}'")
                # Häufige Fehlerquellen
                if ";" in f and '";"' not in f and "';'" not in f:
                    # Semikolon als Argumenttrenner wäre falsch (Speicherformat braucht Komma)
                    if re.search(r"\((?:[^\"()]|\"[^\"]*\")*;", f):
                        errors.append(f"{ws.title}!{c.coordinate}: Semikolon-Trenner? :: {f[:80]}")

# Validierungen prüfen
for ws in wb.worksheets:
    for dv in ws.data_validations.dataValidation:
        if dv.formula1 and dv.formula1.startswith("='"):
            ref = ref_re.findall(dv.formula1)
            for r in ref:
                if r not in sheets:
                    errors.append(f"{ws.title} DV: unbekanntes Sheet '{r}' in {dv.formula1}")

print(f"Sheets: {len(sheets)}, Formeln: {n_formulas}")
if errors:
    print(f"FEHLER: {len(errors)}")
    for e in errors[:40]:
        print(" ", e)
    sys.exit(1)
print("STRUKTUR OK")
