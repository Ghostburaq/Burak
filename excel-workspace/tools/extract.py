#!/usr/bin/env python3
"""Extrahiert alle Quelldaten in ein normalisiertes Pickle für den Workbook-Build.

Strategie:
- data_only=True  -> gecachte Formelwerte (Quelldateien kommen aus echtem Excel)
- Header werden per Suchtext gefunden (robust gegen Titelzeilen)
- Versionsduplikate werden per Schlüssel gemerged (Basis = neueste/vollständigste Datei)
"""
import os, pickle, re, unicodedata
import openpyxl

U = "/root/.claude/uploads/b88840da-82a9-552f-916e-79b3afb44003"
F = {
    "katalog":      f"{U}/1d997271-Aggreko_MiT_Master_Katalog_v1.xlsx",
    "crm_alt":      f"{U}/474b83bc-MiT_CRM_Vorlage_Final.xlsx",
    "crm_neu":      f"{U}/7dee7f2a-MiT_CRM_Vorlage_Final.xlsx",
    "suite":        f"{U}/47b851ef-MiT_Aggreko_DC_Suite_2026.xlsx",
    "monthly":      f"{U}/49d2e6e5-CH_Customer_Monthly_Report_Mai_26.xlsx",
    "preis_intl":   f"{U}/4fdb6b33-MIT_PriceList_Confidential.xlsx",
    "ceo_alt":      f"{U}/5484bfa1-CH_MiT_Strom_Customer_CEO_CFO.xlsx",
    "ceo_neu":      f"{U}/6a0818d2-CH_MiT_Strom_Customer_CEO_CFO.xlsx",
    "preis_chf":    f"{U}/68c35401-202605_MIT_PriceList_CHF.xlsx",
    "rsrg":         f"{U}/822a71f1-RSRG_Kundenanalyse_MiT_2026.xlsx",
    "gesamt":       f"{U}/b977007b-MiT_DC_GESAMTMAPPE_2026_FIXED.xlsx",
    "bau_zsf":      f"{U}/d5c64a29-Zusammenfassung_Bauprojekte.xlsx",
    "strom_cust":   f"{U}/f38e29e4-CH_MiT_Strom_Customer.xlsx",
}

_wb_cache = {}
def wb(key):
    if key not in _wb_cache:
        _wb_cache[key] = openpyxl.load_workbook(F[key], data_only=True, read_only=False)
    return _wb_cache[key]

def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v

def hnorm(v):
    """Header-Normalisierung: Whitespace/Umbrüche zu einem Leerzeichen, Soft-Hyphen raus."""
    if v is None:
        return None
    s = " ".join(str(v).replace("\xad", "").split())
    return s if s else None

def find_header_row(ws, must_contain, max_scan=12):
    """Findet die Zeile, die alle Strings aus must_contain enthält."""
    for r in range(1, max_scan + 1):
        vals = [hnorm(c.value) or "" for c in ws[r]]
        joined = "|".join(vals)
        if all(m in joined for m in must_contain):
            return r, vals
    raise ValueError(f"Header nicht gefunden in '{ws.title}': {must_contain}")

def read_table(ws, must_contain, max_col=None, stop_after_empty=30):
    """Liest Header + Datenzeilen ab der gefundenen Headerzeile."""
    hrow, hvals = find_header_row(ws, must_contain)
    ncol = max_col or ws.max_column
    headers = []
    for c in range(1, ncol + 1):
        h = hnorm(ws.cell(row=hrow, column=c).value)
        headers.append(h)
    # trailing None-Header abschneiden
    while headers and headers[-1] is None:
        headers.pop()
    # doppelte Headernamen eindeutig machen (z.B. 'Ziel' 2x in Akquise)
    seen_h = {}
    for i, h in enumerate(headers):
        if h is None:
            continue
        if h in seen_h:
            seen_h[h] += 1
            headers[i] = f"{h} #{seen_h[h]}"
        else:
            seen_h[h] = 1
    ncol = len(headers)
    rows, empty_run = [], 0
    for r in range(hrow + 1, ws.max_row + 1):
        vals = [clean(ws.cell(row=r, column=c).value) for c in range(1, ncol + 1)]
        if all(v is None for v in vals):
            empty_run += 1
            if empty_run >= stop_after_empty:
                break
            continue
        empty_run = 0
        rows.append(vals)
    return headers, rows

def norm_key(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())

def as_dicts(headers, rows):
    return [dict(zip(headers, r)) for r in rows]

OUT = {}
LOG = []

# ============ 1) PIPELINE (3 Versionen mergen) ============
def get_pipeline(key, sheet):
    ws = wb(key)[sheet]
    h, rows = read_table(ws, ["Kunde / Unternehmen", "Volumen CHF", "Status"])
    # leere Spaltennamen entfernen (Version 5484bfa1 hat eine Leerspalte)
    idx = [i for i, x in enumerate(h) if x is not None]
    h2 = [h[i] for i in idx]
    rows2 = [[r[i] for i in idx] for r in rows]
    # nur Zeilen mit Kundennamen
    ki = h2.index("Kunde / Unternehmen")
    rows2 = [r for r in rows2 if r[ki]]
    return h2, rows2

p_new_h, p_new = get_pipeline("ceo_neu", "MiT Strom Pipeline")
p_old_h, p_old = get_pipeline("ceo_alt", "MiT Strom Pipeline")
p_mon_h, p_mon = get_pipeline("monthly", "MiT Strom Pipeline")
LOG.append(f"Pipeline Zeilen: ceo_neu={len(p_new)} ceo_alt={len(p_old)} monthly={len(p_mon)}")
LOG.append(f"Pipeline Header ceo_neu ({len(p_new_h)}): {p_new_h}")
LOG.append(f"Pipeline Header monthly ({len(p_mon_h)}): {p_mon_h}")

def pipe_key(h, row):
    d = dict(zip(h, row))
    return (norm_key(d.get("Kunde / Unternehmen")), norm_key(d.get("Leistung / Fleet")),
            str(d.get("Start") or ""))

# Basis = Version mit den meisten Zeilen; andere anhängen falls Schlüssel fehlt
candidates = [("ceo_neu", p_new_h, p_new), ("monthly", p_mon_h, p_mon), ("ceo_alt", p_old_h, p_old)]
candidates.sort(key=lambda t: len(t[2]), reverse=True)
base_name, base_h, base_rows = candidates[0]
seen = {pipe_key(base_h, r) for r in base_rows}
merged = [dict(zip(base_h, r)) for r in base_rows]
for name, h, rows in candidates[1:]:
    add = 0
    for r in rows:
        k = pipe_key(h, r)
        if k not in seen:
            seen.add(k)
            merged.append(dict(zip(h, r)))
            add += 1
    LOG.append(f"Pipeline-Merge: +{add} aus {name}")
OUT["pipeline"] = {"basis": base_name, "header": base_h, "rows": merged}
LOG.append(f"Pipeline gesamt: {len(merged)} (Basis {base_name})")

# ============ 2) KUNDEN-CRM (2 Versionen) ============
c_new_h, c_new = read_table(wb("crm_neu")["📋 Kunden-Datenbank"], ["Firmenname", "Segment", "Prio"])
c_old_h, c_old = read_table(wb("crm_alt")["📋 Kunden-Datenbank"], ["Firmenname", "Segment", "Prio"])
fi = c_new_h.index("Firmenname *")
c_new = [r for r in c_new if r[fi]]
c_old = [r for r in c_old if r[c_old_h.index("Firmenname *")]]
seen = {norm_key(r[fi]) for r in c_new}
crm = [dict(zip(c_new_h, r)) for r in c_new]
add = 0
for r in c_old:
    k = norm_key(r[c_old_h.index("Firmenname *")])
    if k not in seen:
        seen.add(k); crm.append(dict(zip(c_old_h, r))); add += 1
OUT["crm"] = {"header": c_new_h, "rows": crm}
LOG.append(f"CRM: neu={len(c_new)} alt={len(c_old)} merge+{add} gesamt={len(crm)}")

# ============ 3) KUNDENKARTEI (Monthly 2000 + strom_cust 499) ============
k_mon_h, k_mon = read_table(wb("monthly")["CH Customer Overview"], ["Kunde / Unternehmen", "Ansprechperson"])
k_str_h, k_str = read_table(wb("strom_cust")["Kundenkartei"], ["Kunde / Unternehmen", "Ansprechperson"])
ki = k_mon_h.index("Kunde / Unternehmen")
k_mon = [r for r in k_mon if r[ki]]
k_str = [r for r in k_str if r[k_str_h.index("Kunde / Unternehmen")]]
seen = {norm_key(r[ki]) for r in k_mon}
kartei = [dict(zip(k_mon_h, r)) for r in k_mon]
add = 0
for r in k_str:
    k = norm_key(r[k_str_h.index("Kunde / Unternehmen")])
    if k not in seen:
        seen.add(k); kartei.append(dict(zip(k_str_h, r))); add += 1
OUT["kartei"] = {"header": k_mon_h, "rows": kartei}
LOG.append(f"Kartei: monthly={len(k_mon)} strom={len(k_str)} merge+{add} gesamt={len(kartei)}")

# ============ 4) DC BETREIBER ============
b_s_h, b_s = read_table(wb("suite")["03 CH BETREIBER"], ["Betreiber", "IT-MW"])
b_g_h, b_g = read_table(wb("gesamt")["05 DC_BETREIBER"], ["Betreiber", "IT-MW"])
bi = b_s_h.index("Betreiber")
b_s = [r for r in b_s if r[bi]]
seen = {norm_key(r[bi]) for r in b_s}
betr = [dict(zip(b_s_h, r)) for r in b_s]
add = 0
for r in b_g:
    if r[b_g_h.index("Betreiber")] and norm_key(r[b_g_h.index("Betreiber")]) not in seen:
        seen.add(norm_key(r[b_g_h.index("Betreiber")]))
        betr.append(dict(zip(b_g_h, r))); add += 1
OUT["betreiber"] = {"header": b_s_h, "rows": betr}
LOG.append(f"Betreiber: suite={len(b_s)} gesamt={len(b_g)} merge+{add} total={len(betr)}")

# ============ 5) DC BAUPROJEKTE (3 Quellen) ============
g_h, g_rows = read_table(wb("gesamt")["02 BAUPROJEKTE_SCHWEIZ"], ["Projekt / Ort", "Bauherr"])
z_h, z_rows = read_table(wb("bau_zsf")["Tabelle1"], ["Projekt / Ort", "Inhaber"])
s_h, s_rows = read_table(wb("suite")["04 CH PIPELINE"], ["Projekt / Betreiber", "IT-MW"])
OUT["bau_gesamt"] = {"header": g_h, "rows": as_dicts(g_h, g_rows)}
OUT["bau_zsf"] = {"header": z_h, "rows": as_dicts(z_h, z_rows)}
OUT["bau_suite"] = {"header": s_h, "rows": as_dicts(s_h, s_rows)}
LOG.append(f"Bauprojekte: gesamt={len(g_rows)} zsf={len(z_rows)} suite={len(s_rows)}")
LOG.append(f"  bau_gesamt Header: {g_h}")
LOG.append(f"  bau_zsf Header: {z_h}")
LOG.append(f"  bau_suite Header: {s_h}")

# ============ 6) DC STANDORTE ============
st_g_h, st_g = read_table(wb("gesamt")["06 STANDORTE_ADRESSEN"], ["Code", "Betreiber", "PLZ"])
st_s_h, st_s = read_table(wb("suite")["12 CH STANDORTE"], ["Code", "Betreiber", "PLZ"])
ci = st_g_h.index("Code")
st_g = [r for r in st_g if r[ci]]
# Dedup-Schlüssel: Betreiber+Ort (Codes variieren zwischen Dateien: ZUR1-3 vs ZUR1/2/3)
def st_key(h, r):
    d = dict(zip(h, r))
    return (norm_key(d.get("Betreiber")), norm_key(d.get("Ort")))
seen = {st_key(st_g_h, r) for r in st_g}
standorte = [dict(zip(st_g_h, r)) for r in st_g]
add = 0
KEEP_CODES = {"kilf"}  # KI-RZ Laufenburg = eigenes Projekt neben TZL
for r in st_s:
    code = r[st_s_h.index("Code")]
    if not code:
        continue
    k = st_key(st_s_h, r)
    if k not in seen or norm_key(code) in KEEP_CODES:
        seen.add(k)
        standorte.append(dict(zip(st_s_h, r))); add += 1
OUT["standorte"] = {"header": st_g_h, "rows": standorte}
LOG.append(f"Standorte: gesamt={len(st_g)} suite={len(st_s)} merge+{add} total={len(standorte)}")
LOG.append(f"  standorte Header g: {st_g_h}")
LOG.append(f"  standorte Header s: {st_s_h}")

# ============ 7) DC KONTAKTE ============
kt_g_h, kt_g = read_table(wb("gesamt")["07 KONTAKTE_CRM"], ["Unternehmen", "Funktion"])
kt_s_h, kt_s = read_table(wb("suite")["05 CH CRM"], ["Unternehmen", "Funktion"])
def kt_key(h, r):
    d = dict(zip(h, r))
    return (norm_key(d.get("Unternehmen")), norm_key(d.get("Name")))
kt_g = [r for r in kt_g if r[kt_g_h.index("Unternehmen")]]
seen = {kt_key(kt_g_h, r) for r in kt_g}
dckontakte = [dict(zip(kt_g_h, r)) for r in kt_g]
add = 0
for r in kt_s:
    if r[kt_s_h.index("Unternehmen")] and kt_key(kt_s_h, r) not in seen:
        seen.add(kt_key(kt_s_h, r))
        dckontakte.append(dict(zip(kt_s_h, r))); add += 1
OUT["dc_kontakte"] = {"rows": dckontakte}
LOG.append(f"DC-Kontakte: gesamt={len(kt_g)} suite={len(kt_s)} merge+{add} total={len(dckontakte)}")
LOG.append(f"  kontakte Header g: {kt_g_h}")
LOG.append(f"  kontakte Header s: {kt_s_h}")

# ============ 8) GLOBAL PROJEKTE + KONTAKTE ============
gp_h, gp = read_table(wb("suite")["13 GLOBAL PROJEKTE"], ["Projektname", "Sektor", "Land"])
gp = [r for r in gp if r[gp_h.index("Projektname")]]
OUT["global_projekte"] = {"header": gp_h, "rows": gp}
LOG.append(f"Global Projekte: {len(gp)} Header: {gp_h}")

gk_h, gk = read_table(wb("suite")["15 GLOBAL KONTAKTE"], ["Projekt-ID", "Nachname", "Firma"])
gk = [r for r in gk if any(r)]
OUT["global_kontakte"] = {"header": gk_h, "rows": gk}
LOG.append(f"Global Kontakte: {len(gk)} Header: {gk_h}")

# ============ 9) KATALOG ============
ka_h, ka = read_table(wb("katalog")["🔍 ALLE PRODUKTE"], ["ITEM CODE / MOVEX", "KATEGORIE"])
ka = [r for r in ka if r[ka_h.index("BEZEICHNUNG")]]
OUT["katalog"] = {"header": ka_h, "rows": ka}
LOG.append(f"Katalog: {len(ka)}")

# ============ 10) PREISLISTEN ============
pc_h, pc = read_table(wb("preis_chf")["Sheet1"], ["Product_Line__c", "Tagespreis"])
pc = [r for r in pc if r[pc_h.index("Produktname")]]
OUT["preis_chf"] = {"header": pc_h, "rows": pc}
LOG.append(f"Preisliste CHF: {len(pc)} Header: {pc_h}")

pi_h, pi = read_table(wb("preis_intl")["Sheet1"], ["Generic_Code__c", "Weekly Floor LC"])
pi = [r for r in pi if r[pi_h.index("Generic_Code__c")]]
# führende leere Spalte entfernen
if pi_h[0] is None:
    pi_h = pi_h[1:]; pi = [r[1:] for r in pi]
OUT["preis_intl"] = {"header": pi_h, "rows": pi}
LOG.append(f"Preisliste INTL: {len(pi)} Header: {pi_h}")

# ============ 11) NORMEN / MARKTVOLUMEN / AKQUISE ============
n_h, n = read_table(wb("suite")["11 NORMEN"], ["Norm / Standard", "Kategorie"])
OUT["normen"] = {"header": n_h, "rows": n}
LOG.append(f"Normen: {len(n)}")

mv_h, mv = read_table(wb("suite")["10 MARKTVOLUMEN"], ["Produktkategorie", "Marktanteil %"])
mv = [r for r in mv if r[0]]
OUT["marktvolumen"] = {"header": mv_h, "rows": mv}
LOG.append(f"Marktvolumen: {len(mv)} Header: {mv_h}")

aq_s_h, aq_s = read_table(wb("suite")["06 CH AKQUISE 90T"], ["Woche", "Aktion", "Kanal"])
aq_g_h, aq_g = read_table(wb("gesamt")["10 AKQUISE_90TAGE"], ["Woche", "Aktion", "Kanal"])
def aq_key(h, r):
    d = dict(zip(h, r))
    return (norm_key(d.get("Woche")), norm_key(d.get("Aktion"))[:40])
aq_s = [r for r in aq_s if r[0]]
seen = {aq_key(aq_s_h, r) for r in aq_s}
akquise = [dict(zip(aq_s_h, r)) for r in aq_s]
add = 0
for r in aq_g:
    if r[0] and aq_key(aq_g_h, r) not in seen:
        seen.add(aq_key(aq_g_h, r))
        akquise.append(dict(zip(aq_g_h, r))); add += 1
OUT["akquise"] = {"rows": akquise}
LOG.append(f"Akquise: suite={len(aq_s)} gesamt={len(aq_g)} merge+{add} total={len(akquise)}")
LOG.append(f"  akquise Header s: {aq_s_h}")
LOG.append(f"  akquise Header g: {aq_g_h}")

# ============ 12) LOOKUPS (Land -> Aggreko Region) ============
ws = wb("suite")["16 LOOKUPS"]
land_region = []
for r in range(5, ws.max_row + 1):
    land = clean(ws.cell(row=r, column=1).value)
    reg = clean(ws.cell(row=r, column=4).value)  # Spalte D = grobe Region (EUROPE etc.)
    if land and reg:
        land_region.append((land, reg))
OUT["land_region"] = land_region
LOG.append(f"Land->Region: {len(land_region)}")

# ============ 13) Block-Kopien (RSRG, Dossiers, Systeme, Monatsreport-Layout) ============
def copy_block(key, sheet, max_r=None, max_c=None):
    ws = wb(key)[sheet]
    mr = min(ws.max_row, max_r or ws.max_row)
    mc = min(ws.max_column, max_c or ws.max_column)
    block = []
    for r in range(1, mr + 1):
        row = [clean(ws.cell(row=r, column=c).value) for c in range(1, mc + 1)]
        block.append(row)
    # trailing komplett leere Zeilen entfernen
    while block and all(v is None for v in block[-1]):
        block.pop()
    return block

OUT["rsrg_profil"] = copy_block("rsrg", "01 Firmenprofil")
OUT["rsrg_matrix"] = copy_block("rsrg", "02 Opportunity Matrix")
OUT["implenia"] = copy_block("gesamt", "03 IMPLENIA_PARTNER")
OUT["flexbase"] = copy_block("gesamt", "04 FLEXBASE_TZL_LAUFENBURG")
OUT["systeme"] = copy_block("suite", "07 SYSTEME")
LOG.append(f"Blöcke: rsrg={len(OUT['rsrg_profil'])}+{len(OUT['rsrg_matrix'])} implenia={len(OUT['implenia'])} flexbase={len(OUT['flexbase'])} systeme={len(OUT['systeme'])}")

# ============ 14) Distinct-Listen für Validierung ============
def distinct(dicts, key):
    vals = sorted({str(d.get(key)).strip() for d in dicts if d.get(key) not in (None, "")})
    return vals

OUT["listen"] = {
    "pipeline_status": distinct(OUT["pipeline"]["rows"], "Status"),
    "pipeline_segment": distinct(OUT["pipeline"]["rows"], "Segment"),
    "pipeline_akquise": distinct(OUT["pipeline"]["rows"], "Akquise Typ"),
    "crm_segment": distinct(OUT["crm"]["rows"], "Segment"),
    "crm_status": distinct(OUT["crm"]["rows"], "Status"),
    "crm_prio": distinct(OUT["crm"]["rows"], "Prio"),
    "crm_produkt": distinct(OUT["crm"]["rows"], "Hauptprodukt"),
    "kartei_segment": distinct(OUT["kartei"]["rows"], "Segment"),
    "kantone": distinct(OUT["pipeline"]["rows"], "Kt."),
    "katalog_kategorien": sorted({r[ka_h.index("KATEGORIE")] for r in ka if r[ka_h.index("KATEGORIE")]}),
}
for k, v in OUT["listen"].items():
    LOG.append(f"Liste {k} ({len(v)}): {v[:15]}")

with open("extracted.pkl", "wb") as f:
    pickle.dump(OUT, f)
print("\n".join(LOG))
print("\nOK -> extracted.pkl")
